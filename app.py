import os, json, uuid, time, threading, requests, functools, sys, re, secrets, tempfile, copy
import math
import tos
from datetime import datetime, timezone
from urllib.parse import quote, urlparse, parse_qs
from flask import Flask, request, jsonify, send_from_directory, redirect, session, Response, has_request_context
from werkzeug.utils import secure_filename
from werkzeug.security import generate_password_hash, check_password_hash
from volcenginesdkarkruntime import Ark

BASE = os.path.dirname(os.path.abspath(__file__))

def load_env_file(path):
    if not os.path.exists(path):
        return
    with open(path, 'r', encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith('#') or '=' not in line:
                continue
            key, value = line.split('=', 1)
            os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))

load_env_file(os.path.join(BASE, '.env'))

DATABASE_URL = os.environ.get('DATABASE_URL', '').strip()
try:
    import psycopg
    from psycopg.types.json import Jsonb
except Exception as e:
    print(f'[Postgres] driver unavailable: {e}', flush=True)
    psycopg = None
    Jsonb = None

app = Flask(__name__, static_folder='static', static_url_path='/static')
app.secret_key = os.environ.get('APP_SECRET_KEY', os.urandom(24).hex())

# ── Upload size limit (default 512MB) ────────────────────────
MAX_UPLOAD_MB = int(os.environ.get('MAX_UPLOAD_MB', '512'))
app.config['MAX_CONTENT_LENGTH'] = MAX_UPLOAD_MB * 1024 * 1024

UPLOAD   = os.environ.get('MANGA_UPLOAD_DIR', os.path.join(BASE, 'static', 'uploads'))
DATA     = os.environ.get('MANGA_DATA_DIR', os.path.join(BASE, 'data'))
os.makedirs(UPLOAD, exist_ok=True)
os.makedirs(DATA,   exist_ok=True)

# ── Secrets (all from environment, no hardcoded fallbacks) ───
def _require_env(key):
    val = os.environ.get(key, '')
    if not val:
        raise RuntimeError(f'缺少环境变量 {key}，请参考 .env.example 配置')
    return val

ARK_API_KEY = os.environ.get('ARK_API_KEY', '')
TOS_AK     = os.environ.get('TOS_AK', '')
TOS_SK     = os.environ.get('TOS_SK', '')
NANO_GPT_API_KEY = os.environ.get('NANO_GPT_API_KEY', '')
AGNES_API_KEY = os.environ.get('AGNES_API_KEY', '')

# TOS config (non-secret)
TOS_ENDPOINT = "tos-cn-beijing.volces.com"
TOS_REGION   = "cn-beijing"
TOS_BUCKET   = "movie1"
TOS_PUBLIC_BASE = f"https://{TOS_BUCKET}.{TOS_ENDPOINT}"

# Cloudflare R2 config. If these are set, new media is saved to R2 first.
R2_ACCOUNT_ID = os.environ.get('R2_ACCOUNT_ID', '')
R2_ACCESS_KEY_ID = os.environ.get('R2_ACCESS_KEY_ID', '')
R2_SECRET_ACCESS_KEY = os.environ.get('R2_SECRET_ACCESS_KEY', '')
R2_BUCKET = os.environ.get('R2_BUCKET', 'image-web-storage')
R2_PUBLIC_BASE = os.environ.get('R2_PUBLIC_BASE', '').rstrip('/')
MODEL_ID = "doubao-seedance-2-0-fast-260128"
TEXT_MODEL_ID = "doubao-seed-character-251128"

# Script text models (for brainstorm + split)
SCRIPT_MODELS = {
    "doubao": {"label": "豆包", "provider": "volcengine", "model_id": "doubao-seed-character-251128"},
    "glm46": {"label": "GPT-4.1 Mini", "provider": "nano", "model_id": "openai/gpt-4.1-mini"},
    "claude46": {"label": "Claude4.6", "provider": "nano", "model_id": "anthropic/claude-opus-4.6"}
}
SCRIPT_MODEL_DEFAULT = "doubao"

# Nano-GPT adapter
NANO_GPT_BASE = "https://nano-gpt.com/api/v1"
NANO_GPT_MODELS = {
    "kling-v30-std": "kling-v30-std",
    "grok-imagine-video": "grok-imagine-video-reference-to-video",
    "vidu-q3": "vidu-q3",
    "seedance-v15-pro": "bytedance-seedance-v1.5-pro",
}
NANO_GPT_NAMES = set(NANO_GPT_MODELS.keys())

# Third-party generic adapter
THIRD_PARTY_API_BASE = os.environ.get("THIRD_PARTY_API_BASE", "")
THIRD_PARTY_API_KEY = os.environ.get("THIRD_PARTY_API_KEY", "")
THIRD_PARTY_MODEL_ID = "third-party"

AGNES_API_BASE = "https://apihub.agnes-ai.com/v1"
AGNES_VIDEO_MODEL_ID = "agnes-video-v2.0"
AGNES_IMAGE_MODEL_ID = "agnes-image-2.1-flash"
MINIMAX_VIDEO_MODEL_ID = "MiniMax-H3"
ATLAS_API_BASE = "https://api.atlascloud.ai/api/v1"

ALL_MODELS = ["seedance", AGNES_VIDEO_MODEL_ID] + sorted(NANO_GPT_NAMES) + ([THIRD_PARTY_MODEL_ID] if THIRD_PARTY_API_KEY or THIRD_PARTY_API_BASE else [])

# Model capabilities
MODEL_CAPS = {
    "seedance": {"supports_first_frame": True, "supports_last_frame": True, "supports_reference_images": True, "supports_reference_audio": True, "supports_reference_video": True, "resolutions": ["480p","720p"]},
    "kling-v30-std": {"supports_first_frame": True, "supports_last_frame": False, "supports_reference_images": True, "supports_reference_audio": False, "supports_reference_video": True, "resolutions": ["720p"]},
    "grok-imagine-video": {"supports_first_frame": False, "supports_last_frame": False, "supports_reference_images": True, "supports_reference_audio": False, "supports_reference_video": False, "resolutions": ["480p","720p","1080p"]},
    "vidu-q3": {"supports_first_frame": True, "supports_last_frame": False, "supports_reference_images": True, "supports_reference_audio": False, "supports_reference_video": False, "resolutions": ["480p","720p","1080p"]},
    "seedance-v15-pro": {
        "supports_first_frame": True,
        "supports_last_frame": False,
        "supports_reference_images": False,
        "supports_reference_audio": False,
        "supports_reference_video": False,
        "resolutions": ["480p", "720p", "1080p"],
        "ratios": ["21:9", "16:9", "4:3", "1:1", "3:4", "9:16"],
        "min_duration": 4,
        "max_duration": 12,
    },
    AGNES_VIDEO_MODEL_ID: {
        "supports_first_frame": False,
        "supports_last_frame": False,
        "supports_reference_images": False,
        "supports_reference_audio": False,
        "supports_reference_video": False,
        "resolutions": ["480p", "720p", "1080p"],
        "ratios": ["16:9", "9:16", "1:1", "4:3", "3:4"],
        "min_duration": 4,
        "max_duration": 15,
    },
    MINIMAX_VIDEO_MODEL_ID: {
        "supports_first_frame": True,
        "supports_last_frame": True,
        "supports_reference_images": True,
        "supports_reference_audio": True,
        "supports_reference_video": True,
        "resolutions": ["768p", "2K"],
        "ratios": ["21:9", "16:9", "4:3", "1:1", "3:4", "9:16"],
        "min_duration": 4,
        "max_duration": 15,
    },
    THIRD_PARTY_MODEL_ID: {"supports_first_frame": False, "supports_last_frame": False, "supports_reference_images": True, "supports_reference_audio": True, "supports_reference_video": True, "resolutions": ["480p","720p","1080p"]},
}

# Image generation configs
NANO_GPT_MIDJOURNEY_MODEL_ID = "midjourney/text-to-image"
NANO_GPT_MIDJOURNEY_VERSION = "8.2"
NANO_GPT_IMAGE_MODELS = {"gpt-image-2", "nano-banana-2", NANO_GPT_MIDJOURNEY_MODEL_ID}
LEGACY_IMAGE_MODEL_ALIASES = {"midjourney": NANO_GPT_MIDJOURNEY_MODEL_ID}
VOLC_IMAGE_MODEL_ID = "doubao-seedream-4-5-251128"
VOLC_SEEDREAM_5_PRO_MODEL_ID = "doubao-seedream-5-0-pro-260628"
VOLC_IMAGE_MODELS = {
    "volc-seedream-4-5": VOLC_IMAGE_MODEL_ID,
    VOLC_SEEDREAM_5_PRO_MODEL_ID: VOLC_SEEDREAM_5_PRO_MODEL_ID,
}
ALL_IMAGE_MODELS = sorted(NANO_GPT_IMAGE_MODELS) + [AGNES_IMAGE_MODEL_ID] + list(VOLC_IMAGE_MODELS)
IMAGE_RATIOS = ["1:1", "2:3", "3:2", "3:4", "4:3", "9:16", "16:9", "4:5", "5:4", "custom"]
# Ratio → pixel size for Nano models (moderate sizes)
RATIO_TO_SIZE_NANO = {
    "1:1": "1024x1024", "2:3": "1024x1536", "3:2": "1536x1024",
    "3:4": "1536x2048", "4:3": "2048x1536",
    "9:16": "768x1344", "16:9": "1344x768",
    "4:5": "1536x1920", "5:4": "1920x1536"
}
RATIO_TO_SIZE_AGNES = {
    "1:1": "1024x1024", "2:3": "768x1152", "3:2": "1152x768",
    "3:4": "768x1024", "4:3": "1024x768", "9:16": "768x1365",
    "16:9": "1365x768", "4:5": "819x1024", "5:4": "1024x819"
}
# GPT Image style models sometimes ignore uncommon size strings and fall back
# to 1024x1024, so use stricter/common dimensions plus aspect_ratio below.
RATIO_TO_SIZE_GPT_IMAGE = {
    "1:1": "1024x1024", "2:3": "1024x1536", "3:2": "1536x1024",
    "3:4": "1024x1365", "4:3": "1365x1024",
    "9:16": "864x1536", "16:9": "1536x864",
    "4:5": "1024x1280", "5:4": "1280x1024"
}
# Ratio → pixel size for Seedream (requires ≥3.6M pixels)
RATIO_TO_SIZE_VOLC = {
    "1:1": "1920x1920", "2:3": "2048x3072", "3:2": "3072x2048",
    "3:4": "1728x2304", "4:3": "2304x1728",
    "9:16": "1440x2560", "16:9": "2560x1440",
    "4:5": "1728x2160", "5:4": "2160x1728"
}
# Seedream 5.0 Pro accepts the same Ark image API, but caps output at
# 4,624,220 pixels. Keep each preset inside that model-specific limit.
RATIO_TO_SIZE_VOLC_SEEDREAM_5_PRO = {
    "1:1": "2048x2048", "2:3": "1664x2496", "3:2": "2496x1664",
    "3:4": "1792x2304", "4:3": "2304x1792",
    "9:16": "1152x2048", "16:9": "2048x1152",
    "4:5": "1792x2240", "5:4": "2240x1792"
}
DEFAULT_RATIO = "1:1"

QUALITY_PROMPT = """【画面质量强制要求】
1. 人体结构：四肢、手指、面部必须保持真实比例，严禁肢体扭曲、穿模变形。
2. 服装：衣物完整贴合身体，严禁穿透或消失。
3. 声音：包含符合场景的环境音，严禁无声输出。
4. 画幅：严格9:16竖屏，人物主体居中，不得裁切头部或脚部。
5. 连贯性：画面流畅，严禁跳帧闪烁。
6. 人物一致性：全程保持参考图脸部特征、发型、肤色不变。
7. 肤色：严格还原参考图正常肤色，严禁绿色、灰色等异常颜色叠加。
"""

JOBS = {}
SCRIPT_JOBS = {}
JOB_OWNERS = {}
HISTORY_LOCK = threading.Lock()

_postgres_ready = False
_pg_pool = None
_pg_pool_lock = threading.Lock()
_pg_down_until = 0.0

def data_key(path):
    return os.path.relpath(path, DATA).replace(os.sep, '/').replace('.json', '')

def _pg_pool_get():
    """Lazily create a shared Postgres connection pool (thread-safe).

    Replaces the old per-call psycopg.connect() — opening a fresh TCP+TLS
    connection for every load_json/save_json was the dominant latency in
    project delete / open flows (100-400ms per handshake on Railway).
    """
    global _pg_pool, _postgres_ready, _pg_down_until
    if _pg_pool is not None:
        return _pg_pool
    if not DATABASE_URL or not psycopg:
        return None
    if time.time() < _pg_down_until:
        return None
    with _pg_pool_lock:
        if _pg_pool is not None:
            return _pg_pool
        if time.time() < _pg_down_until:
            return None
        try:
            from psycopg_pool import ConnectionPool
            pool = ConnectionPool(
                DATABASE_URL,
                min_size=1,
                max_size=8,
                open=False,
                kwargs={'connect_timeout': 10},
            )
            pool.open(wait=False, timeout=10)
            with pool.connection(timeout=10) as conn:
                conn.execute('''
                    CREATE TABLE IF NOT EXISTS kv_store (
                        key TEXT PRIMARY KEY,
                        data JSONB NOT NULL,
                        updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
                    )
                ''')
            _pg_pool = pool
            _postgres_ready = True
            print('[Postgres] connection pool ready', flush=True)
        except Exception as e:
            print(f'[Postgres] pool unavailable, using /app/data: {e}', flush=True)
            _pg_pool = None
            _postgres_ready = False
            _pg_down_until = time.time() + 60  # backoff before retrying the DB
    return _pg_pool

def _pg_drop_pool(reason):
    global _pg_pool, _postgres_ready, _pg_down_until
    with _pg_pool_lock:
        if _pg_pool is not None:
            try:
                _pg_pool.close()
            except Exception:
                pass
            _pg_pool = None
        _postgres_ready = False
        _pg_down_until = time.time() + 30
        print(f'[Postgres] {reason}, re-init after backoff', flush=True)

def _pg_conn():
    """Return a pooled connection (as a context manager) or None."""
    pool = _pg_pool_get()
    if pool is None:
        return None
    try:
        return pool.connection(timeout=3)
    except Exception as e:
        _pg_drop_pool(f'pool acquire failed: {e}')
        return None

def init_postgres():
    _pg_pool_get()
    return _postgres_ready

def postgres_save(key, data, overwrite=True):
    if not _postgres_ready:
        return False
    conflict = 'DO UPDATE SET data = EXCLUDED.data, updated_at = NOW()' if overwrite else 'DO NOTHING'
    conn = _pg_conn()
    if conn is None:
        return False
    try:
        # pool.connection() returns a context manager; "with conn as c"
        # yields the real connection and returns it to the pool on exit.
        with conn as c:
            c.execute(
                f'INSERT INTO kv_store (key, data) VALUES (%s, %s) ON CONFLICT (key) {conflict}',
                (key, Jsonb(data))
            )
        return True
    except Exception as e:
        print(f'[Postgres] write failed for {key}: {e}', flush=True)
        return False

def postgres_save_many(items):
    """Upsert several independent JSON documents in one database round trip."""
    if not _postgres_ready or not items:
        return False
    conn = _pg_conn()
    if conn is None:
        return False
    placeholders = ', '.join(['(%s, %s)'] * len(items))
    params = []
    for key, data in items:
        params.extend((key, Jsonb(data)))
    try:
        with conn as c:
            c.execute(
                f'INSERT INTO kv_store (key, data) VALUES {placeholders} '
                'ON CONFLICT (key) DO UPDATE SET data = EXCLUDED.data, updated_at = NOW()',
                params
            )
        return True
    except Exception as e:
        print(f'[Postgres] batch write failed: {e}', flush=True)
        return False

def migrate_local_json_to_postgres():
    if not _postgres_ready:
        return {'scanned': 0, 'imported': 0}
    scanned = imported = 0
    for root, _, files in os.walk(DATA):
        for filename in files:
            if not filename.endswith('.json'):
                continue
            path = os.path.join(root, filename)
            try:
                with open(path, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                scanned += 1
                if postgres_save(data_key(path), data, overwrite=False):
                    imported += 1
            except Exception as e:
                print(f'[Postgres] migration skipped {path}: {e}', flush=True)
    print(f'[Postgres] local bootstrap scanned={scanned} imported={imported}', flush=True)
    return {'scanned': scanned, 'imported': imported}

_JSON_CACHE = {}
_JSON_CACHE_LOCK = threading.Lock()

def _json_cache_get(path):
    with _JSON_CACHE_LOCK:
        cached = _JSON_CACHE.get(os.path.abspath(path))
        return copy.deepcopy(cached) if cached is not None else None

def _json_cache_set(path, data):
    with _JSON_CACHE_LOCK:
        _JSON_CACHE[os.path.abspath(path)] = copy.deepcopy(data)

def load_json(path, default):
    """Load from Postgres first, falling back to the local JSON copy."""
    cached = _json_cache_get(path)
    if cached is not None:
        return cached
    key = os.path.relpath(path, DATA).replace(os.sep, '/').replace('.json', '')
    if _postgres_ready:
        conn = _pg_conn()
        if conn is not None:
            try:
                with conn as c:
                    row = c.execute('SELECT data FROM kv_store WHERE key = %s', (key,)).fetchone()
                if row:
                    _json_cache_set(path, row[0])
                    return copy.deepcopy(row[0])
            except Exception as e:
                print(f'[Postgres] read failed for {key}, using /app/data: {e}', flush=True)
    try:
        with open(path, 'r', encoding='utf-8') as f:
            data = json.load(f)
    except Exception:
        data = default
    _json_cache_set(path, data)
    return copy.deepcopy(data)

def _save_json_local(path, data):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    _json_cache_set(path, data)

def save_json(path, data):
    """Write the local recovery copy, then upsert the Postgres primary copy."""
    key = data_key(path)
    _save_json_local(path, data)
    postgres_save(key, data)

def save_json_many(items):
    """Write recovery copies, then persist all documents atomically in Postgres."""
    prepared = []
    for path, data in items:
        _save_json_local(path, data)
        prepared.append((data_key(path), data))
    postgres_save_many(prepared)

def current_user_id():
    return session.get('user_id')

def user_data_path(name, user_id=None):
    user_id = user_id or current_user_id()
    if not user_id:
        raise RuntimeError('unauthorized')
    folder = os.path.join(DATA, 'users', secure_filename(user_id))
    os.makedirs(folder, exist_ok=True)
    return os.path.join(folder, f'{name}.json')

def characters_path(user_id=None): return user_data_path('characters', user_id)
def assets_path(cat, user_id=None): return user_data_path(cat, user_id)
def history_path(user_id=None): return user_data_path('history', user_id)
def styles_path(user_id=None): return user_data_path('styles', user_id)
def settings_path(user_id=None): return user_data_path('api_settings', user_id)
def users_path(): return os.path.join(DATA, 'users.json')
def invitations_path(): return os.path.join(DATA, 'invitations.json')
def model_pricing_path(): return os.path.join(DATA, 'model_pricing.json')

_HISTORY_CONTEXT = threading.local()

def current_history_project_id():
    contextual = getattr(_HISTORY_CONTEXT, 'project_id', None)
    if contextual:
        return contextual
    if not has_request_context():
        return None
    project_id = request.headers.get('X-Project-ID') or request.args.get('project_id')
    if not project_id and request.referrer:
        try:
            project_id = parse_qs(urlparse(request.referrer).query).get('project_id', [None])[0]
        except (TypeError, ValueError):
            project_id = None
    project_id = str(project_id or '').strip()
    return project_id if re.fullmatch(r'[A-Za-z0-9_-]{1,128}', project_id) else None

def insert_history(entry, limit=100, user_id=None, project_id=None):
    entry = dict(entry)
    project_id = project_id or current_history_project_id()
    if project_id:
        entry['project_id'] = project_id
    with HISTORY_LOCK:
        path = history_path(user_id)
        hist = load_json(path, [])
        hist.insert(0, entry)
        save_json(path, hist[:limit])

def save_video_history(video_url, script, original_script=None, refined_script=None,
                       model=None, ratio=None, duration=None, resolution=None,
                       ref_count=None, user_id=None):
    entry = {
        'time': datetime.now().strftime('%Y-%m-%d %H:%M'),
        'type': 'video',
        'video_url': video_url,
        'script': (script or '')[:80],
        'original_script': original_script or script or '',
        'refined_script': refined_script or script or ''
    }
    if model: entry['model'] = model
    if ratio: entry['ratio'] = ratio
    if duration: entry['duration'] = duration
    if resolution: entry['resolution'] = resolution
    if ref_count is not None: entry['ref_count'] = ref_count
    insert_history(entry, user_id=user_id)
    return entry

def ensure_default_styles(styles):
    if not isinstance(styles, list):
        styles = []
    merged = list(styles)
    changed = False
    existing_ids = {s.get('id') for s in merged if isinstance(s, dict)}
    existing_names = {s.get('name') for s in merged if isinstance(s, dict)}
    for style in DEFAULT_STYLES:
        if style.get('id') not in existing_ids and style.get('name') not in existing_names:
            merged.append(style)
            changed = True
            existing_ids.add(style.get('id'))
            existing_names.add(style.get('name'))
    default_by_id = {s.get('id'): s for s in DEFAULT_STYLES}
    for i, style in enumerate(merged):
        if not isinstance(style, dict):
            continue
        default = default_by_id.get(style.get('id'))
        if not default:
            continue
        refreshed = dict(style)
        for key in ('name', 'thumbnail_url', 'prompt', 'negative_prompt', 'use_for_image', 'use_for_video'):
            if refreshed.get(key) != default.get(key):
                refreshed[key] = default.get(key)
                changed = True
        merged[i] = refreshed
    return merged, changed

# Ensure data directory and files exist (for fresh Volume mounts)
LIVE_ACTION_STYLE_ID = 'style_1'
LIVE_ACTION_GLOBAL_MARKER = '【全局真人影视风格｜全程生效】'

DEFAULT_STYLES = [{'id': 'style_1',
  'name': '真人短剧',
  'thumbnail_url': 'https://movie1.tos-cn-beijing.volces.com/6765e87c2c4540c09f4ede7be5e2bb2b',
  'prompt': '真人影视质感，8K超高清，RAW胶片质感。人物面部写实，保留原生细腻真实皮肤肌理、毛孔、细纹、泪光和自然刘海，五官稳定且有充分面部细节。根据剧情使用过肩拍摄和轻微俯视视角，前景人物背部或肩部遮挡并虚化。户外柔焦夜景，微弱暖调环境微光，浅景深、背景虚化，低饱和清冷电影色调。轻微呼吸运镜，画面丝滑流畅，动态舒缓自然，表情克制，无夸张大幅度动作，营造细腻、破碎、伤感的情绪氛围。',
  'negative_prompt': '五官扭曲，夸张大哭大笑，浓重滤镜，过度磨皮，塑料假脸，畸形肢体，手指错误，服装穿模，高饱和强光，画面卡顿，水印文字，多余人物，红眼特效，浓妆',
  'use_for_image': True,
  'use_for_video': True,
  'created_at': '2026-06-14 12:00'},
 {'id': 'style_2',
  'name': '电影写实',
  'thumbnail_url': 'https://movie1.tos-cn-beijing.volces.com/d338bd0fb5654fb1879044fa0fa918c0',
  'prompt': '整体呈现电影级写实风格，画面像高质量实拍电影剧照。光影有层次，色彩克制统一，环境材质真实可信，人物与场景融合自然。镜头语言稳重，构图有明确视觉中心，前景、中景、背景层次清楚。人物表情自然，动作真实，避免过度摆拍和广告感。画面保留真实摄影的细节、景深、光线方向和空间质感，适合严肃剧情、情绪戏和高质感叙事镜头。',
  'negative_prompt': '低清晰度，模糊，画面脏乱，构图混乱，人物五官变形，脸部崩坏，年龄漂移，性别变化，多余人物，肢体畸形，手指错误，手部融合，服装穿模，身体比例异常，主体被裁切，背景物体扭曲，字幕，水印，文字，Logo，边框，拼贴，多格漫画',
  'use_for_image': True,
  'use_for_video': True,
  'created_at': '2026-06-14 12:00'},
 {'id': 'style_3',
  'name': '都市偶像剧',
  'thumbnail_url': 'https://movie1.tos-cn-beijing.volces.com/897b8b73295647678b55ca7293f43df7',
  'prompt': '整体呈现现代都市偶像剧风格，画面明亮、干净、精致，人物外貌清爽好看，服装时尚但不夸张。场景具有现代城市生活质感，如咖啡馆、公寓、办公室、街道、商场等，环境整洁有设计感。光线柔和通透，肤色自然，色彩偏温暖或清新，画面浪漫但不过度梦幻。人物互动自然，情绪细腻，适合暧昧、重逢、误会、告白、职场和日常情节。',
  'negative_prompt': '低清晰度，模糊，画面脏乱，构图混乱，人物五官变形，脸部崩坏，年龄漂移，性别变化，多余人物，肢体畸形，手指错误，手部融合，服装穿模，身体比例异常，主体被裁切，背景物体扭曲，字幕，水印，文字，Logo，边框，拼贴，多格漫画',
  'use_for_image': True,
  'use_for_video': True,
  'created_at': '2026-06-14 12:00'},
 {'id': 'style_4',
  'name': '悬疑冷色',
  'thumbnail_url': 'https://movie1.tos-cn-beijing.volces.com/c884c1315ae449fa809b5cc9878ea72c',
  'prompt': '整体呈现悬疑剧冷色调风格，画面低饱和、冷蓝灰、阴影层次明显，氛围紧张压抑。光线方向明确，常带有窗外冷光、走廊灯、监控感顶光或局部硬光，形成强烈明暗对比。场景保持真实可信，细节克制，不夸张恐怖怪诞。人物表情内敛、警觉、压抑或怀疑，镜头构图留有适当空白和不安感。适合秘密、追踪、对峙、调查、背叛和反转情节。',
  'negative_prompt': '低清晰度，模糊，画面脏乱，构图混乱，人物五官变形，脸部崩坏，年龄漂移，性别变化，多余人物，肢体畸形，手指错误，手部融合，服装穿模，身体比例异常，主体被裁切，背景物体扭曲，字幕，水印，文字，Logo，边框，拼贴，多格漫画',
  'use_for_image': True,
  'use_for_video': True,
  'created_at': '2026-06-14 12:00'},
 {'id': 'style_5',
  'name': '古风权谋',
  'thumbnail_url': 'https://movie1.tos-cn-beijing.volces.com/67b94fcde64046a58233c4523acad92d',
  'prompt': '整体呈现古风影视剧权谋风格，画面庄重、精致、克制，具有中式建筑、古代服饰、木质结构、屏风、烛火、庭院、宫殿或书房等视觉元素。色彩以暖金、墨色、深红、青灰、木色为主，光影层次丰富，氛围沉稳大气。人物服装完整考究，发型、配饰、妆容符合古装审美，表情含蓄有张力。构图讲究人物站位和权力关系，适合对峙、密谈、审问、谋划和情感压抑的戏。',
  'negative_prompt': '低清晰度，模糊，画面脏乱，构图混乱，人物五官变形，脸部崩坏，年龄漂移，性别变化，多余人物，肢体畸形，手指错误，手部融合，服装穿模，身体比例异常，主体被裁切，背景物体扭曲，字幕，水印，文字，Logo，边框，拼贴，多格漫画',
  'use_for_image': True,
  'use_for_video': True,
  'created_at': '2026-06-14 12:00'},
 {'id': 'style_6',
  'name': '国漫3D',
  'thumbnail_url': 'https://movie1.tos-cn-beijing.volces.com/595d66045bf74804b0376dd8b366dc6d.jpg',
  'prompt': '整体呈现高质量国漫3D风格，人物为精致三维角色，五官清晰，面部表情细腻，发丝、服装、材质和光影具有影视级CG质感。画面华丽但不塑料，皮肤和布料材质有层次，动作姿态自然。场景空间完整，灯光具有戏剧性和层次感，色彩鲜明但不杂乱。适合奇幻、都市幻想、古风玄幻、热血、情绪爆发和高视觉冲击的AI漫剧内容。',
  'negative_prompt': '低清晰度，模糊，画面脏乱，构图混乱，人物五官变形，脸部崩坏，年龄漂移，性别变化，多余人物，肢体畸形，手指错误，手部融合，服装穿模，身体比例异常，主体被裁切，背景物体扭曲，字幕，水印，文字，Logo，边框，拼贴，多格漫画',
  'use_for_image': True,
  'use_for_video': True,
  'created_at': '2026-06-14 12:00'},
 {'id': 'style_7',
  'name': '日系动画',
  'thumbnail_url': 'https://movie1.tos-cn-beijing.volces.com/a05e6f30728a4c86bf9ba6f667059ade',
  'prompt': '整体呈现日系动画风格，画面线条干净，角色轮廓清晰，色彩柔和明亮，人物表情细腻，动作姿态自然。背景绘制精致，有生活气息或电影动画质感，光线温柔，有明确时间氛围，如清晨、黄昏、雨天、夜晚灯光等。人物服装和发型保持一致，不写实真人化，不变成3D塑料感。适合青春、校园、治愈、恋爱、日常、奇幻和情绪细腻的分镜。',
  'negative_prompt': '低清晰度，模糊，画面脏乱，构图混乱，人物五官变形，脸部崩坏，年龄漂移，性别变化，多余人物，肢体畸形，手指错误，手部融合，服装穿模，身体比例异常，主体被裁切，背景物体扭曲，字幕，水印，文字，Logo，边框，拼贴，多格漫画',
  'use_for_image': True,
  'use_for_video': True,
  'created_at': '2026-06-14 12:00'},
 {'id': 'style_8',
  'name': '赛博霓虹',
  'thumbnail_url': 'https://movie1.tos-cn-beijing.volces.com/1af24b6d7c8b462ca8b6bb158ed7d452',
  'prompt': '整体呈现赛博朋克霓虹风格，画面以夜景城市、霓虹灯牌、湿润街道、玻璃反射、金属材质、电子屏幕和未来科技元素为主。色彩以蓝、紫、粉、青色霓虹光为核心，暗部深邃，高光强烈，反射丰富。人物造型具有未来都市感，但不杂乱堆砌。构图强调城市纵深、光影对比和科技氛围，适合追逐、交易、黑客、夜晚对峙、未来都市短剧和强视觉风格镜头。',
  'negative_prompt': '低清晰度，模糊，画面脏乱，构图混乱，人物五官变形，脸部崩坏，年龄漂移，性别变化，多余人物，肢体畸形，手指错误，手部融合，服装穿模，身体比例异常，主体被裁切，背景物体扭曲，字幕，水印，文字，Logo，边框，拼贴，多格漫画',
  'use_for_image': True,
  'use_for_video': True,
  'created_at': '2026-06-14 12:00'},
 {'id': 'style_9',
  'name': '韩漫短剧',
  'thumbnail_url': 'https://movie1.tos-cn-beijing.volces.com/16e04c823639433e8d83f7cec108c588',
  'prompt': '【强制画风：韩漫 / 韩国 Webtoon 2D 插画】画面必须是精致韩系网络漫画、恋爱条漫、动态漫关键帧风格，不是真人照片，不是3D渲染，不是AI真人写真，不是服装商品图。人物使用干净细腻的黑色或深棕线稿，柔和赛璐璐上色，半透明皮肤阴影，精致眼睫和眼神高光，薄唇，清爽脸型，细腻发丝分束，整体有韩漫男主/女主的修长肩颈、小脸、挺拔身形和高级感。色彩偏粉白、浅蓝、奶油光或清透暖光，背景可以是校园、樱花、都市街道、咖啡馆、公寓、办公室、医院、豪宅等韩漫场景，浅景深、柔光、花瓣或细腻空气感可以使用。构图像高质量韩国恋爱漫画封面或单格大画面，人物情绪细腻，适合暧昧、重逢、误会、追妻、契约恋爱、财阀、复仇、职场和都市情感剧情。即使生成角色三视图、服装设定或分镜构图，也必须保持同一个韩漫2D插画画风，不能变成写实模特照、3D人偶、棚拍白底商品图。可参考韩国条漫的细线稿、柔和阴影、清透肤色、精致五官和浪漫氛围，但不要生成文字、对白框、字幕或水印。',
  'negative_prompt': '真人照片，真实摄影，AI真人写真，写实皮肤纹理，3D渲染，CG人偶，塑料感，服装商品图，模特棚拍，白底真人设定照，证件照，欧美漫画，厚涂油画，低幼Q版，粗糙线条，夸张表情包风格，韩文文字，对白框，字幕，水印，Logo，低清晰度，模糊，画面脏乱，构图混乱，人物五官变形，脸部崩坏，年龄漂移，性别变化，多余人物，肢体畸形，手指错误，手部融合，服装穿模，身体比例异常，主体被裁切，背景物体扭曲，边框，拼贴，多格漫画',
  'use_for_image': True,
  'use_for_video': True,
  'created_at': '2026-06-14 12:00'}]

def init_data():
    os.makedirs(DATA, exist_ok=True)
    for path, default in [
        (users_path(), {}),
        (invitations_path(), []),
    ]:
        if not os.path.exists(path) or os.path.getsize(path) < 10:
            save_json(path, default)

    initial_password = os.environ.get('ADMIN_INITIAL_PASSWORD', '')
    admin_ids = [item.strip().lower() for item in os.environ.get('ADMIN_USERS', '').split(',') if item.strip()]
    if initial_password and admin_ids:
        users = load_json(users_path(), {})
        changed = False
        for admin_id in admin_ids:
            if admin_id not in users:
                users[admin_id] = {
                    'username': admin_id,
                    'password_hash': generate_password_hash(initial_password, method='pbkdf2:sha256'),
                    'created_at': datetime.now(timezone.utc).isoformat(),
                    'model_permissions': ['text', 'image', 'video'],
                    'points': 0
                }
                changed = True
        if changed:
            save_json(users_path(), users)

init_postgres()
migrate_local_json_to_postgres()
init_data()

# ── Object storage upload helper ──────────────────────────────
def get_tos_client():
    return tos.TosClientV2(TOS_AK, TOS_SK, TOS_ENDPOINT, TOS_REGION)

def r2_configured():
    return all([R2_ACCOUNT_ID, R2_ACCESS_KEY_ID, R2_SECRET_ACCESS_KEY, R2_BUCKET, R2_PUBLIC_BASE])

def persistent_storage_configured():
    return r2_configured() or bool(TOS_AK and TOS_SK)

def active_storage_name():
    if r2_configured():
        return 'r2'
    if TOS_AK and TOS_SK:
        return 'tos'
    return 'local'

def storage_name_for_url(url):
    if R2_PUBLIC_BASE and url and url.startswith(R2_PUBLIC_BASE):
        return 'r2'
    if TOS_PUBLIC_BASE and url and url.startswith(TOS_PUBLIC_BASE):
        return 'tos'
    return 'local'

def get_r2_client():
    import boto3
    return boto3.client(
        's3',
        endpoint_url=f'https://{R2_ACCOUNT_ID}.r2.cloudflarestorage.com',
        aws_access_key_id=R2_ACCESS_KEY_ID,
        aws_secret_access_key=R2_SECRET_ACCESS_KEY,
        region_name='auto',
    )

def r2_public_url(object_key):
    return f"{R2_PUBLIC_BASE}/{quote(object_key, safe='/')}"

def upload_to_r2(file_data, object_key, content_type='application/octet-stream', content_length=None):
    """Upload bytes or stream to R2, return (public_url, success)."""
    if not r2_configured():
        return None, False
    if content_length is not None and int(content_length) <= 0:
        print(f'[R2] refused empty object: {object_key}', flush=True)
        return None, False
    try:
        kwargs = {
            'Bucket': R2_BUCKET,
            'Key': object_key,
            'Body': file_data,
            'ContentType': content_type,
        }
        if content_length:
            kwargs['ContentLength'] = content_length
        client = get_r2_client()
        client.put_object(**kwargs)
        stored_size = int(client.head_object(Bucket=R2_BUCKET, Key=object_key).get('ContentLength') or 0)
        if stored_size <= 0 or (content_length is not None and stored_size != int(content_length)):
            client.delete_object(Bucket=R2_BUCKET, Key=object_key)
            raise Exception(f'对象大小校验失败：expected={content_length} stored={stored_size}')
        return r2_public_url(object_key), True
    except Exception as e:
        print(f'[R2] upload failed: {e}', flush=True)
        return None, False

def is_persistent_storage_url(url):
    if not url:
        return False
    if r2_configured():
        return bool(R2_PUBLIC_BASE and url.startswith(R2_PUBLIC_BASE))
    return bool(TOS_PUBLIC_BASE and url.startswith(TOS_PUBLIC_BASE))

def upload_to_tos(file_data, object_key, content_type='application/octet-stream', content_length=None):
    """Upload bytes or stream to durable object storage, return (public_url, success)."""
    if r2_configured():
        return upload_to_r2(file_data, object_key, content_type, content_length)

    try:
        client = get_tos_client()
        if hasattr(file_data, 'read'):
            # Stream: use put_object with streaming body
            client.put_object(TOS_BUCKET, object_key,
                              content=file_data,
                              content_type=content_type,
                              content_length=content_length)
        else:
            client.put_object(TOS_BUCKET, object_key,
                              content=file_data,
                              content_type=content_type)
        return f"{TOS_PUBLIC_BASE}/{object_key}", True
    except Exception as e:
        print(f'[TOS] upload failed: {e}')
        return None, False


# ── User authentication ──────────────────────────────────────
USERS_LOCK = threading.Lock()

class QuotaError(Exception): pass

def minimax_api_root(api_base):
    base_url = (api_base or '').rstrip('/')
    for suffix in ('/v2/video_generation', '/v2'):
        if base_url.lower().endswith(suffix):
            return base_url[:-len(suffix)]
    return base_url

def normalize_full_api_url(value):
    url = str(value or '').strip()
    if not url:
        return ''
    parsed = urlparse(url)
    if parsed.scheme not in ('http', 'https') or not parsed.netloc:
        raise ValueError('完整接口必须是以 http:// 或 https:// 开头的 URL')
    return url

def normalize_api_profile(profile):
    profile = dict(profile or {})
    base_url = (profile.get('base_url') or '').rstrip('/')
    provider = profile.get('provider', '')
    submit_url = normalize_full_api_url(profile.get('submit_url'))
    status_url = normalize_full_api_url(profile.get('status_url')).replace('{taskId}', '{task_id}')
    endpoint_mode = profile.get('endpoint_mode') if profile.get('endpoint_mode') == 'full' else ''
    if not base_url and submit_url:
        parsed_submit = urlparse(submit_url)
        base_url = f'{parsed_submit.scheme}://{parsed_submit.netloc}'
    if 'api.atlascloud.ai' in base_url.lower():
        provider = 'atlas'
        base_url = ATLAS_API_BASE
    if str(profile.get('model') or '').lower() == 'minimax-h3' and any(host in base_url.lower() for host in ('api.minimaxi.com', 'api.minimax.io')):
        provider = 'minimax'
        base_url = minimax_api_root(base_url)
    normalized = {
        'id': profile.get('id') or uuid.uuid4().hex,
        'name': (profile.get('name') or '').strip(),
        'provider': provider, 'base_url': base_url,
        'api_key': profile.get('api_key', ''), 'model': profile.get('model', ''),
        'submit_url': submit_url, 'status_url': status_url,
        'endpoint_mode': endpoint_mode,
        'last_test': profile.get('last_test'),
        'created_at': profile.get('created_at') or datetime.now(timezone.utc).isoformat()
    }
    if not normalized['name']:
        normalized['name'] = ' · '.join(filter(None, [provider or 'Custom API', normalized['model']]))
    return normalized

def get_api_profiles(kind, user_id=None):
    path = settings_path(user_id)
    settings = load_json(path, {})
    stored_profiles = settings.get('api_profiles', {}).get(kind)
    changed = not isinstance(stored_profiles, list)
    if isinstance(stored_profiles, list):
        profiles = [normalize_api_profile(item) for item in stored_profiles if isinstance(item, dict)]
        changed = changed or profiles != stored_profiles
    else:
        legacy = settings.get('apis', {}).get(kind, {})
        profiles = [normalize_api_profile(legacy)] if isinstance(legacy, dict) and any(legacy.get(key) for key in ('provider', 'base_url', 'api_key', 'model')) else []

    selected = settings.get('selected_api_profiles', {}).get(kind)
    ids = {item['id'] for item in profiles}
    if selected not in ids:
        selected = profiles[0]['id'] if profiles else None
        changed = True
    if changed:
        settings.setdefault('api_profiles', {})[kind] = profiles
        settings.setdefault('selected_api_profiles', {})[kind] = selected
        if profiles:
            settings.setdefault('apis', {})[kind] = dict(next(item for item in profiles if item['id'] == selected))
        save_json(path, settings)
    return profiles, selected

def public_api_profile(profile, kind=None):
    public = {
        'id': profile.get('id'), 'name': profile.get('name', ''),
        'provider': profile.get('provider', ''), 'base_url': profile.get('base_url', ''),
        'model': profile.get('model', ''), 'configured': bool(profile.get('api_key')),
        'submit_url': profile.get('submit_url', ''), 'status_url': profile.get('status_url', ''),
        'endpoint_mode': profile.get('endpoint_mode', ''),
        'last_test': profile.get('last_test')
    }
    model_name = str(profile.get('model') or '').lower()
    capabilities = next((caps for name, caps in MODEL_CAPS.items() if name.lower() == model_name), None)
    if not capabilities and kind == 'video' and profile.get('provider') == 'nano':
        capabilities = nano_video_model_capabilities(profile)
    elif not capabilities and kind == 'video' and profile.get('provider') == 'atlas':
        capabilities = atlas_video_model_capabilities(profile)
    elif not capabilities and kind == 'video' and profile.get('provider') == 'ark':
        capabilities = MODEL_CAPS.get('seedance')
    if capabilities:
        public['capabilities'] = capabilities
    return public

_NANO_VIDEO_MODEL_CACHE = {}
_ATLAS_MODEL_CACHE = {'time': 0, 'items': []}
_ATLAS_SCHEMA_CACHE = {}

def atlas_model_items():
    if _ATLAS_MODEL_CACHE['items'] and time.time() - _ATLAS_MODEL_CACHE['time'] < 3600:
        return _ATLAS_MODEL_CACHE['items']
    try:
        response = requests.get(f'{ATLAS_API_BASE}/models', timeout=20)
        if response.status_code != 200:
            return []
        body = response.json()
        items = body.get('data', body) if isinstance(body, dict) else body
        if not isinstance(items, list):
            return []
        _ATLAS_MODEL_CACHE.update(time=time.time(), items=items)
        return items
    except (requests.RequestException, ValueError):
        return []

def atlas_model_schema(model_id):
    if model_id in _ATLAS_SCHEMA_CACHE:
        return _ATLAS_SCHEMA_CACHE[model_id]
    item = next((row for row in atlas_model_items() if row.get('model') == model_id), None)
    schema_url = (item or {}).get('schema')
    if not schema_url:
        return None
    try:
        response = requests.get(schema_url, timeout=20)
        if response.status_code != 200:
            return None
        schema = response.json()
        _ATLAS_SCHEMA_CACHE[model_id] = schema
        return schema
    except (requests.RequestException, ValueError):
        return None

def atlas_input_properties(model_id):
    schema = atlas_model_schema(model_id) or {}
    components = schema.get('components') or {}
    schemas = components.get('schemas') or {}
    input_schema = schemas.get('Input') or schema.get('input_schema') or {}
    return input_schema.get('properties') or {}

def atlas_image_variant(model_id):
    if not model_id:
        return None
    if model_id.endswith('/image-to-video'):
        return model_id
    if model_id.endswith(('/text-to-video', '/reference-to-video')):
        candidate = model_id.rsplit('/', 1)[0] + '/image-to-video'
        if any(row.get('model') == candidate for row in atlas_model_items()):
            return candidate
    return None

def _schema_enum(properties, *names):
    for name in names:
        definition = properties.get(name) or {}
        values = definition.get('enum')
        if isinstance(values, list):
            return values
    return []

def atlas_video_model_capabilities(profile):
    model_id = str((profile or {}).get('model') or '')
    effective_model = atlas_image_variant(model_id) or model_id
    properties = atlas_input_properties(effective_model)
    if not properties:
        return None
    names = set(properties)
    duration_def = properties.get('duration') or {}
    durations = _schema_enum(properties, 'duration')
    numeric_durations = []
    for value in durations:
        try:
            parsed = int(value)
            if parsed > 0:
                numeric_durations.append(parsed)
        except (TypeError, ValueError):
            pass
    caps = {
        'supports_first_frame': bool(names & {'image', 'image_url', 'imageUrl', 'first_frame_url', 'firstFrameUrl'}),
        'supports_last_frame': bool(names & {'last_image', 'lastImage', 'last_frame_url', 'lastFrameUrl'}),
        'supports_reference_images': bool(names & {'reference_images', 'images'}),
        'supports_reference_audio': bool(names & {'reference_audios', 'audio_url'}),
        'supports_reference_video': bool(names & {'reference_videos', 'video_url'}),
    }
    resolutions = [str(value) for value in _schema_enum(properties, 'resolution')]
    ratios = [str(value) for value in _schema_enum(properties, 'ratio', 'aspect_ratio')]
    if resolutions:
        caps['resolutions'] = resolutions
    if ratios:
        caps['ratios'] = ratios
    if numeric_durations:
        caps['min_duration'] = min(numeric_durations)
        caps['max_duration'] = max(numeric_durations)
    else:
        if duration_def.get('minimum') is not None:
            caps['min_duration'] = duration_def['minimum']
        if duration_def.get('maximum') is not None:
            caps['max_duration'] = duration_def['maximum']
    return caps

def nano_video_model_info(profile):
    if not profile or profile.get('provider') != 'nano' or not profile.get('api_key') or not profile.get('model'):
        return None
    api_root = (profile.get('base_url') or NANO_GPT_BASE).rstrip('/')
    if api_root.endswith('/v1'):
        api_root = api_root[:-3]
    cache_key = (api_root, profile['api_key'])
    cached = _NANO_VIDEO_MODEL_CACHE.get(cache_key)
    if cached and time.time() - cached['time'] < 300:
        items = cached['items']
        return next((item for item in items if item.get('id') == profile.get('model')), None)
    try:
        response = requests.get(
            f'{api_root}/v1/video-models',
            headers={'x-api-key': profile['api_key']}, timeout=20
        )
        if response.status_code != 200:
            return None
        items = response.json().get('data', [])
        _NANO_VIDEO_MODEL_CACHE[cache_key] = {'time': time.time(), 'items': items}
        return next((item for item in items if item.get('id') == profile.get('model')), None)
    except (requests.RequestException, ValueError):
        return None

def nano_video_model_capabilities(profile, model_info=None):
    info = model_info or nano_video_model_info(profile)
    if not info:
        return None
    declared = info.get('capabilities') or {}
    parameter_defs = ((info.get('supported_parameters') or {}).get('parameters') or {})
    parameter_names = set(parameter_defs)
    pricing = info.get('pricing') or {}
    resolutions = []
    resolution_def = parameter_defs.get('resolution') or {}
    for option in resolution_def.get('options') or []:
        value = option.get('value') if isinstance(option, dict) else option
        if value is not None:
            resolutions.append(str(value))
    duration_def = parameter_defs.get('duration') or {}
    durations = []
    for option in duration_def.get('options') or []:
        value = option.get('value') if isinstance(option, dict) else option
        try:
            durations.append(int(value))
        except (TypeError, ValueError):
            pass
    ratio_def = parameter_defs.get('aspect_ratio') or {}
    ratios = []
    for option in ratio_def.get('options') or []:
        value = option.get('value') if isinstance(option, dict) else option
        if value is not None:
            ratios.append(str(value))
    caps = {
        'supports_first_frame': bool(declared.get('image_to_video')),
        'supports_last_frame': bool({'last_image', 'lastImage', 'lastFrameUrl', 'last_frame_url'} & parameter_names),
        'supports_reference_images': bool({'referenceImages', 'reference_images'} & parameter_names),
        'supports_reference_audio': bool(declared.get('audio_input')),
        'supports_reference_video': bool(declared.get('video_to_video')),
    }
    if resolutions:
        caps['resolutions'] = resolutions
    if ratios:
        caps['ratios'] = ratios
    if durations:
        caps['min_duration'] = min(durations)
        caps['max_duration'] = max(durations)
    else:
        if pricing.get('min_duration') is not None:
            caps['min_duration'] = pricing['min_duration']
        if pricing.get('max_duration') is not None:
            caps['max_duration'] = pricing['max_duration']
    return caps

def get_personal_api(kind, user_id=None, profile_id=None):
    profiles, selected = get_api_profiles(kind, user_id)
    target_id = profile_id or selected
    profile = next((item for item in profiles if item['id'] == target_id), None)
    return dict(profile or {'provider': '', 'base_url': '', 'api_key': '', 'model': '', 'last_test': None})

def use_personal_api(user_id=None):
    user_id = user_id or current_user_id()
    if is_admin(user_id): return False
    user = load_json(users_path(), {}).get(user_id, {})
    return int(user.get('points') or 0) <= 0

def resolve_api(kind, builtin, user_id=None, force_personal=False, profile_id=None, strict_builtin=False):
    if not force_personal: return builtin
    personal = get_personal_api(kind, user_id, profile_id)
    if personal.get('api_key') and (personal.get('base_url') or personal.get('submit_url')) and personal.get('model'): return personal
    if force_personal:
        raise QuotaError('请选择一个已保存且配置完整的个人 API')
    return builtin

def get_user_api_settings(user_id=None):
    return {
        'ark_api_key': ARK_API_KEY, 'nano_api_key': NANO_GPT_API_KEY,
        'third_party_api_base': THIRD_PARTY_API_BASE, 'third_party_api_key': THIRD_PARTY_API_KEY,
    }

def is_admin(user_id=None):
    user_id = (user_id or current_user_id() or '').lower()
    admins = {item.strip().lower() for item in os.environ.get('ADMIN_USERS', '').split(',') if item.strip()}
    return user_id in admins

def model_pricing_catalog():
    labels = {
        'doubao': '豆包', 'glm46': 'GPT-4.1 Mini', 'claude46': 'Claude 4.6',
        'gpt-image-2': 'GPT Image 2', 'nano-banana-2': 'Banana 2',
        NANO_GPT_MIDJOURNEY_MODEL_ID: 'Midjourney 8.2',
        AGNES_IMAGE_MODEL_ID: 'Agnes Image 2.1 Flash', 'volc-seedream-4-5': 'Seedream 4.5',
        VOLC_SEEDREAM_5_PRO_MODEL_ID: 'Seedream 5.0 Pro',
        'seedance': 'Seedance（火山）', AGNES_VIDEO_MODEL_ID: 'Agnes Video v2.0',
        'kling-v30-std': 'Kling v3.0 Std', 'grok-imagine-video': 'Grok Imagine',
        'vidu-q3': 'Vidu Q3', 'seedance-v15-pro': 'Seedance v1.5 Pro',
        THIRD_PARTY_MODEL_ID: '第三方内置模型',
    }
    groups = {
        'text': list(SCRIPT_MODELS),
        'image': list(ALL_IMAGE_MODELS),
        'video': list(dict.fromkeys(ALL_MODELS)),
    }
    return [
        {'kind': kind, 'model': model, 'label': labels.get(model, model),
         'unit': 'second' if kind == 'video' else 'request'}
        for kind, models in groups.items() for model in models
    ]

def load_model_pricing():
    stored = load_json(model_pricing_path(), {})
    return stored if isinstance(stored, dict) else {}

def model_point_cost(kind, model, quantity=1):
    pricing = load_model_pricing()
    kind_pricing = pricing.get(kind) or {}
    legacy_model = next((old for old, current in LEGACY_IMAGE_MODEL_ALIASES.items() if current == model), None)
    unit_price = max(0, int(kind_pricing.get(model) or (kind_pricing.get(legacy_model) if legacy_model else 0) or 0))
    multiplier = max(1, int(quantity or 1)) if kind == 'video' else 1
    return unit_price * multiplier

def reserve_model_points(kind, model, user_id=None, quantity=1, personal=False):
    user_id = (user_id or current_user_id() or '').lower()
    if personal or is_admin(user_id): return 0
    cost = model_point_cost(kind, model, quantity)
    if cost <= 0: return 0
    with USERS_LOCK:
        users = load_json(users_path(), {})
        user = users.get(user_id)
        if not user: raise QuotaError('用户不存在')
        balance = max(0, int(user.get('points') or 0))
        if balance < cost:
            unit = f'（{quantity}秒 × {cost // max(1, int(quantity or 1))}分）' if kind == 'video' else ''
            raise QuotaError(f'积分不足：本次需要 {cost} 分{unit}，当前剩余 {balance} 分。可选择「自己的 API」')
        user['points'] = balance - cost
        save_json(users_path(), users)
    return cost

def refund_model_points(user_id, cost):
    if not cost or is_admin(user_id): return
    with USERS_LOCK:
        users = load_json(users_path(), {})
        if user_id not in users: return
        users[user_id]['points'] = max(0, int(users[user_id].get('points') or 0)) + int(cost)
        save_json(users_path(), users)

def start_metered_job(target, args, job_id, user_id, point_cost):
    project_id = current_history_project_id()
    def runner():
        _HISTORY_CONTEXT.project_id = project_id
        try:
            target(*args)
        except Exception as exc:
            JOBS[job_id] = {'status': 'failed', 'error': str(exc)}
        finally:
            if point_cost and JOBS.get(job_id, {}).get('status') == 'failed':
                refund_model_points(user_id, point_cost)
            _HISTORY_CONTEXT.project_id = None
    threading.Thread(target=runner, daemon=True).start()

def _mock_generation_job(job_id, kind='image', delay=0.8):
    """Test-only hook (MANGA_MOCK_GENERATION=1): completes jobs without any
    provider call so acceptance runs never touch real paid models."""
    def runner():
        time.sleep(delay)
        if kind == 'image':
            JOBS[job_id] = {'status': 'succeeded', 'url': '/static/site-icon.png', 'name': 'mock.png',
                            'model': 'mock', 'ratio': '1:1', 'mode': 'storyboard'}
        else:
            JOBS[job_id] = {'status': 'succeeded', 'video_url': '/static/site-icon.png', 'error': None,
                            'model': 'mock', 'ratio': '9:16', 'duration': 5, 'resolution': '720p'}
    threading.Thread(target=runner, daemon=True).start()

def admin_required(f):
    @functools.wraps(f)
    def wrap(*a, **kw):
        if not current_user_id(): return jsonify(error='unauthorized'), 401
        if not is_admin(): return jsonify(error='forbidden'), 403
        return f(*a, **kw)
    return wrap

def model_access_required(kind):
    def decorator(f):
        @functools.wraps(f)
        def wrap(*a, **kw):
            if is_admin(): return f(*a, **kw)
            user = load_json(users_path(), {}).get(current_user_id(), {})
            if kind not in user.get('model_permissions', []):
                return jsonify(error=f'当前账号未开放{kind}模型权限，请联系管理员获取邀请码'), 403
            return f(*a, **kw)
        return wrap
    return decorator

def login_required(f):
    @functools.wraps(f)
    def wrap(*a, **kw):
        user_id = current_user_id()
        if not user_id: return jsonify(error='unauthorized'), 401
        if load_json(users_path(), {}).get(user_id, {}).get('disabled'):
            session.clear()
            return jsonify(error='账号已停用'), 403
        return f(*a, **kw)
    return wrap
# ── static pages ──────────────────────────────────────────────
@app.route('/')
def index():
    if not current_user_id():
        return send_from_directory('static', 'login.html')
    return send_from_directory('static', 'projects.html')

@app.route('/classic')
def classic_page():
    if not current_user_id():
        return redirect('/')
    return send_from_directory('static', 'index.html')

@app.route('/assets')
def assets_page():
    if not current_user_id():
        return redirect('/')
    return send_from_directory('static', 'assets.html')

@app.route('/workspace/<project_id>')
def workspace_page(project_id):
    if not current_user_id():
        return redirect('/')
    project = _projects_load(current_user_id()).get('projects', {}).get(project_id)
    if not isinstance(project, dict):
        return jsonify(error='项目不存在'), 404
    return send_from_directory('static', 'workspace.html')

@app.route('/admin')
def admin_page():
    if not current_user_id(): return redirect('/')
    if not is_admin(): return '<h2 style="text-align:center;margin-top:100px">无管理员权限</h2>', 403
    return send_from_directory('static', 'admin.html')

@app.route('/api-settings')
@login_required
def api_settings_page():
    return send_from_directory('static', 'api-settings.html')

# ── Canvas workspace (additive) ──────────────────────────────
CANVAS_LOCK = threading.Lock()
MAX_CANVASES_PER_USER = 100
MAX_CANVAS_NODES = 300
MAX_CANVAS_EDGES = 600
MAX_CANVAS_TITLE_LEN = 100
MAX_CANVAS_ID_LEN = 128
MAX_CANVAS_STR_LEN = 200000
MAX_CANVAS_BODY_BYTES = 512 * 1024
CANVAS_NODE_TYPES = ('script', 'shot', 'asset', 'image', 'video', 'result', 'note')
SENSITIVE_KEYS = ('api_key', 'cookie', 'authorization', 'password', 'token', 'secret')

class CanvasValidationError(Exception):
    pass

def canvases_path(user_id=None):
    return user_data_path('canvases', user_id)

def _canvas_is_finite_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)

def _canvas_valid_id(value):
    return isinstance(value, str) and 1 <= len(value) <= MAX_CANVAS_ID_LEN

def _canvas_valid_str(value, max_len=MAX_CANVAS_STR_LEN):
    return isinstance(value, str) and len(value) <= max_len

def _canvas_valid_opt_str(value, max_len=MAX_CANVAS_STR_LEN):
    return value is None or _canvas_valid_str(value, max_len)

def _canvas_valid_media_url(value):
    if value is None:
        return True
    if not isinstance(value, str) or len(value) > 2000:
        return False
    lowered = value.lower()
    if lowered.startswith('data:') or ';base64' in lowered:
        return False
    return value.startswith('http://') or value.startswith('https://') or value.startswith('/static/')

def _canvas_reject_sensitive_keys(value):
    if isinstance(value, dict):
        for key, item in value.items():
            if str(key).lower() in SENSITIVE_KEYS:
                raise CanvasValidationError('包含敏感字段: ' + str(key))
            _canvas_reject_sensitive_keys(item)
    elif isinstance(value, list):
        for item in value:
            _canvas_reject_sensitive_keys(item)

def _canvas_valid_position(value):
    return (isinstance(value, dict) and set(value.keys()) == {'x', 'y'}
            and _canvas_is_finite_number(value.get('x'))
            and _canvas_is_finite_number(value.get('y')))

def _canvas_valid_viewport(value):
    return (isinstance(value, dict) and set(value.keys()) == {'x', 'y', 'zoom'}
            and _canvas_is_finite_number(value.get('x'))
            and _canvas_is_finite_number(value.get('y'))
            and _canvas_is_finite_number(value.get('zoom'))
            and value.get('zoom') > 0)

def _canvas_valid_enum(value, options):
    return isinstance(value, str) and value in options

def _canvas_valid_bool(value):
    return isinstance(value, bool)

def _canvas_valid_int_range(value, low, high):
    return isinstance(value, int) and not isinstance(value, bool) and low <= value <= high

def _canvas_valid_string_list(value, max_items):
    return (isinstance(value, list) and len(value) <= max_items
            and all(_canvas_valid_id(item) for item in value))

def _canvas_valid_asset_ref(ref):
    if not isinstance(ref, dict):
        return False
    allowed = {'source', 'ref_id', 'name', 'url', 'role_label'}
    if set(ref.keys()) - allowed:
        return False
    if not _canvas_valid_enum(ref.get('source'), ('character', 'outfit', 'scene', 'audio', 'upload', 'style', 'video')):
        return False
    ref_id = ref.get('ref_id')
    if ref_id is not None and not (isinstance(ref_id, (str, int)) and not isinstance(ref_id, bool) and len(str(ref_id)) <= MAX_CANVAS_ID_LEN):
        return False
    if not _canvas_valid_opt_str(ref.get('name'), 200):
        return False
    if not _canvas_valid_media_url(ref.get('url')):
        return False
    if not _canvas_valid_opt_str(ref.get('role_label'), 200):
        return False
    return True

def _canvas_valid_split(split):
    if not isinstance(split, dict):
        return False
    allowed = {'job_id', 'status', 'shots', 'error'}
    if set(split.keys()) - allowed:
        return False
    if not _canvas_valid_opt_str(split.get('job_id'), MAX_CANVAS_ID_LEN):
        return False
    if not _canvas_valid_enum(split.get('status'), ('idle', 'running', 'succeeded', 'failed')):
        return False
    shots = split.get('shots')
    if not isinstance(shots, list) or len(shots) > 200:
        return False
    if not all(isinstance(item, dict) for item in shots):
        return False
    if not _canvas_valid_opt_str(split.get('error'), 2000):
        return False
    return True

def _canvas_valid_segment(segment):
    return segment is None or isinstance(segment, dict)

CANVAS_NODE_DATA_SPECS = {
    'script': {
        'label': lambda v: _canvas_valid_opt_str(v, 100),
        'script': lambda v: _canvas_valid_opt_str(v, MAX_CANVAS_STR_LEN),
        'script_model': lambda v: _canvas_valid_enum(v, ('doubao', 'glm46', 'claude46', 'personal-api')),
        'split_mode': lambda v: _canvas_valid_enum(v, ('smart', 'short', 'long')),
        'style_id': lambda v: _canvas_valid_opt_str(v, MAX_CANVAS_ID_LEN),
        'use_personal_api': _canvas_valid_bool,
        'api_profile_id': lambda v: _canvas_valid_opt_str(v, MAX_CANVAS_ID_LEN),
        'split': _canvas_valid_split,
    },
    'shot': {
        'label': lambda v: _canvas_valid_opt_str(v, 100),
        'script_node_id': lambda v: _canvas_valid_opt_str(v, MAX_CANVAS_ID_LEN),
        'shot_index': lambda v: _canvas_valid_int_range(v, 0, 10000),
        'segment': _canvas_valid_segment,
    },
    'asset': {
        'label': lambda v: _canvas_valid_opt_str(v, 100),
        'asset_type': lambda v: _canvas_valid_enum(v, ('character', 'outfit', 'scene', 'audio', 'upload', 'style', 'video')),
        'refs': lambda v: isinstance(v, list) and len(v) <= 200 and all(_canvas_valid_asset_ref(item) for item in v),
    },
    'image': {
        'label': lambda v: _canvas_valid_opt_str(v, 100),
        'role': lambda v: _canvas_valid_enum(v, ('storyboard', 'first_frame', 'last_frame')),
        'prompt': lambda v: _canvas_valid_opt_str(v, MAX_CANVAS_STR_LEN),
        'image_model': lambda v: _canvas_valid_opt_str(v, 128),
        'ratio': lambda v: _canvas_valid_opt_str(v, 32),
        'style_id': lambda v: _canvas_valid_opt_str(v, MAX_CANVAS_ID_LEN),
        'use_personal_api': _canvas_valid_bool,
        'api_profile_id': lambda v: _canvas_valid_opt_str(v, MAX_CANVAS_ID_LEN),
        'input_asset_node_ids': lambda v: _canvas_valid_string_list(v, 200),
        'job_id': lambda v: _canvas_valid_opt_str(v, MAX_CANVAS_ID_LEN),
        'status': lambda v: _canvas_valid_enum(v, ('idle', 'running', 'succeeded', 'failed')),
        'image_url': _canvas_valid_media_url,
        'error': lambda v: _canvas_valid_opt_str(v, 2000),
    },
    'video': {
        'label': lambda v: _canvas_valid_opt_str(v, 100),
        'video_model': lambda v: _canvas_valid_opt_str(v, 128),
        'ratio': lambda v: _canvas_valid_opt_str(v, 32),
        'duration': lambda v: _canvas_valid_int_range(v, 1, 120),
        'resolution': lambda v: _canvas_valid_opt_str(v, 32),
        'optimize_prompt': _canvas_valid_bool,
        'style_id': lambda v: _canvas_valid_opt_str(v, MAX_CANVAS_ID_LEN),
        'use_personal_api': _canvas_valid_bool,
        'api_profile_id': lambda v: _canvas_valid_opt_str(v, MAX_CANVAS_ID_LEN),
        'shot_node_id': lambda v: _canvas_valid_opt_str(v, MAX_CANVAS_ID_LEN),
        'asset_node_ids': lambda v: _canvas_valid_string_list(v, 200),
        'storyboard_image_node_id': lambda v: _canvas_valid_opt_str(v, MAX_CANVAS_ID_LEN),
        'first_frame_node_id': lambda v: _canvas_valid_opt_str(v, MAX_CANVAS_ID_LEN),
        'last_frame_node_id': lambda v: _canvas_valid_opt_str(v, MAX_CANVAS_ID_LEN),
        'job_id': lambda v: _canvas_valid_opt_str(v, MAX_CANVAS_ID_LEN),
        'status': lambda v: _canvas_valid_enum(v, ('idle', 'running', 'succeeded', 'failed')),
        'error': lambda v: _canvas_valid_opt_str(v, 2000),
    },
    'result': {
        'label': lambda v: _canvas_valid_opt_str(v, 100),
        'kind': lambda v: _canvas_valid_enum(v, ('image', 'video')),
        'media_url': _canvas_valid_media_url,
        'history_id': lambda v: _canvas_valid_opt_str(v, MAX_CANVAS_ID_LEN),
        'source_job_id': lambda v: _canvas_valid_opt_str(v, MAX_CANVAS_ID_LEN),
    },
    'note': {
        'label': lambda v: _canvas_valid_opt_str(v, 100),
        'text': lambda v: _canvas_valid_opt_str(v, MAX_CANVAS_STR_LEN),
        'color': lambda v: _canvas_valid_opt_str(v, 32),
    },
}

def _canvas_check_fields(data, allowed, context):
    if not isinstance(data, dict):
        raise CanvasValidationError(context + ' 必须是对象')
    unknown = set(data.keys()) - set(allowed.keys())
    if unknown:
        raise CanvasValidationError(context + ' 包含未知字段: ' + ', '.join(sorted(str(k) for k in unknown)))
    for key, check in allowed.items():
        if key in data and not check(data[key]):
            raise CanvasValidationError(context + '.' + key + ' 类型或取值不合法')

def _canvas_validate_node(node):
    if not isinstance(node, dict):
        raise CanvasValidationError('node 必须是对象')
    unknown = set(node.keys()) - {'id', 'type', 'position', 'data'}
    if unknown:
        raise CanvasValidationError('node 包含未知字段: ' + ', '.join(sorted(str(k) for k in unknown)))
    if not _canvas_valid_id(node.get('id')):
        raise CanvasValidationError('node.id 非法')
    node_type = node.get('type')
    if node_type not in CANVAS_NODE_TYPES:
        raise CanvasValidationError('未知节点类型: ' + str(node_type))
    if not _canvas_valid_position(node.get('position')):
        raise CanvasValidationError('node.position 非法或包含 NaN/Infinity')
    data = node.get('data')
    if not isinstance(data, dict):
        raise CanvasValidationError('node.data 必须是对象')
    _canvas_check_fields(data, CANVAS_NODE_DATA_SPECS[node_type], node_type + '.data')
    return node

def _canvas_validate_edge(edge):
    if not isinstance(edge, dict):
        raise CanvasValidationError('edge 必须是对象')
    unknown = set(edge.keys()) - {'id', 'source', 'target', 'label'}
    if unknown:
        raise CanvasValidationError('edge 包含未知字段: ' + ', '.join(sorted(str(k) for k in unknown)))
    for field in ('id', 'source', 'target'):
        if not _canvas_valid_id(edge.get(field)):
            raise CanvasValidationError('edge.' + field + ' 非法')
    if 'label' in edge and not _canvas_valid_opt_str(edge.get('label'), 100):
        raise CanvasValidationError('edge.label 非法')
    return edge

def _canvas_validate_edges_references(edges, node_ids):
    for edge in edges:
        if edge['source'] not in node_ids:
            raise CanvasValidationError('edge.source 引用了不存在的节点: ' + edge['source'])
        if edge['target'] not in node_ids:
            raise CanvasValidationError('edge.target 引用了不存在的节点: ' + edge['target'])

def _canvas_validate_payload(body, allow_id=False):
    if not isinstance(body, dict):
        raise CanvasValidationError('请求体必须是对象')
    _canvas_reject_sensitive_keys(body)
    top_allowed = {'id', 'title', 'created_at', 'updated_at', 'viewport', 'nodes', 'edges'} if allow_id else {'title', 'nodes', 'edges', 'viewport'}
    unknown = set(body.keys()) - top_allowed
    if unknown:
        raise CanvasValidationError('请求体包含未知字段: ' + ', '.join(sorted(str(k) for k in unknown)))
    result = {}
    if 'title' in body:
        title = body['title']
        if not isinstance(title, str) or len(title) > MAX_CANVAS_TITLE_LEN:
            raise CanvasValidationError('title 必须是不超过 100 字符的字符串')
        result['title'] = title.strip() or '未命名画布'
    if 'nodes' in body:
        nodes = body['nodes']
        if not isinstance(nodes, list) or len(nodes) > MAX_CANVAS_NODES:
            raise CanvasValidationError('nodes 必须是列表且不超过 ' + str(MAX_CANVAS_NODES) + ' 个')
        seen_ids = set()
        for node in nodes:
            _canvas_validate_node(node)
            node_id = node['id']
            if node_id in seen_ids:
                raise CanvasValidationError('node.id 重复: ' + node_id)
            seen_ids.add(node_id)
        result['nodes'] = nodes
    if 'edges' in body:
        edges = body['edges']
        if not isinstance(edges, list) or len(edges) > MAX_CANVAS_EDGES:
            raise CanvasValidationError('edges 必须是列表且不超过 ' + str(MAX_CANVAS_EDGES) + ' 个')
        seen_ids = set()
        for edge in edges:
            _canvas_validate_edge(edge)
            edge_id = edge['id']
            if edge_id in seen_ids:
                raise CanvasValidationError('edge.id 重复: ' + edge_id)
            seen_ids.add(edge_id)
        result['edges'] = edges
    if 'viewport' in body:
        viewport = body['viewport']
        if not _canvas_valid_viewport(viewport):
            raise CanvasValidationError('viewport 非法或包含 NaN/Infinity')
        result['viewport'] = viewport
    return result

def _canvas_load(user_id):
    data = load_json(canvases_path(user_id), {'version': 1, 'canvases': {}})
    if not isinstance(data, dict):
        data = {'version': 1, 'canvases': {}}
    data.setdefault('version', 1)
    if not isinstance(data.get('canvases'), dict):
        data['canvases'] = {}
    return data

def _canvas_save(user_id, data):
    save_json(canvases_path(user_id), data)

def _canvas_request_body():
    raw = request.get_data(cache=True)
    if len(raw) > MAX_CANVAS_BODY_BYTES:
        return None, (jsonify(error='请求体过大'), 413)
    if not raw:
        return None, (jsonify(error='请求体必须是 JSON'), 400)
    try:
        body = json.loads(raw)
    except (ValueError, TypeError):
        return None, (jsonify(error='请求体必须是 JSON'), 400)
    if not isinstance(body, dict):
        return None, (jsonify(error='请求体必须是对象'), 400)
    return body, None

def _canvas_list():
    with CANVAS_LOCK:
        data = _canvas_load(current_user_id())
    items = [{
        'id': canvas.get('id'),
        'title': canvas.get('title'),
        'created_at': canvas.get('created_at'),
        'updated_at': canvas.get('updated_at'),
    } for canvas in data['canvases'].values() if isinstance(canvas, dict)]
    items.sort(key=lambda item: item.get('created_at') or '', reverse=True)
    return jsonify(canvases=items)

def _canvas_create():
    body, err = _canvas_request_body()
    if err:
        return err
    try:
        parsed = _canvas_validate_payload(body, allow_id=False)
    except CanvasValidationError as exc:
        return jsonify(error=str(exc)), 400
    user_id = current_user_id()
    canvas_id = uuid.uuid4().hex
    with CANVAS_LOCK:
        data = _canvas_load(user_id)
        if len(data['canvases']) >= MAX_CANVASES_PER_USER:
            return jsonify(error='画布数量已达上限 ' + str(MAX_CANVASES_PER_USER)), 400
        nodes = parsed.get('nodes') or []
        edges = parsed.get('edges') or []
        node_ids = {node['id'] for node in nodes}
        try:
            _canvas_validate_edges_references(edges, node_ids)
        except CanvasValidationError as exc:
            return jsonify(error=str(exc)), 400
        now = datetime.now(timezone.utc).isoformat()
        canvas = {
            'id': canvas_id,
            'title': parsed.get('title', '未命名画布'),
            'created_at': now,
            'updated_at': now,
            'viewport': parsed.get('viewport', {'x': 0, 'y': 0, 'zoom': 1}),
            'nodes': nodes,
            'edges': edges,
        }
        data['canvases'][canvas_id] = canvas
        _canvas_save(user_id, data)
    return jsonify(canvas=canvas), 201

def _canvas_get(canvas_id):
    with CANVAS_LOCK:
        data = _canvas_load(current_user_id())
    canvas = data['canvases'].get(canvas_id)
    if not isinstance(canvas, dict):
        return jsonify(error='画布不存在'), 404
    return jsonify(canvas=canvas)

def _canvas_update(canvas_id):
    body, err = _canvas_request_body()
    if err:
        return err
    if 'id' in body and body['id'] != canvas_id:
        return jsonify(error='画布 ID 不匹配'), 400
    try:
        parsed = _canvas_validate_payload(body, allow_id=True)
    except CanvasValidationError as exc:
        return jsonify(error=str(exc)), 400
    user_id = current_user_id()
    with CANVAS_LOCK:
        data = _canvas_load(user_id)
        canvas = data['canvases'].get(canvas_id)
        if not isinstance(canvas, dict):
            return jsonify(error='画布不存在'), 404
        nodes = parsed.get('nodes', canvas.get('nodes') or [])
        edges = parsed.get('edges', canvas.get('edges') or [])
        node_ids = {node['id'] for node in nodes}
        try:
            _canvas_validate_edges_references(edges, node_ids)
        except CanvasValidationError as exc:
            return jsonify(error=str(exc)), 400
        if 'title' in parsed:
            canvas['title'] = parsed['title']
        if 'viewport' in parsed:
            canvas['viewport'] = parsed['viewport']
        if 'nodes' in parsed:
            canvas['nodes'] = nodes
        if 'edges' in parsed:
            canvas['edges'] = edges
        canvas['updated_at'] = datetime.now(timezone.utc).isoformat()
        _canvas_save(user_id, data)
    return jsonify(canvas=canvas)

def _canvas_delete(canvas_id):
    user_id = current_user_id()
    with CANVAS_LOCK:
        data = _canvas_load(user_id)
        if canvas_id not in data['canvases']:
            return jsonify(error='画布不存在'), 404
        del data['canvases'][canvas_id]
        _canvas_save(user_id, data)
    return jsonify(ok=True)

@app.route('/canvas')
def canvas_page():
    if not current_user_id():
        return send_from_directory('static', 'login.html')
    return send_from_directory('static', 'canvas/dist/index.html')

# ── Canvas V2 (replaces the legacy canvas workbench tab) ─────
@app.route('/canvas-v2')
def canvas_v2_page():
    if not current_user_id():
        return send_from_directory('static', 'login.html')
    return send_from_directory('static', 'canvas-v2/dist/index.html')

# ── Canvas V2 preview alias → real workbench (single entry) ──
@app.route('/canvas-v2-preview')
def canvas_v2_preview_page():
    if not current_user_id():
        return send_from_directory('static', 'login.html')
    return redirect('/canvas-v2')

@app.route('/api/canvas-v2/rollback')
@login_required
def canvas_v2_rollback():
    return jsonify(enabled=os.environ.get('CANVAS_V2_ROLLBACK', '') == '1')

@app.route('/api/canvas', methods=['GET', 'POST'])
@login_required
def canvas_collection():
    if request.method == 'POST':
        return _canvas_create()
    return _canvas_list()

@app.route('/api/canvas/<canvas_id>', methods=['GET', 'PUT', 'DELETE'])
@login_required
def canvas_item(canvas_id):
    if not _canvas_valid_id(canvas_id):
        return jsonify(error='无效的画布 ID'), 400
    if request.method == 'GET':
        return _canvas_get(canvas_id)
    if request.method == 'PUT':
        return _canvas_update(canvas_id)
    return _canvas_delete(canvas_id)

# ── Projects (additive) ─────────────────────────────────────
PROJECTS_LOCK = threading.Lock()
MAX_PROJECTS_PER_USER = 200
MAX_PROJECT_TITLE_LEN = 100
MAX_PROJECT_BODY_BYTES = 512 * 1024
MAX_PROJECT_CLASSIC_STATE_JSON_BYTES = 200 * 1024

def projects_path(user_id=None):
    return user_data_path('projects', user_id)

def _projects_load(user_id):
    with PROJECTS_LOCK:
        data = load_json(projects_path(user_id), {'version': 1, 'projects': {}})
        if not isinstance(data, dict):
            data = {'version': 1, 'projects': {}}
        data.setdefault('version', 1)
        if not isinstance(data.get('projects'), dict):
            data['projects'] = {}
        return data

def _projects_save(user_id, data):
    with PROJECTS_LOCK:
        save_json(projects_path(user_id), data)

def _ensure_projects_migrated(user_id):
    data = _projects_load(user_id)
    if data['projects']:
        return data
    canvases_data = _canvas_load(user_id)
    canvases = canvases_data.get('canvases') if isinstance(canvases_data, dict) else {}
    existing_canvas_ids = {p.get('canvas_id') for p in data['projects'].values() if isinstance(p, dict)}
    changed = False
    for cid, canvas in canvases.items():
        if not isinstance(canvas, dict):
            continue
        if cid in existing_canvas_ids:
            continue
        now = datetime.now(timezone.utc).isoformat()
        pid = uuid.uuid4().hex
        data['projects'][pid] = {
            'id': pid,
            'title': canvas.get('title') or '未命名项目',
            'canvas_id': cid,
            'cover_url': None,
            'created_at': canvas.get('created_at') or now,
            'updated_at': canvas.get('updated_at') or now,
            'last_opened_at': now,
            'classic_state': {},
        }
        existing_canvas_ids.add(cid)
        changed = True
    if changed:
        _projects_save(user_id, data)
    return data

def _projects_request_body():
    raw = request.get_data(cache=True)
    if len(raw) > MAX_PROJECT_BODY_BYTES:
        return None, (jsonify(error='请求体过大'), 413)
    if not raw:
        return {}, None
    try:
        body = json.loads(raw)
    except (ValueError, TypeError):
        return None, (jsonify(error='请求体必须是 JSON'), 400)
    if not isinstance(body, dict):
        return None, (jsonify(error='请求体必须是对象'), 400)
    return body, None

def _validate_project_payload(body):
    _canvas_reject_sensitive_keys(body)
    allowed = {'title', 'cover_url', 'classic_state', 'last_mode', 'canvas_v2_id', 'canvas_v2_migrated_at'}
    unknown = set(body.keys()) - allowed
    if unknown:
        raise CanvasValidationError('请求体包含未知字段: ' + ', '.join(sorted(str(k) for k in unknown)))
    result = {}
    if 'title' in body:
        title = body['title']
        if not isinstance(title, str) or len(title) > MAX_PROJECT_TITLE_LEN:
            raise CanvasValidationError('title 不合法')
        result['title'] = title.strip() or '未命名项目'
    if 'cover_url' in body:
        if not _canvas_valid_media_url(body['cover_url']):
            raise CanvasValidationError('cover_url 不合法')
        result['cover_url'] = body['cover_url']
    if 'classic_state' in body:
        cs = body['classic_state']
        if not isinstance(cs, dict):
            raise CanvasValidationError('classic_state 必须是对象')
        if len(json.dumps(cs, ensure_ascii=False)) > MAX_PROJECT_CLASSIC_STATE_JSON_BYTES:
            raise CanvasValidationError('classic_state 过大')
        result['classic_state'] = cs
    if 'last_mode' in body:
        if body['last_mode'] not in ('classic', 'canvas'):
            raise CanvasValidationError('last_mode 不合法')
        result['last_mode'] = body['last_mode']
    if 'canvas_v2_id' in body:
        if body['canvas_v2_id'] is not None and not _canvas_valid_id(body['canvas_v2_id']):
            raise CanvasValidationError('canvas_v2_id 不合法')
        result['canvas_v2_id'] = body['canvas_v2_id']
    if 'canvas_v2_migrated_at' in body:
        if not _canvas_valid_opt_str(body['canvas_v2_migrated_at'], 64):
            raise CanvasValidationError('canvas_v2_migrated_at 不合法')
        result['canvas_v2_migrated_at'] = body['canvas_v2_migrated_at']
    return result

def _projects_list():
    user_id = current_user_id()
    data = _ensure_projects_migrated(user_id)
    trash_only = request.args.get('trash') == '1'
    items = []
    for project in data['projects'].values():
        if not isinstance(project, dict):
            continue
        deleted = project.get('deleted_at') not in (None, '')
        if trash_only != deleted:
            continue
        items.append({
            'id': project.get('id'),
            'title': project.get('title'),
            'canvas_id': project.get('canvas_id'),
            'canvas_v2_id': project.get('canvas_v2_id'),
            'cover_url': project.get('cover_url'),
            'created_at': project.get('created_at'),
            'updated_at': project.get('updated_at'),
            'last_opened_at': project.get('last_opened_at'),
            'last_mode': project.get('last_mode', 'classic'),
            'deleted_at': project.get('deleted_at'),
        })
    items.sort(key=lambda item: item.get('last_opened_at') or item.get('updated_at') or '', reverse=True)
    return jsonify(projects=items)

def _projects_create():
    body, err = _projects_request_body()
    if err:
        return err
    title = body.get('title')
    if title is None:
        title = '未命名项目'
    elif not isinstance(title, str) or len(title) > MAX_PROJECT_TITLE_LEN:
        return jsonify(error='项目标题不合法'), 400
    else:
        title = title.strip() or '未命名项目'
    initial_mode = body.get('initial_mode')
    if initial_mode not in ('classic', 'canvas'):
        initial_mode = 'classic'
    user_id = current_user_id()
    data = _ensure_projects_migrated(user_id)
    if len(data['projects']) >= MAX_PROJECTS_PER_USER:
        return jsonify(error='项目数量已达上限 ' + str(MAX_PROJECTS_PER_USER)), 400
    now = datetime.now(timezone.utc).isoformat()
    canvas_id = uuid.uuid4().hex
    canvas = {'id': canvas_id, 'title': title, 'created_at': now, 'updated_at': now, 'viewport': {'x': 0, 'y': 0, 'zoom': 0.85}, 'nodes': [], 'edges': []}
    pid = uuid.uuid4().hex
    project = {'id': pid, 'title': title, 'canvas_id': canvas_id, 'cover_url': None, 'created_at': now, 'updated_at': now, 'last_opened_at': now, 'classic_state': {}, 'last_mode': initial_mode, 'deleted_at': None}
    data['projects'][pid] = project
    with CANVAS_LOCK:
        cdata = _canvas_load(user_id)
        cdata['canvases'][canvas_id] = canvas
        save_json_many(((canvases_path(user_id), cdata), (projects_path(user_id), data)))
    return jsonify(project=project), 201

def _projects_get(project_id):
    user_id = current_user_id()
    data = _ensure_projects_migrated(user_id)
    project = data['projects'].get(project_id)
    if not isinstance(project, dict):
        return jsonify(error='项目不存在'), 404
    # Mirror _projects_list so clients always see a known last_mode. Without
    # this, projects migrated from legacy canvas data (which never set
    # last_mode) expose last_mode=None, and the workspace shell's boot() can
    # leave the mode-chooser overlay stuck on top of everything.
    out = dict(project)
    if out.get('last_mode') not in ('classic', 'canvas'):
        out['last_mode'] = 'classic'
    return jsonify(project=out)

def _projects_update(project_id):
    body, err = _projects_request_body()
    if err:
        return err
    user_id = current_user_id()
    data = _ensure_projects_migrated(user_id)
    project = data['projects'].get(project_id)
    if not isinstance(project, dict):
        return jsonify(error='项目不存在'), 404
    try:
        parsed = _validate_project_payload(body)
    except CanvasValidationError as exc:
        return jsonify(error=str(exc)), 400
    if 'title' in parsed:
        project['title'] = parsed['title']
    if 'cover_url' in parsed:
        project['cover_url'] = parsed['cover_url']
    if 'classic_state' in parsed:
        project['classic_state'] = parsed['classic_state']
    if 'last_mode' in parsed:
        project['last_mode'] = parsed['last_mode']
    project['updated_at'] = datetime.now(timezone.utc).isoformat()
    _projects_save(user_id, data)
    return jsonify(project=project)

@app.route('/api/projects', methods=['GET', 'POST'])
@login_required
def projects_collection():
    if request.method == 'POST':
        return _projects_create()
    return _projects_list()

@app.route('/api/projects/<project_id>', methods=['GET', 'PUT'])
@login_required
def project_item(project_id):
    if not _canvas_valid_id(project_id):
        return jsonify(error='无效的项目 ID'), 400
    if request.method == 'GET':
        return _projects_get(project_id)
    return _projects_update(project_id)

@app.route('/api/projects/<project_id>/open', methods=['POST'])
@login_required
def project_open(project_id):
    if not _canvas_valid_id(project_id):
        return jsonify(error='无效的项目 ID'), 400
    user_id = current_user_id()
    data = _ensure_projects_migrated(user_id)
    project = data['projects'].get(project_id)
    if not isinstance(project, dict):
        return jsonify(error='项目不存在'), 404
    project['last_opened_at'] = datetime.now(timezone.utc).isoformat()
    _projects_save(user_id, data)
    # Return the full project (mirroring _projects_get's last_mode default) so
    # the client can skip the follow-up GET /api/projects/<id> round trip.
    out = dict(project)
    if out.get('last_mode') not in ('classic', 'canvas'):
        out['last_mode'] = 'classic'
    return jsonify(ok=True, project=out, last_opened_at=project['last_opened_at'])

@app.route('/api/projects/<project_id>', methods=['DELETE'])
@login_required
def project_delete(project_id):
    if not _canvas_valid_id(project_id):
        return jsonify(error='无效的项目 ID'), 400
    user_id = current_user_id()
    data = _ensure_projects_migrated(user_id)
    project = data['projects'].get(project_id)
    if not isinstance(project, dict):
        return jsonify(error='项目不存在'), 404
    project['deleted_at'] = datetime.now(timezone.utc).isoformat()
    project['updated_at'] = project['deleted_at']
    _projects_save(user_id, data)
    return jsonify(ok=True)

@app.route('/api/projects/<project_id>/restore', methods=['POST'])
@login_required
def project_restore(project_id):
    if not _canvas_valid_id(project_id):
        return jsonify(error='无效的项目 ID'), 400
    user_id = current_user_id()
    data = _ensure_projects_migrated(user_id)
    project = data['projects'].get(project_id)
    if not isinstance(project, dict):
        return jsonify(error='项目不存在'), 404
    project['deleted_at'] = None
    project['updated_at'] = datetime.now(timezone.utc).isoformat()
    _projects_save(user_id, data)
    return jsonify(ok=True)

@app.route('/api/projects/<project_id>/permanent', methods=['POST'])
@login_required
def project_permanent_delete(project_id):
    if not _canvas_valid_id(project_id):
        return jsonify(error='无效的项目 ID'), 400
    user_id = current_user_id()
    data = _ensure_projects_migrated(user_id)
    project = data['projects'].get(project_id)
    if not isinstance(project, dict):
        return jsonify(error='项目不存在'), 404
    canvas_id = project.get('canvas_id')
    if canvas_id:
        with CANVAS_LOCK:
            cdata = _canvas_load(user_id)
            if canvas_id in cdata.get('canvases', {}):
                del cdata['canvases'][canvas_id]
                _canvas_save(user_id, cdata)
    del data['projects'][project_id]
    _projects_save(user_id, data)
    return jsonify(ok=True)

# ── Canvas V2 (additive) ────────────────────────────────────
# New @xyflow/react canvas documents. Kept in a separate per-user store so
# the legacy canvases.json stays untouched and rollback stays possible.
CANVAS_V2_LOCK = threading.Lock()
MAX_CANVAS_V2_PER_USER = 100
MAX_CANVAS_V2_NODES = 600
MAX_CANVAS_V2_EDGES = 1200
MAX_CANVAS_V2_TITLE_LEN = 100
MAX_CANVAS_V2_NODE_DATA_BYTES = 256 * 1024
MAX_CANVAS_V2_BODY_BYTES = 2 * 1024 * 1024
CANVAS_V2_NODE_KEYS = ('id', 'type', 'position', 'data')
CANVAS_V2_EDGE_KEYS = ('id', 'source', 'target', 'label', 'type', 'sourceHandle', 'targetHandle')

def canvas_v2_path(user_id=None):
    return user_data_path('canvas_v2', user_id)

def _canvas_v2_load(user_id):
    with CANVAS_V2_LOCK:
        data = load_json(canvas_v2_path(user_id), {'version': 1, 'canvases': {}})
    if not isinstance(data, dict):
        data = {'version': 1, 'canvases': {}}
    data.setdefault('version', 1)
    if not isinstance(data.get('canvases'), dict):
        data['canvases'] = {}
    return data

def _canvas_v2_save(user_id, data):
    with CANVAS_V2_LOCK:
        save_json(canvas_v2_path(user_id), data)

def _canvas_v2_request_body():
    raw = request.get_data(cache=True)
    if len(raw) > MAX_CANVAS_V2_BODY_BYTES:
        return None, (jsonify(error='请求体过大'), 413)
    if not raw:
        return None, (jsonify(error='请求体必须是 JSON'), 400)
    try:
        body = json.loads(raw)
    except (ValueError, TypeError):
        return None, (jsonify(error='请求体必须是 JSON'), 400)
    if not isinstance(body, dict):
        return None, (jsonify(error='请求体必须是对象'), 400)
    return body, None

def _canvas_v2_reject_data_urls(value):
    """Reject data:/base64 payloads anywhere inside a v2 node's data so the
    persisted store only ever holds server-hosted URLs."""
    if isinstance(value, str):
        lowered = value.strip().lower()
        if lowered.startswith('data:') or ';base64' in lowered[:200]:
            raise CanvasValidationError('数据中包含 data: URL，不允许持久化')
        return
    if isinstance(value, dict):
        for v in value.values():
            _canvas_v2_reject_data_urls(v)
    elif isinstance(value, list):
        for v in value:
            _canvas_v2_reject_data_urls(v)

def _canvas_v2_validate_node(node):
    if not isinstance(node, dict):
        raise CanvasValidationError('node 必须是对象')
    unknown = set(node.keys()) - set(CANVAS_V2_NODE_KEYS)
    if unknown:
        raise CanvasValidationError('node 包含未知字段: ' + ', '.join(sorted(str(k) for k in unknown)))
    if not _canvas_valid_id(node.get('id')):
        raise CanvasValidationError('node.id 非法')
    node_type = node.get('type')
    if not isinstance(node_type, str) or not node_type.strip() or len(node_type) > 64:
        raise CanvasValidationError('node.type 非法')
    if not _canvas_valid_position(node.get('position')):
        raise CanvasValidationError('node.position 非法或包含 NaN/Infinity')
    data = node.get('data')
    if not isinstance(data, dict):
        raise CanvasValidationError('node.data 必须是对象')
    if len(json.dumps(data, ensure_ascii=False)) > MAX_CANVAS_V2_NODE_DATA_BYTES:
        raise CanvasValidationError('node.data 过大')
    _canvas_reject_sensitive_keys(data)
    _canvas_v2_reject_data_urls(data)
    return node

def _canvas_v2_validate_edge(edge):
    if not isinstance(edge, dict):
        raise CanvasValidationError('edge 必须是对象')
    unknown = set(edge.keys()) - set(CANVAS_V2_EDGE_KEYS)
    if unknown:
        raise CanvasValidationError('edge 包含未知字段: ' + ', '.join(sorted(str(k) for k in unknown)))
    for field in ('id', 'source', 'target'):
        if not _canvas_valid_id(edge.get(field)):
            raise CanvasValidationError('edge.' + field + ' 非法')
    if 'label' in edge and not _canvas_valid_opt_str(edge.get('label'), 200):
        raise CanvasValidationError('edge.label 非法')
    for field in ('type', 'sourceHandle', 'targetHandle'):
        if field in edge and not _canvas_valid_opt_str(edge.get(field), 64):
            raise CanvasValidationError('edge.' + field + ' 非法')
    return edge

def _canvas_v2_validate_payload(body, allow_id=False):
    if not isinstance(body, dict):
        raise CanvasValidationError('请求体必须是对象')
    _canvas_reject_sensitive_keys(body)
    top_allowed = {'id', 'title', 'viewport', 'nodes', 'edges'} if allow_id else {'title', 'viewport', 'nodes', 'edges'}
    unknown = set(body.keys()) - top_allowed
    if unknown:
        raise CanvasValidationError('请求体包含未知字段: ' + ', '.join(sorted(str(k) for k in unknown)))
    result = {}
    if 'title' in body:
        title = body['title']
        if not isinstance(title, str) or len(title) > MAX_CANVAS_V2_TITLE_LEN:
            raise CanvasValidationError('title 必须是不超过 100 字符的字符串')
        result['title'] = title.strip() or '未命名画布'
    if 'nodes' in body:
        nodes = body['nodes']
        if not isinstance(nodes, list) or len(nodes) > MAX_CANVAS_V2_NODES:
            raise CanvasValidationError('nodes 必须是列表且不超过 ' + str(MAX_CANVAS_V2_NODES) + ' 个')
        seen_ids = set()
        for node in nodes:
            _canvas_v2_validate_node(node)
            node_id = node['id']
            if node_id in seen_ids:
                raise CanvasValidationError('node.id 重复: ' + node_id)
            seen_ids.add(node_id)
        result['nodes'] = nodes
    if 'edges' in body:
        edges = body['edges']
        if not isinstance(edges, list) or len(edges) > MAX_CANVAS_V2_EDGES:
            raise CanvasValidationError('edges 必须是列表且不超过 ' + str(MAX_CANVAS_V2_EDGES) + ' 个')
        seen_ids = set()
        for edge in edges:
            _canvas_v2_validate_edge(edge)
            edge_id = edge['id']
            if edge_id in seen_ids:
                raise CanvasValidationError('edge.id 重复: ' + edge_id)
            seen_ids.add(edge_id)
        result['edges'] = edges
    if 'viewport' in body:
        viewport = body['viewport']
        if not _canvas_valid_viewport(viewport):
            raise CanvasValidationError('viewport 非法或包含 NaN/Infinity')
        result['viewport'] = viewport
    return result

def _canvas_v2_list():
    data = _canvas_v2_load(current_user_id())
    items = [{
        'id': canvas.get('id'),
        'title': canvas.get('title'),
        'created_at': canvas.get('created_at'),
        'updated_at': canvas.get('updated_at'),
    } for canvas in data['canvases'].values() if isinstance(canvas, dict)]
    items.sort(key=lambda item: item.get('created_at') or '', reverse=True)
    return jsonify(canvases=items)

def _canvas_v2_create():
    body, err = _canvas_v2_request_body()
    if err:
        return err
    try:
        parsed = _canvas_v2_validate_payload(body, allow_id=False)
    except CanvasValidationError as exc:
        return jsonify(error=str(exc)), 400
    user_id = current_user_id()
    canvas_id = uuid.uuid4().hex
    data = _canvas_v2_load(user_id)
    if len(data['canvases']) >= MAX_CANVAS_V2_PER_USER:
        return jsonify(error='画布数量已达上限 ' + str(MAX_CANVAS_V2_PER_USER)), 400
    nodes = parsed.get('nodes') or []
    edges = parsed.get('edges') or []
    node_ids = {node['id'] for node in nodes}
    try:
        for edge in edges:
            if edge['source'] not in node_ids:
                raise CanvasValidationError('edge.source 引用了不存在的节点: ' + edge['source'])
            if edge['target'] not in node_ids:
                raise CanvasValidationError('edge.target 引用了不存在的节点: ' + edge['target'])
    except CanvasValidationError as exc:
        return jsonify(error=str(exc)), 400
    now = datetime.now(timezone.utc).isoformat()
    canvas = {
        'id': canvas_id,
        'title': parsed.get('title', '未命名画布'),
        'created_at': now,
        'updated_at': now,
        'viewport': parsed.get('viewport', {'x': 0, 'y': 0, 'zoom': 1}),
        'nodes': nodes,
        'edges': edges,
    }
    data['canvases'][canvas_id] = canvas
    _canvas_v2_save(user_id, data)
    return jsonify(canvas=canvas), 201

def _canvas_v2_get(canvas_id):
    data = _canvas_v2_load(current_user_id())
    canvas = data['canvases'].get(canvas_id)
    if not isinstance(canvas, dict):
        return jsonify(error='画布不存在'), 404
    return jsonify(canvas=canvas)

def _canvas_v2_update(canvas_id):
    body, err = _canvas_v2_request_body()
    if err:
        return err
    if 'id' in body and body['id'] != canvas_id:
        return jsonify(error='画布 ID 不匹配'), 400
    try:
        parsed = _canvas_v2_validate_payload(body, allow_id=True)
    except CanvasValidationError as exc:
        return jsonify(error=str(exc)), 400
    user_id = current_user_id()
    data = _canvas_v2_load(user_id)
    canvas = data['canvases'].get(canvas_id)
    if not isinstance(canvas, dict):
        return jsonify(error='画布不存在'), 404
    nodes = parsed.get('nodes', canvas.get('nodes') or [])
    edges = parsed.get('edges', canvas.get('edges') or [])
    node_ids = {node['id'] for node in nodes}
    try:
        for edge in edges:
            if edge['source'] not in node_ids:
                raise CanvasValidationError('edge.source 引用了不存在的节点: ' + edge['source'])
            if edge['target'] not in node_ids:
                raise CanvasValidationError('edge.target 引用了不存在的节点: ' + edge['target'])
    except CanvasValidationError as exc:
        return jsonify(error=str(exc)), 400
    if 'title' in parsed:
        canvas['title'] = parsed['title']
    if 'viewport' in parsed:
        canvas['viewport'] = parsed['viewport']
    if 'nodes' in parsed:
        canvas['nodes'] = nodes
    if 'edges' in parsed:
        canvas['edges'] = edges
    canvas['updated_at'] = datetime.now(timezone.utc).isoformat()
    _canvas_v2_save(user_id, data)
    return jsonify(canvas=canvas)

def _canvas_v2_delete(canvas_id):
    user_id = current_user_id()
    data = _canvas_v2_load(user_id)
    if canvas_id not in data['canvases']:
        return jsonify(error='画布不存在'), 404
    del data['canvases'][canvas_id]
    _canvas_v2_save(user_id, data)
    return jsonify(ok=True)

@app.route('/api/canvas-v2', methods=['GET', 'POST'])
@login_required
def canvas_v2_collection():
    if request.method == 'POST':
        return _canvas_v2_create()
    return _canvas_v2_list()

@app.route('/api/canvas-v2/<canvas_id>', methods=['GET', 'PUT', 'DELETE'])
@login_required
def canvas_v2_item(canvas_id):
    if not _canvas_valid_id(canvas_id):
        return jsonify(error='无效的画布 ID'), 400
    if request.method == 'GET':
        return _canvas_v2_get(canvas_id)
    if request.method == 'PUT':
        return _canvas_v2_update(canvas_id)
    return _canvas_v2_delete(canvas_id)

def _project_v2_get_or_create(project_id):
    user_id = current_user_id()
    data = _ensure_projects_migrated(user_id)
    project = data['projects'].get(project_id)
    if not isinstance(project, dict):
        return None, jsonify(error='项目不存在'), 404
    canvas_v2_id = project.get('canvas_v2_id')
    if canvas_v2_id:
        canvas = _canvas_v2_load(user_id)['canvases'].get(canvas_v2_id)
        if isinstance(canvas, dict):
            return project, jsonify(canvas=canvas), 200
    # Create the v2 document on first open (legacy canvases.json untouched).
    canvas_id = uuid.uuid4().hex
    now = datetime.now(timezone.utc).isoformat()
    canvas = {'id': canvas_id, 'title': project.get('title') or '未命名画布', 'created_at': now, 'updated_at': now,
              'viewport': {'x': 0, 'y': 0, 'zoom': 0.85}, 'nodes': [], 'edges': []}
    cdata = _canvas_v2_load(user_id)
    cdata['canvases'][canvas_id] = canvas
    project['canvas_v2_id'] = canvas_id
    save_json_many(((canvas_v2_path(user_id), cdata), (projects_path(user_id), data)))
    return project, jsonify(canvas=canvas), 201

@app.route('/api/projects/<project_id>/canvas-v2', methods=['GET', 'PUT'])
@login_required
def project_canvas_v2(project_id):
    if not _canvas_valid_id(project_id):
        return jsonify(error='无效的项目 ID'), 400
    if request.method == 'GET':
        project, resp, status = _project_v2_get_or_create(project_id)
        if resp.status_code >= 400:
            return resp, status
        # One-time compatible read of the legacy canvas for client-side migration.
        # Only needed before the v2 document has been migrated; skip the
        # (expensive) legacy canvases load on every subsequent open.
        legacy = None
        if not project.get('canvas_v2_migrated_at'):
            legacy_id = project.get('canvas_id')
            if legacy_id:
                legacy_data = _canvas_load(current_user_id())
                legacy = legacy_data['canvases'].get(legacy_id)
        result = resp.get_json()
        result['legacy_canvas'] = legacy
        result['canvas_v2_migrated_at'] = project.get('canvas_v2_migrated_at')
        return jsonify(**result), status
    body, err = _canvas_v2_request_body()
    if err:
        return err
    canvas_v2_id = (body.get('canvas_v2_id') or '').strip()
    if not _canvas_valid_id(canvas_v2_id):
        return jsonify(error='canvas_v2_id 非法'), 400
    user_id = current_user_id()
    data = _ensure_projects_migrated(user_id)
    project = data['projects'].get(project_id)
    if not isinstance(project, dict):
        return jsonify(error='项目不存在'), 404
    canvas = _canvas_v2_load(user_id)['canvases'].get(canvas_v2_id)
    if not isinstance(canvas, dict):
        return jsonify(error='画布不存在'), 404
    project['canvas_v2_id'] = canvas_v2_id
    _projects_save(user_id, data)
    return jsonify(project=project)

# ── Skills (additive, Phase 1) ──────────────────────────────
SKILLS_DIR = os.path.join(BASE, 'prompts', 'skills')
SKILLS_REGISTRY_FILE = os.path.join(SKILLS_DIR, 'registry.json')
SKILL_REGISTRY_CACHE = None

def _load_skill_registry():
    global SKILL_REGISTRY_CACHE
    if SKILL_REGISTRY_CACHE is not None:
        return SKILL_REGISTRY_CACHE
    try:
        with open(SKILLS_REGISTRY_FILE, 'r', encoding='utf-8') as fh:
            data = json.load(fh)
    except Exception:
        data = {'version': 1, 'skills': []}
    skills = data.get('skills') if isinstance(data, dict) else []
    if not isinstance(skills, list):
        skills = []
    SKILL_REGISTRY_CACHE = skills
    return skills

def _public_skill_meta(skill):
    if not isinstance(skill, dict):
        return {}
    allowed = ('id', 'source', 'version', 'title', 'description', 'input_schema', 'output_type', 'target_node_type', 'consumes_text_points')
    return {k: skill.get(k) for k in allowed}

@app.route('/api/skills', methods=['GET'])
@login_required
def skills_list():
    return jsonify(skills=[_public_skill_meta(s) for s in _load_skill_registry() if isinstance(s, dict)])

@app.route('/api/skills/<skill_id>', methods=['GET'])
@login_required
def skills_get(skill_id):
    skill = next((s for s in _load_skill_registry() if isinstance(s, dict) and s.get('id') == skill_id), None)
    if not skill:
        return jsonify(error='技能不存在'), 404
    return jsonify(skill=_public_skill_meta(skill))

@app.route('/api/skills/<skill_id>/run', methods=['POST'])
@login_required
@model_access_required('text')
def skills_run(skill_id):
    """Run a Seedance / Drama skill with the text model only. The result is
    returned as an editable draft; it never triggers image or video calls."""
    body = request.json or {}
    skill = next((s for s in _load_skill_registry() if isinstance(s, dict) and s.get('id') == skill_id), None)
    if not skill:
        return jsonify(error='技能不存在'), 404
    prompt_file = skill.get('prompt_file')
    if not prompt_file:
        return jsonify(error='技能未配置提示词'), 400
    prompt_path = os.path.join(SKILLS_DIR, prompt_file)
    if not os.path.isfile(prompt_path):
        return jsonify(error='技能提示词缺失'), 500
    try:
        with open(prompt_path, 'r', encoding='utf-8') as fh:
            system_prompt = fh.read()
    except Exception as e:
        return jsonify(error=f'技能提示词读取失败: {e}'), 500
    input_data = body.get('input')
    if not isinstance(input_data, dict):
        return jsonify(error='input 必须是对象'), 400
    user_input = json.dumps(input_data, ensure_ascii=False, sort_keys=True)
    script_model = body.get('script_model', SCRIPT_MODEL_DEFAULT)
    force_personal = bool(body.get('use_personal_api')) or script_model == 'personal-api'
    if not force_personal and script_model not in SCRIPT_MODELS:
        return jsonify(error='无效的文本模型'), 400
    point_cost = 0
    try:
        api_cfg = resolve_api(
            'text', builtin_text_api(SCRIPT_MODEL_DEFAULT if force_personal else script_model),
            current_user_id(), force_personal=force_personal, profile_id=body.get('api_profile_id'),
            strict_builtin='use_personal_api' in body and not force_personal
        )
        point_cost = reserve_model_points('text', script_model, current_user_id(), personal=force_personal)
        text = call_script_text_model(script_model, system_prompt, user_input, temperature=0.7, max_tokens=2000, api_cfg=api_cfg)
        return jsonify(text=text, points=point_cost)
    except QuotaError as e:
        return jsonify(error=str(e)), 402
    except Exception as e:
        refund_model_points(current_user_id(), point_cost)
        return jsonify(error=f'生成失败: {str(e)}'), 500



# ── Canvas import/export (additive, Phase 2) ────────────────
MAX_IMPORT_BYTES = 512 * 1024

def _validate_canvas_import_body(body):
    _canvas_reject_sensitive_keys(body)
    if not isinstance(body, dict):
        raise CanvasValidationError('导入内容必须是对象')
    if body.get('format') != 'manga-studio-canvas':
        raise CanvasValidationError('不支持的导入格式')
    if body.get('schema_version') != 1:
        raise CanvasValidationError('不支持的 schema_version')
    inner = {}
    if 'title' in body:
        inner['title'] = body['title']
    for key in ('nodes', 'edges', 'viewport'):
        if key in body:
            inner[key] = body[key]
    return _canvas_validate_payload(inner, allow_id=False)

def _create_project_with_canvas(user_id, title, canvas_payload, initial_mode='canvas'):
    now = datetime.now(timezone.utc).isoformat()
    canvas_id = uuid.uuid4().hex
    canvas = {
        'id': canvas_id,
        'title': title,
        'created_at': now,
        'updated_at': now,
        'viewport': canvas_payload.get('viewport', {'x': 0, 'y': 0, 'zoom': 0.85}),
        'nodes': canvas_payload.get('nodes', []),
        'edges': canvas_payload.get('edges', []),
    }
    with CANVAS_LOCK:
        cdata = _canvas_load(user_id)
        cdata['canvases'][canvas_id] = canvas
        _canvas_save(user_id, cdata)
    project_id = uuid.uuid4().hex
    project = {
        'id': project_id,
        'title': title,
        'canvas_id': canvas_id,
        'cover_url': None,
        'created_at': now,
        'updated_at': now,
        'last_opened_at': now,
        'classic_state': {},
        'last_mode': initial_mode,
        'deleted_at': None,
    }
    try:
        data = _projects_load(user_id)
        data['projects'][project_id] = project
        _projects_save(user_id, data)
    except Exception:
        with CANVAS_LOCK:
            cdata = _canvas_load(user_id)
            cdata['canvases'].pop(canvas_id, None)
            _canvas_save(user_id, cdata)
        raise
    return project

@app.route('/api/canvas/<canvas_id>/export', methods=['GET'])
@login_required
def canvas_export(canvas_id):
    if not _canvas_valid_id(canvas_id):
        return jsonify(error='无效的画布 ID'), 400
    user_id = current_user_id()
    with CANVAS_LOCK:
        data = _canvas_load(user_id)
    canvas = data['canvases'].get(canvas_id)
    if not isinstance(canvas, dict):
        return jsonify(error='画布不存在'), 404
    payload = {
        'format': 'manga-studio-canvas',
        'schema_version': 1,
        'exported_at': datetime.now(timezone.utc).isoformat(),
        'title': canvas.get('title') or '未命名画布',
        'viewport': canvas.get('viewport') or {'x': 0, 'y': 0, 'zoom': 0.85},
        'nodes': canvas.get('nodes') or [],
        'edges': canvas.get('edges') or [],
    }
    return jsonify(payload)

@app.route('/api/projects/import', methods=['POST'])
@login_required
def project_import():
    raw = request.get_data(cache=True)
    if len(raw) > MAX_IMPORT_BYTES:
        return jsonify(error='文件过大'), 413
    if not raw:
        return jsonify(error='导入内容为空'), 400
    try:
        body = json.loads(raw)
    except (ValueError, TypeError):
        return jsonify(error='导入内容不是合法 JSON'), 400
    try:
        parsed = _validate_canvas_import_body(body)
    except CanvasValidationError as exc:
        return jsonify(error=str(exc)), 400
    user_id = current_user_id()
    title = parsed.get('title', '导入项目')
    try:
        project = _create_project_with_canvas(user_id, title, parsed, initial_mode='canvas')
    except Exception as exc:
        return jsonify(error='导入失败: ' + str(exc)), 500
    return jsonify(project=project), 201


@app.route('/api/auth/register', methods=['POST'])
def register():
    body = request.json or {}
    username = (body.get('username') or '').strip()
    password = body.get('password') or ''
    invite_code = (body.get('invite_code') or '').strip().upper()
    valid_name = 3 <= len(username) <= 120 and re.fullmatch(r'[A-Za-z0-9._@+-]+', username)
    valid_email = '@' not in username or re.fullmatch(r'[^@\s]+@[^@\s]+\.[^@\s]+', username)
    if not valid_name or not valid_email:
        return jsonify(error='请输入有效的邮箱或用户名'), 400
    if len(password) < 8:
        return jsonify(error='密码至少需要 8 位'), 400
    user_id = username.lower()
    with USERS_LOCK:
        users = load_json(users_path(), {})
        if user_id in users:
            return jsonify(error='用户名已存在'), 409
        invitations = load_json(invitations_path(), [])
        invite = next((item for item in invitations if item.get('code') == invite_code and item.get('active', True) and item.get('used', 0) < item.get('max_uses', 1)), None)
        if not invite:
            return jsonify(error='邀请码无效或已用完'), 400
        invite['used'] = invite.get('used', 0) + 1
        users[user_id] = {'username': username, 'password_hash': generate_password_hash(password, method='pbkdf2:sha256'), 'created_at': datetime.now(timezone.utc).isoformat(), 'model_permissions': invite.get('permissions', []), 'points': invite.get('points', 0), 'invite_code': invite_code}
        save_json(users_path(), users)
        save_json(invitations_path(), invitations)
    session.clear()
    session['user_id'] = user_id
    return jsonify(ok=True, username=username)

@app.route('/api/auth/login', methods=['POST'])
def login():
    body = request.json or {}
    user_id = (body.get('username') or '').strip().lower()
    user = load_json(users_path(), {}).get(user_id)
    if not user or not check_password_hash(user.get('password_hash', ''), body.get('password') or ''):
        return jsonify(error='用户名或密码错误'), 401
    if user.get('disabled'):
        return jsonify(error='账号已停用，请联系管理员'), 403
    session.clear()
    session['user_id'] = user_id
    return jsonify(ok=True, username=user.get('username', user_id))

@app.route('/api/auth/logout', methods=['POST'])
def logout():
    session.clear()
    return jsonify(ok=True)

@app.route('/api/auth/me')
@login_required
def auth_me():
    user = load_json(users_path(), {}).get(current_user_id(), {})
    admin = is_admin()
    return jsonify(username=user.get('username', current_user_id()), is_admin=admin, model_permissions=user.get('model_permissions', []), points=None if admin else user.get('points', 0), unlimited_points=admin)

@app.route('/api/settings', methods=['GET', 'POST'])
@login_required
def api_settings():
    if request.method == 'POST':
        body = request.json or {}
        kind = body.get('kind')
        if kind not in ('text', 'image', 'video'): return jsonify(error='无效接口类型'), 400
        profiles, _ = get_api_profiles(kind)
        settings = load_json(settings_path(), {})
        profile_id = str(body.get('profile_id') or '').strip()
        is_new_profile = not profile_id
        if profile_id:
            cfg = next((item for item in profiles if item['id'] == profile_id), None)
            if not cfg: return jsonify(error='接口配置不存在'), 404
        else:
            cfg = normalize_api_profile({'id': uuid.uuid4().hex})
            profiles.append(cfg)
            profile_id = cfg['id']
        if 'name' in body:
            cfg['name'] = str(body.get('name') or '').strip()
        for key in ('provider', 'base_url', 'model', 'submit_url', 'status_url'):
            if key in body: cfg[key] = str(body.get(key) or '').strip().rstrip('/') if key == 'base_url' else str(body.get(key) or '').strip()
        if body.get('api_key'): cfg['api_key'] = str(body['api_key']).strip()
        custom_full_endpoint = cfg.get('provider') == 'custom' and (
            is_new_profile or cfg.get('endpoint_mode') == 'full' or bool(cfg.get('submit_url'))
        )
        if custom_full_endpoint:
            if not cfg.get('submit_url'):
                return jsonify(error='新建 Custom API 必须填写完整提交接口 URL'), 400
            if kind == 'video' and not cfg.get('status_url'):
                return jsonify(error='新建视频 Custom API 必须填写完整查询接口 URL，并用 {task_id} 标记任务 ID'), 400
            if kind == 'video' and '{task_id}' not in str(cfg.get('status_url')).replace('{taskId}', '{task_id}'):
                return jsonify(error='视频查询接口必须包含 {task_id} 占位符'), 400
            cfg['endpoint_mode'] = 'full'
        try:
            cfg = normalize_api_profile(cfg)
        except ValueError as exc:
            return jsonify(error=str(exc)), 400
        profiles = [cfg if item['id'] == profile_id else item for item in profiles]
        settings.setdefault('api_profiles', {})[kind] = profiles
        settings.setdefault('selected_api_profiles', {})[kind] = profile_id
        settings.setdefault('apis', {})[kind] = dict(cfg)
        save_json(settings_path(), settings)
        return jsonify(ok=True, profile=public_api_profile(cfg, kind), selected_profile_id=profile_id)

    apis = {}
    api_profiles = {}
    selected_profiles = {}
    for kind in ('text', 'image', 'video'):
        profiles, selected = get_api_profiles(kind)
        selected_cfg = next((item for item in profiles if item['id'] == selected), {})
        apis[kind] = public_api_profile(selected_cfg, kind)
        api_profiles[kind] = [public_api_profile(item, kind) for item in profiles]
        selected_profiles[kind] = selected
    admin = is_admin()
    return jsonify(
        apis=apis, api_profiles=api_profiles, selected_api_profiles=selected_profiles,
        using_personal=use_personal_api(),
        points=None if admin else load_json(users_path(), {}).get(current_user_id(), {}).get('points', 0),
        unlimited_points=admin
    )

@app.route('/api/settings/<kind>/<profile_id>/select', methods=['POST'])
@login_required
def select_personal_api(kind, profile_id):
    if kind not in ('text', 'image', 'video'): return jsonify(error='无效接口类型'), 400
    profiles, _ = get_api_profiles(kind)
    cfg = next((item for item in profiles if item['id'] == profile_id), None)
    if not cfg: return jsonify(error='接口配置不存在'), 404
    settings = load_json(settings_path(), {})
    settings.setdefault('selected_api_profiles', {})[kind] = profile_id
    settings.setdefault('apis', {})[kind] = dict(cfg)
    save_json(settings_path(), settings)
    return jsonify(ok=True, selected_profile_id=profile_id)

@app.route('/api/settings/<kind>/<profile_id>', methods=['DELETE'])
@login_required
def delete_personal_api(kind, profile_id):
    if kind not in ('text', 'image', 'video'): return jsonify(error='无效接口类型'), 400
    profiles, selected = get_api_profiles(kind)
    if not any(item['id'] == profile_id for item in profiles): return jsonify(error='接口配置不存在'), 404
    profiles = [item for item in profiles if item['id'] != profile_id]
    if selected == profile_id:
        selected = profiles[0]['id'] if profiles else None
    settings = load_json(settings_path(), {})
    settings.setdefault('api_profiles', {})[kind] = profiles
    settings.setdefault('selected_api_profiles', {})[kind] = selected
    if selected:
        settings.setdefault('apis', {})[kind] = dict(next(item for item in profiles if item['id'] == selected))
    else:
        settings.setdefault('apis', {}).pop(kind, None)
    save_json(settings_path(), settings)
    return jsonify(ok=True, selected_profile_id=selected)

def test_personal_api(kind, profile_id=None):
    profiles, selected = get_api_profiles(kind)
    target_id = profile_id or selected
    cfg = next((item for item in profiles if item['id'] == target_id), None)
    if not cfg: return jsonify(error='接口配置不存在'), 404
    if not cfg.get('api_key') or not (cfg.get('base_url') or cfg.get('submit_url')) or not cfg.get('model'): return jsonify(error='请先保存完整接口配置'), 400
    try:
        if cfg.get('endpoint_mode') == 'full':
            if not cfg.get('submit_url'):
                return jsonify(error='缺少完整提交接口 URL'), 400
            if kind == 'video' and (not cfg.get('status_url') or '{task_id}' not in cfg.get('status_url')):
                return jsonify(error='视频查询接口必须包含 {task_id} 占位符'), 400
            tested_at = datetime.now(timezone.utc).isoformat()
            cfg['last_test'] = tested_at
            settings = load_json(settings_path(), {})
            settings.setdefault('api_profiles', {})[kind] = [cfg if item['id'] == target_id else item for item in profiles]
            if selected == target_id:
                settings.setdefault('apis', {})[kind] = dict(cfg)
            save_json(settings_path(), settings)
            return jsonify(ok=True, message='完整接口格式有效；为避免计费，测试不会提交生成请求', last_test=tested_at)
        headers = {'Authorization': f"Bearer {cfg['api_key']}", 'x-api-key': cfg['api_key']}
        if cfg.get('provider') == 'anthropic':
            headers['anthropic-version'] = '2023-06-01'
        base_url = ATLAS_API_BASE if cfg.get('provider') == 'atlas' else cfg['base_url']
        if kind == 'video' and cfg.get('provider') == 'minimax':
            base_url = minimax_api_root(base_url)
            r = requests.get(f"{base_url}/v2/query/video_generation/0", headers=headers, timeout=20)
            response_text = r.text.lower()
            task_not_found = 'record not found' in response_text and '(1000)' in response_text
            if r.status_code in (401, 403) or (r.status_code >= 500 and not task_not_found) or (r.status_code == 404 and 'page not found' in response_text):
                return jsonify(error=f'连接失败：HTTP {r.status_code} {r.text[:160]}'), 502
        else:
            r = requests.get(f"{base_url}/models", headers=headers, timeout=20)
            if r.status_code not in (200, 201): return jsonify(error=f'连接失败：HTTP {r.status_code} {r.text[:160]}'), 502
        tested_at = datetime.now(timezone.utc).isoformat()
        cfg['last_test'] = tested_at
        settings = load_json(settings_path(), {})
        settings.setdefault('api_profiles', {})[kind] = [cfg if item['id'] == target_id else item for item in profiles]
        if selected == target_id:
            settings.setdefault('apis', {})[kind] = dict(cfg)
        save_json(settings_path(), settings)
        return jsonify(ok=True, message='连接成功，额度用完后将使用此接口', last_test=tested_at)
    except requests.RequestException as e:
        return jsonify(error=f'连接失败：{str(e)}'), 502

@app.route('/api/settings/<kind>/test', methods=['POST'])
@login_required
def personal_api_test(kind):
    if kind not in ('text', 'image', 'video'): return jsonify(error='无效接口类型'), 400
    return test_personal_api(kind)

@app.route('/api/settings/<kind>/<profile_id>/test', methods=['POST'])
@login_required
def personal_api_profile_test(kind, profile_id):
    if kind not in ('text', 'image', 'video'): return jsonify(error='无效接口类型'), 400
    return test_personal_api(kind, profile_id)

@app.route('/api/admin/users')
@admin_required
def admin_users():
    users = load_json(users_path(), {})
    result = []
    for user_id, user in users.items():
        history = load_json(history_path(user_id), [])
        admin = is_admin(user_id)
        result.append({'id': user_id, 'username': user.get('username', user_id), 'created_at': user.get('created_at', ''), 'disabled': bool(user.get('disabled')), 'is_admin': admin, 'history_count': len(history) if isinstance(history, list) else 0, 'permissions': user.get('model_permissions', []), 'points': None if admin else user.get('points', 0)})
    result.sort(key=lambda item: item['created_at'], reverse=True)
    return jsonify(users=result, total=len(result))

@app.route('/api/admin/users/<user_id>/status', methods=['POST'])
@admin_required
def admin_user_status(user_id):
    user_id = user_id.lower()
    users = load_json(users_path(), {})
    if user_id not in users: return jsonify(error='用户不存在'), 404
    disabled = bool((request.json or {}).get('disabled'))
    if is_admin(user_id) and disabled: return jsonify(error='不能停用管理员账号'), 400
    users[user_id]['disabled'] = disabled
    save_json(users_path(), users)
    return jsonify(ok=True)

@app.route('/api/admin/users/<user_id>/points', methods=['POST'])
@admin_required
def admin_user_points(user_id):
    user_id = user_id.lower()
    users = load_json(users_path(), {})
    if user_id not in users: return jsonify(error='用户不存在'), 404
    if is_admin(user_id): return jsonify(error='管理员账号不使用积分'), 400
    try:
        points = int((request.json or {}).get('points'))
    except (TypeError, ValueError):
        return jsonify(error='积分必须是整数'), 400
    if points < 0: return jsonify(error='积分不能小于 0'), 400
    with USERS_LOCK:
        users = load_json(users_path(), {})
        if user_id not in users: return jsonify(error='用户不存在'), 404
        users[user_id]['points'] = points
        save_json(users_path(), users)
    return jsonify(ok=True, points=points)

@app.route('/api/admin/model-pricing', methods=['GET', 'POST'])
@admin_required
def admin_model_pricing():
    catalog = model_pricing_catalog()
    allowed = {(item['kind'], item['model']) for item in catalog}
    if request.method == 'POST':
        body = request.json or {}
        kind = str(body.get('kind') or '')
        model = str(body.get('model') or '')
        if (kind, model) not in allowed: return jsonify(error='模型不在平台内置列表中'), 400
        try:
            points = int(body.get('points'))
        except (TypeError, ValueError):
            return jsonify(error='积分必须是整数'), 400
        if points < 0: return jsonify(error='积分不能小于 0'), 400
        pricing = load_model_pricing()
        pricing.setdefault(kind, {})[model] = points
        save_json(model_pricing_path(), pricing)
    pricing = load_model_pricing()
    rows = [dict(item, points=max(0, int((pricing.get(item['kind']) or {}).get(item['model']) or 0))) for item in catalog]
    return jsonify(models=rows)

@app.route('/api/admin/personal-apis')
@admin_required
def admin_personal_apis():
    users = load_json(users_path(), {})
    rows = []
    for user_id, user in users.items():
        for kind in ('text', 'image', 'video'):
            profiles, selected = get_api_profiles(kind, user_id)
            for profile in profiles:
                rows.append({
                    'user_id': user_id,
                    'username': user.get('username', user_id),
                    'kind': kind,
                    'id': profile['id'],
                    'name': profile.get('name', ''),
                    'provider': profile.get('provider', ''),
                    'base_url': profile.get('base_url', ''),
                    'model': profile.get('model', ''),
                    'configured': bool(profile.get('api_key')),
                    'selected': profile['id'] == selected,
                    'created_at': profile.get('created_at', '')
                })
    rows.sort(key=lambda item: item.get('created_at', ''), reverse=True)
    return jsonify(profiles=rows, total=len(rows))

@app.route('/api/admin/users/<user_id>/apis/<kind>/<profile_id>', methods=['DELETE'])
@admin_required
def admin_delete_personal_api(user_id, kind, profile_id):
    user_id = user_id.lower()
    if user_id not in load_json(users_path(), {}): return jsonify(error='用户不存在'), 404
    if kind not in ('text', 'image', 'video'): return jsonify(error='无效接口类型'), 400
    profiles, selected = get_api_profiles(kind, user_id)
    if not any(item['id'] == profile_id for item in profiles): return jsonify(error='接口配置不存在'), 404
    profiles = [item for item in profiles if item['id'] != profile_id]
    if selected == profile_id:
        selected = profiles[0]['id'] if profiles else None
    path = settings_path(user_id)
    settings = load_json(path, {})
    settings.setdefault('api_profiles', {})[kind] = profiles
    settings.setdefault('selected_api_profiles', {})[kind] = selected
    if selected:
        settings.setdefault('apis', {})[kind] = dict(next(item for item in profiles if item['id'] == selected))
    else:
        settings.setdefault('apis', {}).pop(kind, None)
    save_json(path, settings)
    return jsonify(ok=True, selected_profile_id=selected)

@app.route('/api/admin/invitations', methods=['GET', 'POST'])
@admin_required
def admin_invitations():
    invitations = load_json(invitations_path(), [])
    if request.method == 'POST':
        body = request.json or {}
        permissions = [kind for kind in ('text', 'image', 'video') if kind in body.get('permissions', [])]
        code = (body.get('code') or secrets.token_hex(4)).strip().upper()
        if any(item.get('code') == code for item in invitations): return jsonify(error='邀请码已存在'), 409
        invitations.insert(0, {'code': code, 'permissions': permissions, 'points': int(body.get('points') or 0), 'max_uses': max(1, int(body.get('max_uses') or 1)), 'used': 0, 'active': True, 'created_at': datetime.now(timezone.utc).isoformat()})
        save_json(invitations_path(), invitations)
    return jsonify(invitations=invitations)

@app.route('/api/admin/invitations/<code>/status', methods=['POST'])
@admin_required
def admin_invitation_status(code):
    invitations = load_json(invitations_path(), [])
    invite = next((item for item in invitations if item.get('code') == code.upper()), None)
    if not invite: return jsonify(error='邀请码不存在'), 404
    invite['active'] = bool((request.json or {}).get('active'))
    save_json(invitations_path(), invitations)
    return jsonify(ok=True)

@app.route('/api/admin/apis')
@admin_required
def admin_apis():
    r2_fields = {
        'R2_ACCOUNT_ID': R2_ACCOUNT_ID,
        'R2_ACCESS_KEY_ID': R2_ACCESS_KEY_ID,
        'R2_SECRET_ACCESS_KEY': R2_SECRET_ACCESS_KEY,
        'R2_BUCKET': R2_BUCKET,
        'R2_PUBLIC_BASE': R2_PUBLIC_BASE,
    }
    return jsonify(apis={
        'text': {'configured': bool(ARK_API_KEY or NANO_GPT_API_KEY)},
        'image': {'configured': bool(ARK_API_KEY or NANO_GPT_API_KEY or AGNES_API_KEY)},
        'video': {'configured': bool(ARK_API_KEY or NANO_GPT_API_KEY or AGNES_API_KEY or THIRD_PARTY_API_KEY)}
    }, storage={
        'backend': active_storage_name(),
        'persistent': persistent_storage_configured(),
        'r2_configured': r2_configured(),
        'r2_missing': [key for key, value in r2_fields.items() if not value],
    })

# ── uploads ───────────────────────────────────────────────────
@app.route('/api/tos-presign', methods=['POST'])
@login_required
def tos_presign():
    body = request.json or {}
    filename = secure_filename(body.get('filename') or 'upload')
    content_type = body.get('content_type') or 'application/octet-stream'
    ext = os.path.splitext(filename)[1].lower()
    date_prefix = datetime.now().strftime('%Y%m%d')
    object_key = f"uploads/{date_prefix}/{uuid.uuid4().hex}{ext}"
    try:
        if r2_configured():
            upload_url = get_r2_client().generate_presigned_url(
                'put_object',
                Params={'Bucket': R2_BUCKET, 'Key': object_key, 'ContentType': content_type},
                ExpiresIn=3600,
            )
            return jsonify({
                'upload_url': upload_url,
                'public_url': r2_public_url(object_key),
                'object_key': object_key,
                'storage': 'r2'
            })

        client = get_tos_client()
        from tos.enum import HttpMethodType
        signed = client.pre_signed_url(
            HttpMethodType.Http_Method_Put,
            TOS_BUCKET,
            object_key,
            expires=3600,
            header={'Content-Type': content_type}
        )
        upload_url = getattr(signed, 'signed_url', None) or getattr(signed, 'url', None) or str(signed)
        return jsonify({
            'upload_url': upload_url,
            'public_url': f"{TOS_PUBLIC_BASE}/{object_key}",
            'object_key': object_key,
            'storage': 'tos'
        })
    except Exception as e:
        print(f'[TOS] presign failed: {e}')
        return jsonify(error=str(e)), 500

@app.route('/api/upload', methods=['POST'])
@login_required
def upload():
    f = request.files.get('file')
    if not f: return jsonify(error='no file'), 400
    ext  = os.path.splitext(secure_filename(f.filename))[1].lower()
    name = uuid.uuid4().hex + ext
    ct = f.content_type or 'application/octet-stream'
    file_bytes = f.read()

    # Upload to durable object storage
    public_url, ok = upload_to_tos(file_bytes, name, ct)
    if ok:
        return jsonify(url=public_url, name=name, storage=storage_name_for_url(public_url))

    if persistent_storage_configured():
        return jsonify(error='文件存储失败，请稍后重试'), 503

    # Fallback to local
    path = os.path.join(UPLOAD, name)
    with open(path, 'wb') as fw:
        fw.write(file_bytes)
    return jsonify(url=f'/static/uploads/{name}', name=name, storage='local')

# ── characters ────────────────────────────────────────────────
@app.route('/api/characters', methods=['GET'])
@login_required
def get_chars():
    return jsonify(load_json(characters_path(), {}))

@app.route('/api/characters', methods=['POST'])
@login_required
def save_chars():
    data = request.json
    save_json(characters_path(), data)
    return jsonify(ok=True)

# ── asset libraries (outfits / scenes / audios) ───────────────
@app.route('/api/assets/<cat>', methods=['GET'])
@login_required
def get_assets(cat):
    return jsonify(load_json(assets_path(cat), []))

@app.route('/api/assets/<cat>', methods=['POST'])
@login_required
def save_assets_api(cat):
    data = request.json
    save_json(assets_path(cat), data)
    return jsonify(ok=True)


# ── asset item CRUD (additive) ───────────────────────────────
ASSET_CATS = ('outfits', 'scenes', 'audios', 'uploads', 'styles')

def _assets_request_body():
    raw = request.get_data(cache=True)
    if len(raw) > 512 * 1024:
        return None, (jsonify(error='请求体过大'), 413)
    if not raw:
        return None, (jsonify(error='请求体必须是 JSON'), 400)
    try:
        body = json.loads(raw)
    except (ValueError, TypeError):
        return None, (jsonify(error='请求体必须是 JSON'), 400)
    if not isinstance(body, dict):
        return None, (jsonify(error='请求体必须是对象'), 400)
    return body, None

def _valid_asset_name(v):
    return isinstance(v, str) and 1 <= len(v.strip()) <= 200

def _valid_asset_url(v):
    if not isinstance(v, str) or len(v) > 2000:
        return False
    lowered = v.lower()
    if lowered.startswith('data:') or ';base64' in lowered:
        return False
    return v.startswith('http://') or v.startswith('https://') or v.startswith('/static/')

@app.route('/api/assets/<cat>/item', methods=['POST'])
@login_required
def create_asset_item(cat):
    body, err = _assets_request_body()
    if err:
        return err
    user_id = current_user_id()
    if cat in ('character', 'characters'):
        name = body.get('name')
        images = body.get('images') or []
        media = body.get('media') or []
        if not _valid_asset_name(name):
            return jsonify(error='名称不合法'), 400
        if not isinstance(images, list) or not all(_valid_asset_url(u) for u in images):
            return jsonify(error='图片 URL 不合法'), 400
        if not isinstance(media, list) or not all(
            isinstance(item, dict)
            and item.get('kind') in ('video', 'audio')
            and _valid_asset_url(item.get('url'))
            for item in media
        ):
            return jsonify(error='媒体 URL 不合法'), 400
        chars = load_json(characters_path(user_id), {})
        if not isinstance(chars, dict):
            chars = {}
        item_id = uuid.uuid4().hex
        chars[item_id] = {'name': name.strip(), 'images': [{'url': u} for u in images]}
        if media:
            chars[item_id]['media'] = [
                {'url': item['url'], 'kind': item['kind']}
                for item in media
            ]
        save_json(characters_path(user_id), chars)
        return jsonify(ok=True, id=item_id), 201
    if cat not in ASSET_CATS:
        return jsonify(error='不支持的分类'), 400
    name = body.get('name')
    url = body.get('url') or body.get('thumbnail_url') or ''
    if not _valid_asset_name(name):
        return jsonify(error='名称不合法'), 400
    if not _valid_asset_url(url):
        return jsonify(error='URL 不合法'), 400
    items = load_json(assets_path(cat, user_id), [])
    if not isinstance(items, list):
        items = []
    media_kind = body.get('kind') if body.get('kind') in ('image', 'video', 'audio') else 'image'
    item = {'id': uuid.uuid4().hex, 'name': name.strip(), 'url': url, 'kind': media_kind}
    if cat == 'styles':
        item = {'id': uuid.uuid4().hex, 'name': name.strip(), 'thumbnail_url': url, 'prompt': str(body.get('prompt') or '')[:5000], 'negative_prompt': str(body.get('negative_prompt') or '')[:5000], 'use_for_image': bool(body.get('use_for_image', True)), 'use_for_video': bool(body.get('use_for_video', True))}
    items.append(item)
    save_json(assets_path(cat, user_id), items)
    return jsonify(ok=True, id=item['id']), 201

@app.route('/api/assets/<cat>/item/<item_id>/media', methods=['POST'])
@login_required
def append_asset_item_media(cat, item_id):
    body, err = _assets_request_body()
    if err:
        return err
    url = body.get('url') or ''
    media_kind = body.get('kind') if body.get('kind') in ('image', 'video', 'audio') else 'image'
    if not _valid_asset_url(url):
        return jsonify(error='URL 不合法'), 400

    user_id = current_user_id()
    if cat in ('character', 'characters'):
        chars = load_json(characters_path(user_id), {})
        if not isinstance(chars, dict) or item_id not in chars or chars[item_id].get('deleted_at'):
            return jsonify(error='资产不存在'), 404
        character = chars[item_id]
        if media_kind == 'image':
            images = character.get('images') if isinstance(character.get('images'), list) else []
            if not any(isinstance(item, dict) and item.get('url') == url for item in images):
                images.append({'url': url})
            character['images'] = images
        else:
            media = character.get('media') if isinstance(character.get('media'), list) else []
            if not any(isinstance(item, dict) and item.get('url') == url for item in media):
                media.append({'url': url, 'kind': media_kind})
            character['media'] = media
        save_json(characters_path(user_id), chars)
        return jsonify(ok=True, id=item_id)

    if cat not in ASSET_CATS:
        return jsonify(error='不支持的分类'), 400
    items = load_json(assets_path(cat, user_id), [])
    if not isinstance(items, list):
        items = []
    parent = next((item for item in items if isinstance(item, dict) and item.get('id') == item_id and not item.get('deleted_at')), None)
    if not parent:
        return jsonify(error='资产不存在'), 404
    child = {
        'id': uuid.uuid4().hex,
        'parent_id': parent.get('parent_id') or parent.get('id'),
        'name': str(parent.get('name') or '未命名素材')[:200],
        'url': url,
        'kind': media_kind,
    }
    items.append(child)
    save_json(assets_path(cat, user_id), items)
    return jsonify(ok=True, id=child['id']), 201

@app.route('/api/assets/<cat>/item/<item_id>', methods=['PUT', 'DELETE'])
@login_required
def update_asset_item(cat, item_id):
    user_id = current_user_id()
    if cat in ('character', 'characters'):
        chars = load_json(characters_path(user_id), {})
        if not isinstance(chars, dict) or item_id not in chars:
            return jsonify(error='资产不存在'), 404
        if request.method == 'DELETE':
            chars[item_id]['deleted_at'] = datetime.now(timezone.utc).isoformat()
            save_json(characters_path(user_id), chars)
            return jsonify(ok=True)
        body, err = _assets_request_body()
        if err:
            return err
        name = body.get('name')
        if not _valid_asset_name(name):
            return jsonify(error='名称不合法'), 400
        chars[item_id]['name'] = name.strip()
        save_json(characters_path(user_id), chars)
        return jsonify(ok=True)
    if cat not in ASSET_CATS:
        return jsonify(error='不支持的分类'), 400
    items = load_json(assets_path(cat, user_id), [])
    if not isinstance(items, list):
        items = []
    item = next((x for x in items if isinstance(x, dict) and x.get('id') == item_id), None)
    if not item:
        return jsonify(error='资产不存在'), 404
    if request.method == 'DELETE':
        item['deleted_at'] = datetime.now(timezone.utc).isoformat()
        save_json(assets_path(cat, user_id), items)
        return jsonify(ok=True)
    body, err = _assets_request_body()
    if err:
        return err
    name = body.get('name')
    if name is not None:
        if not _valid_asset_name(name):
            return jsonify(error='名称不合法'), 400
        item['name'] = name.strip()
    url = body.get('url') or body.get('thumbnail_url')
    if url is not None:
        if not _valid_asset_url(url):
            return jsonify(error='URL 不合法'), 400
        if cat == 'styles':
            item['thumbnail_url'] = url
        else:
            item['url'] = url
    if cat == 'styles' and body.get('prompt') is not None:
        item['prompt'] = str(body.get('prompt'))[:5000]
    save_json(assets_path(cat, user_id), items)
    return jsonify(ok=True)


@app.route('/api/assets/<cat>/item/<item_id>/restore', methods=['POST'])
@login_required
def restore_asset_item(cat, item_id):
    user_id = current_user_id()
    if cat in ('character', 'characters'):
        chars = load_json(characters_path(user_id), {})
        if not isinstance(chars, dict) or item_id not in chars:
            return jsonify(error='资产不存在'), 404
        chars[item_id].pop('deleted_at', None)
        save_json(characters_path(user_id), chars)
        return jsonify(ok=True)
    if cat not in ASSET_CATS:
        return jsonify(error='不支持的分类'), 400
    items = load_json(assets_path(cat, user_id), [])
    if not isinstance(items, list):
        items = []
    item = next((x for x in items if isinstance(x, dict) and x.get('id') == item_id), None)
    if not item:
        return jsonify(error='资产不存在'), 404
    item.pop('deleted_at', None)
    save_json(assets_path(cat, user_id), items)
    return jsonify(ok=True)

# ── history ───────────────────────────────────────────────────
@app.route('/api/history', methods=['GET'])
@login_required
def get_history():
    history = load_json(history_path(), [])
    project_id = current_history_project_id()
    if project_id:
        history = [entry for entry in history if isinstance(entry, dict) and entry.get('project_id') in (None, '', project_id)]
    return jsonify(history)

@app.route('/api/history/media')
@login_required
def history_media_by_url():
    media_url = request.args.get('url', '')
    history = load_json(history_path(), [])
    item = next((entry for entry in history if entry.get('video_url') == media_url or entry.get('image_url') == media_url), None)
    if not item:
        return jsonify(error='媒体记录不存在'), 404
    try:
        upstream_headers = {'User-Agent': 'Mozilla/5.0'}
        if request.headers.get('Range'):
            upstream_headers['Range'] = request.headers['Range']
        upstream = requests.get(media_url, headers=upstream_headers, timeout=180, stream=True)
    except requests.RequestException as e:
        return jsonify(error=f'媒体读取失败：{e}'), 502
    if upstream.status_code not in (200, 206):
        upstream.close()
        return jsonify(error=f'媒体读取失败：HTTP {upstream.status_code}'), 502

    def stream_download():
        try:
            for chunk in upstream.iter_content(256 * 1024):
                if chunk:
                    yield chunk
        finally:
            upstream.close()

    ext = os.path.splitext(media_url.split('?', 1)[0])[1] or '.mp4'
    headers = {
        'Content-Type': upstream.headers.get('Content-Type', 'video/mp4'),
        'Accept-Ranges': upstream.headers.get('Accept-Ranges', 'bytes'),
        'Cache-Control': 'private, max-age=3600',
    }
    for key in ('Content-Length', 'Content-Range'):
        if upstream.headers.get(key):
            headers[key] = upstream.headers[key]
    if request.args.get('download') == '1':
        requested_name = secure_filename(request.args.get('name', ''))
        download_name = requested_name or f'generation{ext}'
        headers['Content-Disposition'] = f'attachment; filename="{download_name}"'
    return Response(stream_download(), status=upstream.status_code, headers=headers)

# ── models ────────────────────────────────────────────────────
@app.route('/api/models', methods=['GET'])
@login_required
def get_models():
    cfg = get_user_api_settings()
    has_nano = bool(cfg['nano_api_key'])
    has_third = bool(cfg['third_party_api_key'] and cfg['third_party_api_base'])
    return jsonify({
        'models': ALL_MODELS + ([THIRD_PARTY_MODEL_ID] if has_third and THIRD_PARTY_MODEL_ID not in ALL_MODELS else []),
        'nano_available': has_nano, 'agnes_available': bool(AGNES_API_KEY),
        'third_party_available': has_third, 'default': 'seedance', 'caps': MODEL_CAPS
    })

@app.route('/api/image-models', methods=['GET'])
@login_required
def get_image_models():
    return jsonify({
        'models': ALL_IMAGE_MODELS,
        'ratios': IMAGE_RATIOS,
        'agnes_available': bool(AGNES_API_KEY),
        'default_model': 'gpt-image-2',
        'default_ratio': DEFAULT_RATIO
    })

@app.route('/api/text-models', methods=['GET'])
@login_required
def get_text_models():
    return jsonify({
        'models': list(SCRIPT_MODELS),
        'default': SCRIPT_MODEL_DEFAULT
    })

@app.route('/api/text/generate', methods=['POST'])
@login_required
@model_access_required('text')
def generate_canvas_text():
    """Generic Canvas V2 text generation through the same provider resolver,
    permissions and point accounting as the classic workbench."""
    body = request.json or {}
    prompt = body.get('prompt')
    if not isinstance(prompt, str) or not prompt.strip():
        return jsonify(error='请输入文本任务'), 400
    if len(prompt) > 100000:
        return jsonify(error='文本任务过长'), 413
    script_model = body.get('script_model', SCRIPT_MODEL_DEFAULT)
    force_personal = bool(body.get('use_personal_api')) or script_model == 'personal-api'
    if not force_personal and script_model not in SCRIPT_MODELS:
        return jsonify(error='无效的文本模型'), 400
    point_cost = 0
    try:
        api_cfg = resolve_api(
            'text', builtin_text_api(SCRIPT_MODEL_DEFAULT if force_personal else script_model),
            current_user_id(), force_personal=force_personal,
            profile_id=body.get('api_profile_id'),
            strict_builtin='use_personal_api' in body and not force_personal,
        )
        point_cost = reserve_model_points(
            'text', script_model, current_user_id(), personal=force_personal
        )
        text = call_script_text_model(
            script_model,
            '你是 Manga Studio 的创作助手。准确完成用户给出的文本任务。直接输出结果，不要解释执行过程。',
            prompt.strip(),
            temperature=0.7,
            max_tokens=4000,
            api_cfg=api_cfg,
        )
        return jsonify(text=text, points=point_cost)
    except QuotaError as e:
        return jsonify(error=str(e)), 402
    except Exception as e:
        refund_model_points(current_user_id(), point_cost)
        return jsonify(error=f'生成失败: {str(e)}'), 500

# ── prompt refinement ──────────────────────────────────────────
def builtin_text_api(model_key=SCRIPT_MODEL_DEFAULT):
    model_cfg = SCRIPT_MODELS.get(model_key, SCRIPT_MODELS[SCRIPT_MODEL_DEFAULT])
    if model_cfg['provider'] == 'volcengine':
        return {'provider': 'ark', 'base_url': 'https://ark.cn-beijing.volces.com/api/v3', 'api_key': ARK_API_KEY, 'model': model_cfg['model_id']}
    return {'provider': 'nano', 'base_url': NANO_GPT_BASE, 'api_key': NANO_GPT_API_KEY, 'model': model_cfg['model_id']}

def call_platform_text(system_prompt, user_content, temperature=0.7, max_tokens=4000, cfg=None, user_id=None):
    cfg = cfg or resolve_api('text', builtin_text_api(), user_id)
    if not cfg.get('api_key'): raise Exception('文本 API 尚未配置')
    messages = [{'role': 'system', 'content': system_prompt}, {'role': 'user', 'content': user_content}]
    if cfg.get('provider') == 'ark':
        client = Ark(api_key=cfg['api_key'])
        resp = client.chat.completions.create(model=cfg['model'], messages=messages, temperature=temperature, max_tokens=max_tokens)
        return resp.choices[0].message.content.strip()
    headers = {'Content-Type': 'application/json', 'Authorization': f"Bearer {cfg['api_key']}", 'x-api-key': cfg['api_key']}
    if cfg.get('provider') == 'anthropic':
        headers['anthropic-version'] = '2023-06-01'
    payload = {'model': cfg['model'], 'messages': messages, 'temperature': temperature, 'max_tokens': max_tokens}
    endpoint = cfg.get('submit_url') or f"{cfg['base_url']}/chat/completions"
    r = requests.post(endpoint, headers=headers, json=payload, timeout=120)
    if r.status_code != 200: raise Exception(f'文本 API 调用失败: {r.status_code} {r.text[:300]}')
    return r.json()['choices'][0]['message']['content'].strip()

def load_prompt_style(style_id, user_id=None):
    if not style_id:
        return None
    styles, changed = ensure_default_styles(load_json(styles_path(user_id), []))
    if changed:
        save_json(styles_path(user_id), styles)
    return next((item for item in styles if item.get('id') == style_id), None)


def refine_prompt(script, images, ratio, duration, ark_api_key=None, user_id=None, style_id=None):
    """Use text model to optimize the user's script into a video-ready Chinese prompt."""
    role_desc = '、'.join([img.get('role_label', '参考图') for img in images]) if images else '无参考图'
    live_style = load_prompt_style(style_id, user_id) if style_id == LIVE_ACTION_STYLE_ID else None

    if live_style and LIVE_ACTION_GLOBAL_MARKER in script and '精准分镜时序脚本】' in script:
        return script

    live_action_rule = ''
    output_rule = '5. 输出纯中文，不要英文，不要 markdown，不超过 500 字'
    max_tokens = 1000
    if live_style:
        live_action_rule = (
            '5. 当前已选“真人短剧”风格。必须严格按以下四段输出：\n'
            f'{LIVE_ACTION_GLOBAL_MARKER}\n{live_style.get("prompt", "").strip()}\n'
            f'【负面词】\n{live_style.get("negative_prompt", "").strip()}\n'
            f'【{duration}秒精准分镜时序脚本】\n'
            '00:00-00:01｜1秒｜景别、机位/运镜、人物表情、视线、动作、情绪和画面重点。\n'
            f'按每1秒一条连续写到00:{int(duration):02d}，时间不得缺口、重叠或超出{duration}秒；每秒都必须是可拍摄、可视化的具体变化，不得只写“保持”。\n'
            '【连续性与质量约束】\n'
            '写明人物身份、服装、场景、光线的连续性，以及动作和运镜的自然衔接。\n'
            '6. 不要解释，不要 markdown，不要省略任何一秒。'
        )
        output_rule = ''
        max_tokens = 3000

    system_prompt = (
        '你是一个专业的漫剧分镜优化师。用户会提供一段分镜描述和参考素材信息，'
        '你需要将其优化为适合 AI 视频生成模型理解的中文 prompt。\n'
        '优化规则：\n'
        '1. 描述要具体、视觉化，包含画面构图、人物动作、镜头运动、光线氛围\n'
        '2. 必须提及画幅比例和时长信息\n'
        '3. 明确区分各参考图的用途（人物/服装/场景），不得混用\n'
        '4. 保留用户原始描述中的风格要求；原文没有风格要求时，不得自行添加真人拍摄、真实摄影、动漫、漫画、3D或其他视觉风格\n'
        + output_rule + ('\n' if output_rule and live_action_rule else '') + live_action_rule
    )

    user_prompt = (
        f'画面比例：{ratio}\n'
        f'视频时长：{duration} 秒\n'
        f'参考素材：{role_desc}\n'
        f'用户分镜描述：\n{script}\n\n'
        f'请输出优化后的分镜 prompt：'
    )

    try:
        refined = call_platform_text(system_prompt, user_prompt, temperature=0.7, max_tokens=max_tokens, user_id=user_id)
        if refined:
            if live_style and LIVE_ACTION_GLOBAL_MARKER not in refined:
                refined = (
                    f'{LIVE_ACTION_GLOBAL_MARKER}\n{live_style.get("prompt", "").strip()}\n'
                    f'【负面词】\n{live_style.get("negative_prompt", "").strip()}\n'
                    f'【{duration}秒精准分镜时序脚本】\n{refined}'
                )
            return refined
    except Exception as e:
        print(f'[refine_prompt] 优化失败，使用原始脚本: {e}')
    return script  # fallback to original


# ── Nano-GPT video generation adapter ─────────────────────────
def nano_gpt_generate(job_id, model_key, script, images, audio_url, video_url, first_frame_url, last_frame_url, ratio, duration, host_url,
                      resolution='720p', original_script=None, optimize=False, api_key=None, user_id=None, base_url=None):
    """Generate video via Nano-GPT API. Runs in a background thread."""
    try:
        if not api_key:
            JOBS[job_id] = {'status': 'failed', 'video_url': None, 'error': 'NANO_GPT_API_KEY 未设置，请配置环境变量'}
            return

        model_real = NANO_GPT_MODELS.get(model_key, model_key)
        api_root = (base_url or NANO_GPT_BASE).rstrip('/')
        if api_root.endswith('/v1'): api_root = api_root[:-3]
        headers = {
            'x-api-key': api_key,
            'Content-Type': 'application/json'
        }

        payload = {
            'model': model_real,
            'prompt': script,
            'aspect_ratio': ratio,
            'duration': str(duration),
            'resolution': resolution,
            'showExplicitContent': True,
            'wan27_has_video_input': bool(video_url),
            'wan27_has_reference_images': len(images) > 0,
            'voice': 'af_bella',
            'generateAudio': True,
            'camera_fixed': False,
        }
        model_info = nano_video_model_info({
            'provider': 'nano', 'base_url': base_url or NANO_GPT_BASE,
            'api_key': api_key, 'model': model_real
        })
        parameter_names = set((((model_info or {}).get('supported_parameters') or {}).get('parameters') or {}))
        def public_media_url(url):
            return host_url + url if url and url.startswith('/static/') else url

        if first_frame_url:
            payload['imageUrl'] = public_media_url(first_frame_url)
            payload['mode'] = 'image-to-video'
        if last_frame_url:
            last_field = next((name for name in ('last_image', 'lastImage', 'lastFrameUrl', 'last_frame_url') if name in parameter_names), None)
            if not last_field:
                raise ValueError(f'{model_real} 的 Nano API 没有尾帧参数，不能严格保持尾帧')
            payload[last_field] = public_media_url(last_frame_url)

        ref_images = []
        seen_urls = set()
        def append_ref(url, label):
            url = public_media_url(url)
            if url and url not in seen_urls:
                ref_images.append({'url': url, 'label': label})
                seen_urls.add(url)
        for img in images:
            append_ref(img.get('url', ''), img.get('role_label', 'reference'))
        if ref_images:
            reference_field = 'referenceImages' if 'referenceImages' in parameter_names else 'reference_images' if 'reference_images' in parameter_names else None
            if reference_field:
                payload[reference_field] = [img['url'] for img in ref_images]
        if video_url:
            payload['reference_video'] = video_url
            payload['video_url'] = video_url
        if audio_url:
            final_audio = host_url + audio_url if audio_url.startswith('/static/') else audio_url
            payload['reference_audio'] = final_audio

        JOBS[job_id]['status'] = 'running'
        print(f'[nano-start] model={model_key}', flush=True)

        # Submit generation task
        r = requests.post(f'{api_root}/generate-video', headers=headers, json=payload, timeout=30)
        print(f'[nano-submit] status={r.status_code} body={r.text[:200]}', flush=True)
        if r.status_code not in (200, 202):
            JOBS[job_id] = {'status': 'failed', 'video_url': None, 'error': f'Nano-GPT 提交失败: {r.status_code} {r.text[:200]}'}
            return

        data = r.json()
        task_id = data.get('runId') or data.get('id')
        if not task_id:
            JOBS[job_id] = {'status': 'failed', 'video_url': None, 'error': f'Nano-GPT 返回无 runId: {data}'}
            return

        # The submit response may include an eventsUrl, but Nano can return only
        # {"reason":"snapshot_complete"} there. Use the official status endpoint
        # as the authoritative result poller, with a legacy fallback.
        status_endpoints = [
            (f'{api_root}/video/status', {'requestId': task_id}),
            (f'{api_root}/generate-video/status', {'runId': task_id, 'model': model_real, 'modelSlug': model_real}),
        ]
        events_url_path = data.get('eventsUrl', '')
        if events_url_path and events_url_path.startswith('/'):
            events_url = f'https://nano-gpt.com{events_url_path}'
        else:
            events_url = None
        print(f'[nano-poll] status_endpoints={status_endpoints} events_url={events_url}', flush=True)
        for _ in range(240):  # 20 minutes max
            def _parse_poll_response(resp):
                try:
                    return resp.json()
                except json.JSONDecodeError:
                    pass

                # Nano may return text/event-stream from eventsUrl. Use the latest JSON data event.
                last_data = None
                for line in resp.text.splitlines():
                    line = line.strip()
                    if not line.startswith('data:'):
                        continue
                    data_text = line[5:].strip()
                    if not data_text or data_text == '[DONE]':
                        continue
                    last_data = data_text
                if last_data:
                    try:
                        return json.loads(last_data)
                    except json.JSONDecodeError:
                        return {'status': 'running', 'raw_event': last_data}
                return {'status': 'running', 'raw_text': resp.text[:1000]}

            pd = None
            last_http_error = None
            for status_url, params in status_endpoints:
                try:
                    pr = requests.get(status_url, headers=headers, params=params, timeout=30)
                except requests.RequestException as e:
                    print(f'[nano-poll] request error: {e}', flush=True)
                    last_http_error = {'request_error': str(e), 'url': status_url, 'params': params}
                    continue

                if pr.status_code in (200, 202):
                    pd = _parse_poll_response(pr)
                    if pd.get('reason') != 'snapshot_complete':
                        break
                else:
                    print(f'[nano-poll] http {pr.status_code} url={status_url} body={pr.text[:200]}', flush=True)
                    last_http_error = {'http_status': pr.status_code, 'url': status_url, 'params': params, 'body': pr.text[:500]}

            if pd is None:
                JOBS[job_id]['_last_poll'] = last_http_error or {'error': 'no poll response'}
                time.sleep(5)
                continue

            status_obj = pd.get('data') if isinstance(pd.get('data'), dict) else pd
            st = str(status_obj.get('status') or status_obj.get('state') or status_obj.get('stage') or '').lower()
            # Store raw poll response in job for frontend debugging
            JOBS[job_id]['_last_poll'] = {'status': st, 'keys': list(pd.keys()), 'raw': pd}
            sys.stdout.write(f'[nano-poll] status={st!r} keys={list(pd.keys())} raw={json.dumps(pd, ensure_ascii=False)[:500]}\n')
            sys.stdout.flush()

            # Try to find video URL anywhere in response (flat or nested)
            def _find_url(obj, depth=0):
                if depth > 8:
                    return None
                if isinstance(obj, str) and obj.startswith('http') and (
                    '.mp4' in obj.lower() or '.webm' in obj.lower() or '.mov' in obj.lower()
                ):
                    return obj
                if isinstance(obj, dict):
                    preferred_keys = (
                        'video_url', 'videoUrl', 'video', 'fileUrl', 'file_url', 'assetUrl', 'asset_url',
                        'mp4', 'output_url', 'outputUrl', 'result_url', 'resultUrl', 'download_url',
                        'downloadUrl', 'media_url', 'mediaUrl', 'url'
                    )
                    for key in preferred_keys:
                        if obj.get(key) and isinstance(obj[key], str) and obj[key].startswith('http'):
                            return obj[key]
                    for key, value in obj.items():
                        if isinstance(value, str) and value.startswith('http'):
                            key_l = str(key).lower()
                            if any(x in key_l for x in ('video', 'output', 'result', 'download', 'file', 'asset', 'media')):
                                return value
                    for v in obj.values():
                        found = _find_url(v, depth + 1)
                        if found:
                            return found
                if isinstance(obj, list):
                    for item in obj:
                        found = _find_url(item, depth + 1)
                        if found:
                            return found
                return None

            vurl = _find_url(pd)

            if st in ('completed', 'succeeded', 'done', 'success', 'complete', 'finished') or vurl:
                if vurl:
                    stored_vurl, _ = download_and_save_video(vurl)
                    save_video_history(
                        stored_vurl, script,
                        original_script=original_script or script,
                        refined_script=script if optimize else (original_script or script),
                        model=model_key, ratio=ratio, duration=duration,
                        resolution=resolution, ref_count=len(images), user_id=user_id)
                    JOBS[job_id] = {'status': 'succeeded', 'video_url': stored_vurl, 'source_video_url': vurl, 'error': None}
                    return
                else:
                    JOBS[job_id] = {'status': 'failed', 'video_url': None, 'error': f'Nano-GPT 成功但找不到视频URL，返回字段: {list(pd.keys())}，最后返回: {json.dumps(pd, ensure_ascii=False)[:500]}'}
                    return
            elif st in ('failed', 'error', 'cancelled', 'canceled'):
                err = status_obj.get('error') or status_obj.get('message') or pd.get('error') or pd.get('message') or '未知错误'
                JOBS[job_id] = {'status': 'failed', 'video_url': None, 'error': f'Nano-GPT 失败: {err}'}
                return
            time.sleep(5)

        JOBS[job_id] = {'status': 'failed', 'video_url': None, 'error': 'Nano-GPT 轮询超时（20分钟）'}

    except Exception as e:
        JOBS[job_id] = {'status': 'failed', 'video_url': None, 'error': f'Nano-GPT 异常: {str(e)}'}


# ── Third-party generic video adapter ─────────────────────────
def _find_media_url(obj, depth=0):
    if depth > 8:
        return None
    if isinstance(obj, str) and obj.startswith('http') and any(ext in obj.lower() for ext in ('.mp4', '.webm', '.mov')):
        return obj
    if isinstance(obj, dict):
        for key in ('video_url', 'videoUrl', 'output_url', 'outputUrl', 'download_url'):
            value = obj.get(key)
            if isinstance(value, str) and value.startswith('http'):
                return value
        value = obj.get('url')
        if isinstance(value, str) and value.startswith('http') and any(ext in value.lower() for ext in ('.mp4', '.webm', '.mov')):
            return value
        for value in obj.values():
            found = _find_media_url(value, depth + 1)
            if found:
                return found
    if isinstance(obj, list):
        for value in obj:
            found = _find_media_url(value, depth + 1)
            if found:
                return found
    return None


AGNES_VIDEO_SIZES = {
    '480p': {
        '16:9': (864, 480), '9:16': (480, 864), '1:1': (480, 480),
        '4:3': (640, 480), '3:4': (480, 640),
    },
    '720p': {
        '16:9': (1280, 720), '9:16': (720, 1280), '1:1': (720, 720),
        '4:3': (960, 720), '3:4': (720, 960),
    },
    '1080p': {
        '16:9': (1920, 1080), '9:16': (1080, 1920), '1:1': (1080, 1080),
        '4:3': (1440, 1080), '3:4': (1080, 1440),
    },
}

def agnes_num_frames(duration, frame_rate=24):
    """Return the closest Agnes-valid frame count (8n + 1, at most 441)."""
    target = max(1, int(duration)) * frame_rate
    frames = 8 * round((target - 1) / 8) + 1
    return min(441, max(9, frames))

def agnes_video_meta(payload, requested_ratio, requested_duration, requested_resolution):
    data = payload.get('data') if isinstance(payload, dict) and isinstance(payload.get('data'), dict) else payload
    data = data if isinstance(data, dict) else {}
    metadata = data.get('metadata') if isinstance(data.get('metadata'), dict) else {}
    mapping = metadata.get('size_mapping') if isinstance(metadata.get('size_mapping'), dict) else {}
    seconds = data.get('seconds') or metadata.get('seconds')
    try:
        actual_duration = round(float(seconds), 2)
    except (TypeError, ValueError):
        actual_duration = requested_duration
    size = data.get('size') or metadata.get('size') or mapping.get('size')
    if isinstance(size, dict):
        size = f"{size.get('width')}x{size.get('height')}" if size.get('width') and size.get('height') else None
    if not size and mapping.get('width') and mapping.get('height'):
        size = f"{mapping['width']}x{mapping['height']}"
    return {
        'ratio': mapping.get('ratio') or data.get('ratio') or requested_ratio,
        'duration': actual_duration,
        'resolution': mapping.get('resolution') or data.get('resolution') or requested_resolution,
        'size': size,
    }

def agnes_video_generate(job_id, script, ratio, duration, resolution='720p', original_script=None,
                         optimize=False, api_key=None, user_id=None, api_base=AGNES_API_BASE,
                         model_id=AGNES_VIDEO_MODEL_ID):
    """Create and poll an Agnes OpenAI-compatible video task."""
    try:
        if not api_key:
            raise Exception('AGNES_API_KEY 未设置')
        if resolution not in AGNES_VIDEO_SIZES:
            resolution = '720p'
        if ratio not in AGNES_VIDEO_SIZES[resolution]:
            ratio = '9:16'
        width, height = AGNES_VIDEO_SIZES[resolution][ratio]
        frame_rate = 24
        num_frames = agnes_num_frames(duration, frame_rate)
        headers = {'Authorization': f'Bearer {api_key}', 'Content-Type': 'application/json'}
        payload = {
            'model': model_id,
            'prompt': script,
            'width': width,
            'height': height,
            'num_frames': num_frames,
            'frame_rate': frame_rate,
        }
        api_base = api_base.rstrip('/')
        JOBS[job_id]['status'] = 'running'
        response = requests.post(f'{api_base}/videos', headers=headers, json=payload, timeout=60)
        if response.status_code not in (200, 201, 202):
            try:
                error_data = response.json()
                message = error_data.get('message') or error_data.get('error', {}).get('message') or response.text[:300]
            except (ValueError, AttributeError):
                message = response.text[:300]
            raise Exception(f'Agnes 视频提交失败: HTTP {response.status_code} {message}')

        created = response.json()
        created_data = created.get('data') if isinstance(created.get('data'), dict) else {}
        task_id = created.get('id') or created.get('task_id') or created.get('job_id') or created_data.get('id') or created_data.get('task_id')
        direct_url = _find_media_url(created)
        if direct_url:
            actual = agnes_video_meta(created, ratio, duration, resolution)
            stored_url, _ = download_and_save_video(direct_url)
            save_video_history(stored_url, script, original_script=original_script or script,
                               refined_script=script if optimize else (original_script or script),
                               model=model_id, ratio=actual['ratio'], duration=actual['duration'], resolution=actual['resolution'],
                               ref_count=0, user_id=user_id)
            JOBS[job_id] = {'status': 'succeeded', 'video_url': stored_url, 'source_video_url': direct_url,
                            'ratio': actual['ratio'], 'duration': actual['duration'],
                            'resolution': actual['resolution'], 'size': actual['size'], 'error': None}
            return
        if not task_id:
            raise Exception(f'Agnes 返回中没有任务 ID: {json.dumps(created, ensure_ascii=False)[:500]}')

        for _ in range(240):
            poll = requests.get(f'{api_base}/videos/{task_id}', headers=headers, timeout=30)
            if poll.status_code not in (200, 202):
                raise Exception(f'Agnes 视频查询失败: HTTP {poll.status_code} {poll.text[:300]}')
            result = poll.json()
            result_data = result.get('data') if isinstance(result.get('data'), dict) else result
            status = str(result_data.get('status') or result_data.get('state') or '').lower()
            video_url = _find_media_url(result)
            if video_url:
                actual = agnes_video_meta(result, ratio, duration, resolution)
                stored_url, _ = download_and_save_video(video_url)
                save_video_history(stored_url, script, original_script=original_script or script,
                                   refined_script=script if optimize else (original_script or script),
                                   model=model_id, ratio=actual['ratio'], duration=actual['duration'], resolution=actual['resolution'],
                                   ref_count=0, user_id=user_id)
                JOBS[job_id] = {'status': 'succeeded', 'video_url': stored_url, 'source_video_url': video_url,
                                'ratio': actual['ratio'], 'duration': actual['duration'],
                                'resolution': actual['resolution'], 'size': actual['size'], 'error': None}
                return
            if status in ('failed', 'error', 'cancelled', 'canceled'):
                message = result_data.get('error') or result_data.get('message') or '未知错误'
                raise Exception(f'Agnes 视频生成失败: {message}')
            time.sleep(5)
        raise Exception('Agnes 视频轮询超时（20分钟）')
    except Exception as e:
        JOBS[job_id] = {'status': 'failed', 'video_url': None, 'error': str(e)}


def minimax_video_generate(job_id, script, images, audio_url, video_url, first_frame_url, last_frame_url,
                           ratio, duration, resolution, host_url, model_id, api_key, api_base,
                           original_script=None, optimize=False, user_id=None):
    """Create and poll a MiniMax H3 V2 multimodal video task."""
    try:
        api_root = minimax_api_root(api_base)
        headers = {'Authorization': f'Bearer {api_key}', 'Content-Type': 'application/json'}
        content = [{'type': 'text', 'text': script}]
        seen_urls = set()

        def public_url(url):
            if url and url.startswith('/static/'):
                return host_url + url
            return url

        def append_image(url, role):
            url = public_url(url)
            if url and url not in seen_urls:
                content.append({'type': 'image_url', 'image_url': {'url': url}, 'role': role})
                seen_urls.add(url)

        append_image(first_frame_url, 'first_frame')
        append_image(last_frame_url, 'last_frame')
        for image in images or []:
            append_image(image.get('url'), 'reference_image')
        if video_url:
            content.append({'type': 'video_url', 'video_url': {'url': public_url(video_url)}, 'role': 'reference_video'})
        if audio_url:
            content.append({'type': 'audio_url', 'audio_url': {'url': public_url(audio_url)}, 'role': 'reference_audio'})

        minimax_resolution = {'768p': '768P', '768P': '768P', '2K': '2K', '2k': '2K'}.get(resolution)
        if not minimax_resolution:
            raise Exception(f'MiniMax H3 不支持分辨率 {resolution}，仅支持 768p 和 2K')
        minimax_ratio = ratio if ratio in ('21:9', '16:9', '4:3', '1:1', '3:4', '9:16') else '16:9'
        minimax_duration = max(4, min(15, int(duration)))
        payload = {
            'model': model_id or 'MiniMax-H3',
            'content': content,
            'resolution': minimax_resolution,
            'duration': minimax_duration,
            'ratio': minimax_ratio,
        }

        JOBS[job_id]['status'] = 'running'
        response = requests.post(f'{api_root}/v2/video_generation', headers=headers, json=payload, timeout=60)
        if response.status_code not in (200, 201, 202):
            raise Exception(f'MiniMax 视频提交失败: HTTP {response.status_code} {response.text[:300]}')
        created = response.json()
        task_id = created.get('task_id') or created.get('id')
        if not task_id:
            raise Exception(f'MiniMax 返回中没有 task_id: {json.dumps(created, ensure_ascii=False)[:400]}')

        for _ in range(240):
            poll = requests.get(f'{api_root}/v2/query/video_generation/{task_id}', headers=headers, timeout=30)
            if poll.status_code not in (200, 202):
                raise Exception(f'MiniMax 视频查询失败: HTTP {poll.status_code} {poll.text[:300]}')
            result = poll.json()
            task = result.get('task') if isinstance(result.get('task'), dict) else result
            status = str(task.get('status') or '').lower()
            output_url = _find_media_url(result)
            if status == 'succeeded' or output_url:
                if not output_url:
                    raise Exception('MiniMax 任务已完成但没有返回视频地址')
                stored_url, _ = download_and_save_video(output_url)
                save_video_history(
                    stored_url, script, original_script=original_script or script,
                    refined_script=script if optimize else (original_script or script),
                    model=model_id or 'MiniMax-H3', ratio=task.get('ratio') or minimax_ratio,
                    duration=task.get('duration') or minimax_duration,
                    resolution=task.get('resolution') or minimax_resolution,
                    ref_count=len(seen_urls), user_id=user_id
                )
                JOBS[job_id] = {
                    'status': 'succeeded', 'video_url': stored_url, 'source_video_url': output_url,
                    'ratio': task.get('ratio') or minimax_ratio,
                    'duration': task.get('duration') or minimax_duration,
                    'resolution': task.get('resolution') or minimax_resolution,
                    'error': None
                }
                return
            if status in ('failed', 'cancelled', 'canceled'):
                error = task.get('error') or result.get('error') or '未知错误'
                raise Exception(f'MiniMax 视频生成失败: {error}')
            time.sleep(5)
        raise Exception('MiniMax 视频轮询超时（20分钟）')
    except Exception as e:
        JOBS[job_id] = {'status': 'failed', 'video_url': None, 'error': str(e)}


def third_party_video_adapter(job_id, script, images, audio_url, video_url, first_frame_url, last_frame_url, ratio, duration, host_url,
                              model_key=THIRD_PARTY_MODEL_ID, resolution='720p',
                              original_script=None, optimize=False, api_base=None, api_key=None, user_id=None,
                              submit_url=None, status_url=None):
    """Generic adapter for any third-party video generation API.
    Reads api_base and api_key from env vars. Sends prompt + refs, polls for result."""
    try:
        api_base = api_base or THIRD_PARTY_API_BASE
        api_key = api_key or THIRD_PARTY_API_KEY
        if not api_base:
            JOBS[job_id] = {'status': 'failed', 'video_url': None, 'error': 'THIRD_PARTY_API_BASE 未设置'}
            return
        if not api_key:
            JOBS[job_id] = {'status': 'failed', 'video_url': None, 'error': 'THIRD_PARTY_API_KEY 未设置'}
            return

        headers = {
            'Authorization': f'Bearer {api_key}',
            'x-api-key': api_key,
            'Content-Type': 'application/json'
        }

        # Build reference images
        ref_images = []
        for img in images:
            url = img['url']
            if url.startswith('/static/'):
                url = host_url + url
            ref_images.append({'url': url, 'label': img.get('role_label', 'reference')})

        payload = {
            'model': model_key,
            'prompt': script,
            'aspect_ratio': ratio,
            'duration': duration,
            'reference_images': ref_images,
        }
        # Prefer the OpenAI-compatible /videos route; keep the legacy custom route as fallback.
        width, height = {'1:1': (768, 768), '9:16': (768, 1152), '16:9': (1152, 768)}.get(ratio, (768, 1152))
        openai_payload = {'model': model_key, 'prompt': script, 'width': width, 'height': height}
        if first_frame_url:
            first_frame_url = host_url + first_frame_url if first_frame_url.startswith('/static/') else first_frame_url
            payload['first_frame_url'] = first_frame_url
            openai_payload['first_frame_url'] = first_frame_url
        if last_frame_url:
            last_frame_url = host_url + last_frame_url if last_frame_url.startswith('/static/') else last_frame_url
            payload['last_frame_url'] = last_frame_url
            openai_payload['last_frame_url'] = last_frame_url
        if audio_url:
            final_audio = host_url + audio_url if audio_url.startswith('/static/') else audio_url
            payload['reference_audio'] = final_audio
        if video_url:
            payload['reference_video'] = video_url

        JOBS[job_id]['status'] = 'running'

        exact_submit_url = str(submit_url or '').strip()
        exact_status_url = str(status_url or '').strip()
        if exact_submit_url:
            request_payload = openai_payload if urlparse(exact_submit_url).path.rstrip('/').endswith('/videos') else payload
            r = requests.post(exact_submit_url, headers=headers, json=request_payload, timeout=30)
            resolved_status_url = exact_status_url
        else:
            r = requests.post(f'{api_base}/videos', headers=headers, json=openai_payload, timeout=30)
            resolved_status_url = None
        if not exact_submit_url and r.status_code in (404, 405):
            r = requests.post(f'{api_base}/video/generate', headers=headers, json=payload, timeout=30)
            resolved_status_url = f'{api_base}/video/status/{{task_id}}'
        if r.status_code not in (200, 201, 202):
            JOBS[job_id] = {'status': 'failed', 'video_url': None, 'error': f'第三方提交失败: {r.status_code} {r.text[:200]}'}
            return

        data = r.json()
        data_body = data.get('data') if isinstance(data.get('data'), dict) else {}
        direct_url = _find_media_url(data)
        task_id = data.get('task_id') or data.get('id') or data.get('job_id') or data_body.get('task_id') or data_body.get('id')
        if direct_url:
            stored_url, _ = download_and_save_video(direct_url)
            save_video_history(stored_url, script, original_script=original_script or script,
                               refined_script=script if optimize else (original_script or script), model=model_key,
                               ratio=ratio, duration=duration, resolution=resolution, ref_count=len(images), user_id=user_id)
            JOBS[job_id] = {'status': 'succeeded', 'video_url': stored_url, 'source_video_url': direct_url, 'error': None}
            return
        if not task_id:
            JOBS[job_id] = {'status': 'failed', 'video_url': None, 'error': f'第三方返回无 task_id: {data}'}
            return
        resolved_status_url = resolved_status_url or f'{api_base}/videos/{{task_id}}'

        # Poll
        for _ in range(240):
            pr = requests.get(resolved_status_url.format(task_id=task_id), headers=headers, timeout=30)
            if pr.status_code not in (200, 202):
                JOBS[job_id] = {'status': 'failed', 'video_url': None, 'error': f'第三方查询失败: {pr.status_code}'}
                return

            pd = pr.json()
            pd_body = pd.get('data') if isinstance(pd.get('data'), dict) else pd
            status = str(pd_body.get('status') or pd_body.get('state') or '').lower()
            vurl = _find_media_url(pd)
            if status in ('completed', 'succeeded', 'done', 'success', 'finished') or vurl:
                if not vurl:
                    JOBS[job_id] = {'status': 'failed', 'video_url': None, 'error': '第三方完成但无 video_url'}
                    return
                stored_url, _ = download_and_save_video(vurl)
                save_video_history(
                    stored_url, script,
                    original_script=original_script or script,
                    refined_script=script if optimize else (original_script or script),
                    model=model_key,
                    ratio=ratio,
                    duration=duration,
                    resolution=resolution,
                    ref_count=len(images),
                    user_id=user_id
                )
                JOBS[job_id] = {'status': 'succeeded', 'video_url': stored_url, 'source_video_url': vurl, 'error': None}
                return
            elif status in ('failed', 'error', 'cancelled'):
                err = pd.get('error') or pd.get('message', '未知错误')
                JOBS[job_id] = {'status': 'failed', 'video_url': None, 'error': f'第三方生成失败: {err}'}
                return
            else:
                time.sleep(5)
        JOBS[job_id] = {'status': 'failed', 'video_url': None, 'error': '第三方视频轮询超时（20分钟）'}

    except Exception as e:
        JOBS[job_id] = {'status': 'failed', 'video_url': None, 'error': f'第三方异常: {str(e)}'}


def atlas_output_url(payload):
    if isinstance(payload, str) and payload.startswith('http'):
        return payload
    if isinstance(payload, list):
        for item in payload:
            found = atlas_output_url(item)
            if found:
                return found
    if isinstance(payload, dict):
        for key in ('outputs', 'output', 'url', 'image_url', 'video_url'):
            found = atlas_output_url(payload.get(key))
            if found:
                return found
    return None


def atlas_generate_media(kind, payload, api_key, api_base=None):
    """Submit an Atlas Cloud image/video task and return its output URL."""
    api_root = ATLAS_API_BASE if 'api.atlascloud.ai' in (api_base or '').lower() else (api_base or ATLAS_API_BASE).rstrip('/')
    headers = {'Authorization': f'Bearer {api_key}', 'Content-Type': 'application/json'}
    endpoint = 'generateImage' if kind == 'image' else 'generateVideo'
    response = requests.post(f'{api_root}/model/{endpoint}', headers=headers, json=payload, timeout=60)
    if kind == 'image' and response.status_code in (404, 405):
        response = requests.post(f'{api_root}/model/generateVideo', headers=headers, json=payload, timeout=60)
    if response.status_code not in (200, 201, 202):
        raise Exception(f'Atlas Cloud 提交失败: HTTP {response.status_code} {response.text[:300]}')

    created = response.json()
    created_data = created.get('data') if isinstance(created, dict) and isinstance(created.get('data'), dict) else created
    created_data = created_data if isinstance(created_data, dict) else {}
    direct_url = atlas_output_url(created_data.get('outputs'))
    if direct_url:
        return direct_url
    task_id = created_data.get('id') or created_data.get('task_id')
    if not task_id:
        raise Exception(f'Atlas Cloud 返回中没有任务 ID: {json.dumps(created, ensure_ascii=False)[:400]}')

    interval = 2 if kind == 'image' else 5
    attempts = 180 if kind == 'image' else 240
    status_paths = ('prediction', 'result')
    for _ in range(attempts):
        result = None
        for status_path in status_paths:
            poll = requests.get(f'{api_root}/model/{status_path}/{task_id}', headers=headers, timeout=30)
            if poll.status_code in (404, 405):
                continue
            if poll.status_code not in (200, 202):
                raise Exception(f'Atlas Cloud 查询失败: HTTP {poll.status_code} {poll.text[:240]}')
            result = poll.json()
            break
        if result is None:
            raise Exception('Atlas Cloud 查询接口不可用')
        result_data = result.get('data') if isinstance(result, dict) and isinstance(result.get('data'), dict) else result
        result_data = result_data if isinstance(result_data, dict) else {}
        status = str(result_data.get('status') or result_data.get('state') or '').lower()
        output_url = atlas_output_url(result_data.get('outputs') or result_data.get('output'))
        if output_url:
            return output_url
        if status in ('failed', 'error', 'cancelled', 'canceled', 'timeout'):
            error = result_data.get('error') or result_data.get('message') or '未知错误'
            raise Exception(f'Atlas Cloud 生成失败: {error}')
        if status in ('completed', 'succeeded', 'done', 'success', 'finished'):
            raise Exception('Atlas Cloud 任务完成但没有返回媒体地址')
        time.sleep(interval)
    raise Exception(f'Atlas Cloud {"图片" if kind == "image" else "视频"}生成超时')


ATLAS_IMAGE_SIZES = {
    '1:1': '2048*2048', '4:3': '2304*1728', '3:4': '1728*2304',
    '16:9': '2720*1530', '9:16': '1530*2720',
    '3:2': '2496*1664', '2:3': '1664*2496',
    '5:4': '2304*1728', '4:5': '1728*2304',
}


def atlas_image_generate(prompt, model_id, ratio, custom_size, api_key, api_base, input_images, host_url):
    references = []
    for image in input_images or []:
        url = image.get('url', '')
        if url.startswith('/static/'):
            url = host_url + url
        if url:
            references.append(url)

    model = model_id
    if references and model.endswith('/text-to-image'):
        model = model.rsplit('/', 1)[0] + '/edit'
    elif not references and model.endswith('/edit'):
        model = model.rsplit('/', 1)[0] + '/text-to-image'

    size = custom_size.replace('x', '*').replace('X', '*') if ratio == 'custom' and custom_size else ATLAS_IMAGE_SIZES.get(ratio)
    payload = {'model': model, 'prompt': prompt, 'output_format': 'png'}
    if size:
        payload['size'] = size
    if references:
        payload['images'] = references[:10]
    output_url = atlas_generate_media('image', payload, api_key, api_base)
    return download_and_save_image(output_url)


def atlas_video_generate(job_id, script, images, audio_url, video_url, first_frame_url, last_frame_url, ratio, duration, resolution,
                         host_url, model_id, api_key, api_base, original_script, optimize, user_id):
    try:
        reference_images = []
        for image in images or []:
            url = image.get('url', '')
            if url.startswith('/static/'):
                url = host_url + url
            if url:
                reference_images.append(url)
        reference_videos = []
        if video_url:
            reference_videos.append(host_url + video_url if video_url.startswith('/static/') else video_url)
        reference_audios = []
        if audio_url:
            reference_audios.append(host_url + audio_url if audio_url.startswith('/static/') else audio_url)

        first_frame = host_url + first_frame_url if first_frame_url and first_frame_url.startswith('/static/') else first_frame_url
        last_frame = host_url + last_frame_url if last_frame_url and last_frame_url.startswith('/static/') else last_frame_url
        has_references = bool(reference_images or reference_videos or reference_audios)
        model = model_id
        if first_frame:
            model = atlas_image_variant(model) or model
        elif has_references and model.endswith('/text-to-video'):
            model = model.rsplit('/', 1)[0] + '/reference-to-video'
        elif not has_references and model.endswith('/reference-to-video'):
            model = model.rsplit('/', 1)[0] + '/text-to-video'

        properties = atlas_input_properties(model)
        allowed_resolutions = [str(value) for value in _schema_enum(properties, 'resolution')]
        allowed_ratios = [str(value) for value in _schema_enum(properties, 'ratio', 'aspect_ratio')]
        allowed_durations = [int(value) for value in _schema_enum(properties, 'duration') if str(value).lstrip('-').isdigit() and int(value) > 0]
        resolution_fallback = {
            '480p': '480p', '720p': '720p', '768p': '720p',
            '1080p': '1080p-SR', '1440p': '1440p-SR',
        }.get(resolution, '720p')
        atlas_resolution = resolution if resolution in allowed_resolutions else resolution_fallback
        if allowed_resolutions and atlas_resolution not in allowed_resolutions:
            raise Exception(f'{model} 不支持分辨率 {resolution}')
        atlas_ratio = ratio
        if allowed_ratios and atlas_ratio not in allowed_ratios:
            raise Exception(f'{model} 不支持比例 {ratio}')
        atlas_duration = int(duration)
        duration_def = properties.get('duration') or {}
        if allowed_durations and atlas_duration not in allowed_durations:
            raise Exception(f'{model} 不支持 {duration} 秒时长')
        if duration_def.get('minimum') is not None and atlas_duration < duration_def['minimum'] or duration_def.get('maximum') is not None and atlas_duration > duration_def['maximum']:
            raise Exception(f'{model} 不支持 {duration} 秒时长')
        payload = {
            'model': model,
            'prompt': script,
            'duration': atlas_duration,
            'resolution': atlas_resolution,
        }
        if 'generate_audio' in properties:
            payload['generate_audio'] = True
        if 'watermark' in properties:
            payload['watermark'] = False
        if 'aspect_ratio' in properties:
            payload['aspect_ratio'] = atlas_ratio
        else:
            payload['ratio'] = atlas_ratio
        if first_frame:
            first_field = next((name for name in ('image', 'image_url', 'imageUrl', 'first_frame_url', 'firstFrameUrl') if name in properties), None)
            if not first_field:
                raise Exception(f'{model} 的 Atlas API 没有首帧字段，已停止提交')
            payload[first_field] = first_frame
            if last_frame:
                last_field = next((name for name in ('last_image', 'lastImage', 'last_frame_url', 'lastFrameUrl') if name in properties), None)
                if not last_field:
                    raise Exception(f'{model} 的 Atlas API 没有尾帧字段，已停止提交')
                payload[last_field] = last_frame
        elif last_frame:
            raise Exception('Atlas 严格尾帧模式必须同时提供首帧图')
        elif reference_images:
            payload['reference_images'] = reference_images[:9]
        if reference_videos:
            payload['reference_videos'] = reference_videos[:3]
        if reference_audios and (reference_images or reference_videos):
            payload['reference_audios'] = reference_audios[:3]

        JOBS[job_id]['status'] = 'running'
        output_url = atlas_generate_media('video', payload, api_key, api_base)
        stored_url, _ = download_and_save_video(output_url)
        save_video_history(
            stored_url, script, original_script=original_script or script,
            refined_script=script if optimize else (original_script or script),
            model=model, ratio=atlas_ratio, duration=atlas_duration, resolution=atlas_resolution,
            ref_count=len(reference_images) + int(bool(first_frame)) + int(bool(last_frame)), user_id=user_id
        )
        JOBS[job_id] = {
            'status': 'succeeded', 'video_url': stored_url, 'source_video_url': output_url,
            'ratio': atlas_ratio, 'duration': atlas_duration, 'resolution': atlas_resolution, 'error': None
        }
    except Exception as e:
        print(f'[atlas-video] failed: {e}', flush=True)
        JOBS[job_id] = {'status': 'failed', 'video_url': None, 'error': str(e)}


# ── Image generation helpers ──────────────────────────────────
def download_and_save_image(image_url):
    """Download image from URL and keep it in durable storage when configured."""
    attempts = 3
    response = None
    retryable_statuses = {429, 500, 502, 503, 504}
    for attempt in range(attempts):
        try:
            candidate = requests.get(
                image_url,
                headers={'User-Agent': 'Mozilla/5.0', 'Accept': 'image/*,*/*;q=0.8'},
                timeout=(20, 180),
            )
            if candidate.status_code == 200:
                response = candidate
                break
            if candidate.status_code not in retryable_statuses:
                raise Exception(f'下载图片失败: {candidate.status_code}')
            failure = f'HTTP {candidate.status_code}'
        except (requests.exceptions.Timeout, requests.exceptions.ConnectionError) as exc:
            failure = type(exc).__name__
        if attempt + 1 < attempts:
            print(f'[image-cache] download retry {attempt + 1}/{attempts}: {failure}', flush=True)
            time.sleep(2 ** attempt)

    if response is None:
        raise Exception('图片已生成，但从供应商临时存储下载失败；系统已自动重试 3 次，请稍后重试')

    ct = response.headers.get('Content-Type', 'image/png')
    ext = '.png'
    if 'jpeg' in ct or 'jpg' in ct:
        ext = '.jpg'
    elif 'webp' in ct:
        ext = '.webp'
    name = uuid.uuid4().hex + ext
    img_bytes = response.content
    if not img_bytes:
        raise Exception('图片已生成，但供应商返回了空图片文件')

    # Try durable object storage
    public_url, ok = upload_to_tos(img_bytes, name, ct)
    if ok:
        return public_url, name

    if persistent_storage_configured():
        raise Exception('图片生成成功，但文件存储失败，请稍后重试')

    # Fallback to local
    path = os.path.join(UPLOAD, name)
    with open(path, 'wb') as f:
        f.write(img_bytes)
    return f'/static/uploads/{name}', name


def download_and_save_video(video_url):
    """Download a remote generated video and persist it to configured storage."""
    if not video_url or not video_url.startswith('http'):
        return video_url, None
    if is_persistent_storage_url(video_url):
        return video_url, os.path.basename(video_url.split('?', 1)[0])

    headers = {'User-Agent': 'Mozilla/5.0'}
    try:
        r = requests.get(video_url, headers=headers, timeout=180, stream=True)
        if r.status_code != 200:
            print(f'[video-cache] download failed: {r.status_code} {r.text[:120]}', flush=True)
            return video_url, None

        ct = r.headers.get('Content-Type', 'video/mp4').split(';', 1)[0] or 'video/mp4'
        ext = '.mp4'
        if 'webm' in ct:
            ext = '.webm'
        elif 'quicktime' in ct or 'mov' in ct:
            ext = '.mov'
        name = f"{uuid.uuid4().hex}{ext}"
        content_length = r.headers.get('Content-Length')
        content_length = int(content_length) if content_length and content_length.isdigit() else None

        # urllib3 response bodies are not seekable. boto3 may need to rewind a
        # body while calculating checksums/retrying, so download into a seekable
        # file first. Iterate explicitly: some response.raw wrappers can expose
        # an empty body even though iter_content() yields the media correctly.
        with tempfile.TemporaryFile(mode='w+b') as media_file:
            for chunk in r.iter_content(1024 * 1024):
                if chunk:
                    media_file.write(chunk)
            actual_length = media_file.tell()
            if actual_length <= 0:
                raise Exception('供应商返回了空视频文件（0 字节）')
            media_file.seek(0)
            public_url, ok = upload_to_tos(media_file, name, ct, content_length=actual_length)
        if ok:
            print(f'[video-cache] cached to {storage_name_for_url(public_url)}: {public_url}', flush=True)
            return public_url, name
    except Exception as e:
        print(f'[video-cache] failed: {e}', flush=True)

    if persistent_storage_configured():
        print('[video-cache] durable storage unavailable; preserving provider result URL', flush=True)
    return video_url, None

def migrate_tos_history_videos_to_r2():
    """Copy existing TOS history videos to R2 without deleting old objects."""
    if not r2_configured():
        return
    users = load_json(users_path(), {})
    migrated = failed = 0
    for user_id in users:
        snapshot = load_json(history_path(user_id), [])
        tos_urls = list(dict.fromkeys(
            item.get('video_url') for item in snapshot
            if item.get('type') == 'video' and str(item.get('video_url') or '').startswith(TOS_PUBLIC_BASE)
        ))
        for old_url in tos_urls:
            try:
                probe = requests.head(old_url, headers={'User-Agent': 'Mozilla/5.0'}, timeout=30, allow_redirects=True)
                if probe.status_code >= 400 or int(probe.headers.get('Content-Length') or 0) <= 0:
                    raise Exception('源视频为空或不可读取，保留原历史地址')
                new_url, _ = download_and_save_video(old_url)
                if not new_url or not new_url.startswith(R2_PUBLIC_BASE):
                    raise Exception('R2 未返回公开地址')
                with HISTORY_LOCK:
                    current = load_json(history_path(user_id), [])
                    changed = False
                    for item in current:
                        if item.get('type') == 'video' and item.get('video_url') == old_url:
                            item['video_url'] = new_url
                            changed = True
                    if changed:
                        save_json(history_path(user_id), current)
                        migrated += 1
            except Exception as e:
                failed += 1
                print(f'[R2-migrate] skipped {old_url}: {e}', flush=True)
    print(f'[R2-migrate] history videos migrated={migrated} failed={failed}', flush=True)


def build_image_content(prompt, input_images, host_url):
    """Build Ark image generation content list."""
    content = [{'type': 'text', 'text': prompt}]
    for img in (input_images or []):
        url = img['url']
        if url.startswith('/static/'):
            url = host_url + url
        content.append({'type': 'image_url', 'image_url': {'url': url}})
    return content


def parse_image_size(size):
    try:
        w, h = size.lower().split('x', 1)
        return int(w), int(h)
    except Exception:
        return None, None


def image_size_for_nano(model_id, ratio, custom_size=''):
    if ratio == 'custom' and custom_size:
        return custom_size
    if model_id == AGNES_IMAGE_MODEL_ID:
        return RATIO_TO_SIZE_AGNES.get(ratio, "1024x1024")
    if model_id == 'gpt-image-2':
        return RATIO_TO_SIZE_GPT_IMAGE.get(ratio, "1024x1024")
    return RATIO_TO_SIZE_NANO.get(ratio, "1024x1024")


def image_ratio_instruction(ratio, size):
    if ratio == 'custom' and size:
        return f"【画幅强制】最终图片必须严格输出为 {size} 像素，不要输出 1:1 正方形。"
    if ratio and ratio != "1:1":
        return f"【画幅强制】最终图片必须严格输出为 {ratio} 画幅，不要输出 1:1 正方形，不要自动裁成方图。"
    return "【画幅强制】最终图片使用 1:1 正方形画幅。"


# ── Nano image generation ─────────────────────────────────────
def nano_image_generate(prompt, model_id, ratio, custom_size='', api_key=None, base_url=None,
                        input_images=None, host_url='', submit_url=None):
    """Call Nano-GPT images/generations API."""
    size = image_size_for_nano(model_id, ratio, custom_size)
    width, height = parse_image_size(size)
    final_prompt = image_ratio_instruction(ratio, size) + "\n" + prompt
    headers = {
        'Content-Type': 'application/json',
        'x-api-key': api_key or NANO_GPT_API_KEY,
        'Authorization': f'Bearer {api_key or NANO_GPT_API_KEY}'
    }
    api_root = (base_url or NANO_GPT_BASE).rstrip('/')
    is_nano_midjourney = (
        model_id == NANO_GPT_MIDJOURNEY_MODEL_ID
        and 'nano-gpt.com' in api_root.lower()
    )
    if is_nano_midjourney:
        if input_images:
            raise Exception('Nano-GPT Midjourney 目前仅支持文生图，请移除参考图后重试')
        supported_ratio = ratio if ratio in {'1:1', '16:9', '9:16', '21:9', '9:21', '4:3', '3:4', '3:2', '2:3'} else '1:1'
        payload = {
            'model': model_id,
            'prompt': final_prompt,
            'resolution': supported_ratio,
            'aspect_ratio': supported_ratio,
            'version': NANO_GPT_MIDJOURNEY_VERSION,
            'n': 4,
        }
        endpoint = submit_url or f'{api_root}/images'
    else:
        payload = {
            'model': model_id,
            'prompt': final_prompt,
            'size': size,
            'dimensions': size
        }
        if ratio != 'custom':
            payload['aspect_ratio'] = ratio
        if width and height:
            payload['width'] = width
            payload['height'] = height
        endpoint = submit_url or f'{api_root}/images/generations'
    reference_urls = []
    for img in (input_images or []):
        url = img.get('url', '')
        if url.startswith('/static/'):
            url = host_url + url
        if url:
            reference_urls.append(url)
    if reference_urls:
        payload['image'] = reference_urls[0] if len(reference_urls) == 1 else reference_urls
    r = requests.post(endpoint, headers=headers, json=payload, timeout=360)
    if r.status_code not in (200, 201):
        raise Exception(f'图片 API 生成失败: {r.status_code} {r.text[:200]}')
    data = r.json()
    # Handle url or base64 response
    items = data.get('data', [])
    if isinstance(items, list) and len(items) > 0:
        item = items[0]
        if item.get('url'):
            return download_and_save_image(item['url'])
        if item.get('b64_json'):
            import base64
            img_bytes = base64.b64decode(item['b64_json'])
            name = uuid.uuid4().hex + '.png'
            public_url, ok = upload_to_tos(img_bytes, name, 'image/png')
            if ok:
                return public_url, name
            if persistent_storage_configured():
                raise Exception('图片生成成功，但文件存储失败，请稍后重试')
            path = os.path.join(UPLOAD, name)
            with open(path, 'wb') as f:
                f.write(img_bytes)
            return f'/static/uploads/{name}', name
    raise Exception(f'图片 API 返回无图片: {str(data)[:200]}')


# ── Volc Seedream image generation ────────────────────────────
def volc_image_generate(prompt, input_images, host_url, ratio, custom_size='', api_key=None, model_id=None):
    """Call Volc Ark Seedream for image generation."""
    if ratio == 'custom' and custom_size:
        size = custom_size
    elif model_id == VOLC_SEEDREAM_5_PRO_MODEL_ID:
        size = RATIO_TO_SIZE_VOLC_SEEDREAM_5_PRO.get(ratio, "2048x2048")
    else:
        size = RATIO_TO_SIZE_VOLC.get(ratio, "1920x1920")
    client = Ark(api_key=api_key or ARK_API_KEY)
    ref_urls = []
    for img in (input_images or []):
        url = img['url']
        if url.startswith('/static/'):
            url = host_url + url
        ref_urls.append(url)
    kwargs = {'model': model_id or VOLC_IMAGE_MODEL_ID, 'prompt': prompt, 'size': size, 'watermark': False}
    if ref_urls:
        kwargs['image'] = ref_urls[0] if len(ref_urls) == 1 else ref_urls
    resp = client.images.generate(**kwargs)
    img_url = None
    if hasattr(resp, 'data') and resp.data:
        img_url = resp.data[0].url
    if not img_url:
        raise Exception('Seedream 返回无图片 URL')
    return download_and_save_image(img_url)


# ── Generate image endpoint ───────────────────────────────────
@app.route('/api/generate-image', methods=['POST'])
@login_required
@model_access_required('image')
def generate_image():
    body = request.json or {}
    prompt = body.get('prompt', '').strip()
    selected_model = LEGACY_IMAGE_MODEL_ALIASES.get(
        body.get('image_model', 'gpt-image-2'),
        body.get('image_model', 'gpt-image-2')
    )
    force_personal = bool(body.get('use_personal_api')) or selected_model == 'personal-api'
    if not force_personal and selected_model not in ALL_IMAGE_MODELS:
        return jsonify(error='无效的图片模型'), 400
    if selected_model == AGNES_IMAGE_MODEL_ID:
        builtin = {'provider': 'agnes', 'base_url': AGNES_API_BASE, 'api_key': AGNES_API_KEY, 'model': AGNES_IMAGE_MODEL_ID}
    elif selected_model in NANO_GPT_IMAGE_MODELS:
        builtin = {'provider': 'nano', 'base_url': NANO_GPT_BASE, 'api_key': NANO_GPT_API_KEY, 'model': selected_model}
    else:
        builtin = {
            'provider': 'ark',
            'base_url': 'https://ark.cn-beijing.volces.com/api/v3',
            'api_key': ARK_API_KEY,
            'model': VOLC_IMAGE_MODELS[selected_model],
        }
    ratio = body.get('ratio', DEFAULT_RATIO)
    custom_size = body.get('custom_size', '')
    mode = body.get('mode', 'storyboard')
    input_images = body.get('input_images') or []
    style_id = body.get('style_id')
    host_url = request.host_url.rstrip('/')
    user_id = current_user_id()

    if not prompt:
        return jsonify(error='prompt 不能为空'), 400
    try:
        image_cfg = resolve_api(
            'image', builtin, user_id, force_personal=force_personal,
            profile_id=body.get('api_profile_id'),
            strict_builtin='use_personal_api' in body and not force_personal
        )
    except QuotaError as e:
        return jsonify(error=str(e)), 402
    image_model = image_cfg['model']

    # Inject style
    if style_id:
        style = load_prompt_style(style_id, user_id)
        if style:
            prompt = style.get('prompt', '') + '\n' + prompt
            if style.get('negative_prompt'):
                prompt += '\n\n避免：' + style.get('negative_prompt', '')

    try:
        point_cost = reserve_model_points('image', selected_model, user_id, personal=force_personal)
    except QuotaError as e:
        return jsonify(error=str(e)), 402

    job_id = uuid.uuid4().hex
    JOB_OWNERS[job_id] = user_id
    JOBS[job_id] = {'status': 'pending', 'url': None, 'name': None, 'error': None,
                     'model': image_model, 'ratio': ratio, 'mode': mode}

    if os.environ.get('MANGA_MOCK_GENERATION') == '1':
        _mock_generation_job(job_id, 'image')
        return jsonify(job_id=job_id)

    def run():
        try:
            if image_cfg.get('provider') == 'atlas':
                local_url, filename = atlas_image_generate(
                    prompt, image_model, ratio, custom_size,
                    image_cfg['api_key'], image_cfg['base_url'], input_images, host_url
                )
            elif image_cfg.get('provider') != 'ark':
                local_url, filename = nano_image_generate(
                    prompt, image_model, ratio, custom_size,
                    image_cfg['api_key'], image_cfg['base_url'], input_images, host_url,
                    submit_url=image_cfg.get('submit_url')
                )
            elif image_cfg.get('provider') == 'ark':
                local_url, filename = volc_image_generate(prompt, input_images, host_url, ratio, custom_size, image_cfg['api_key'], image_model)
            else:
                JOBS[job_id] = {'status': 'failed', 'url': None, 'error': f'不支持的图片模型: {image_model}'}
                return

            insert_history({
                'time': datetime.now().strftime('%Y-%m-%d %H:%M'),
                'type': 'image',
                'image_url': local_url,
                'script': prompt[:80],
                'original_script': prompt,
                'model': image_model,
                'ratio': ratio
            }, user_id=user_id)
            JOBS[job_id] = {'status': 'succeeded', 'url': local_url, 'name': filename,
                             'model': image_model, 'ratio': ratio, 'mode': mode}
        except Exception as e:
            JOBS[job_id] = {'status': 'failed', 'url': None, 'error': str(e)}

    start_metered_job(run, (), job_id, user_id, point_cost)
    return jsonify(job_id=job_id)


# ── Image job status (polling) ────────────────────────────────
@app.route('/api/image-status/<job_id>', methods=['GET'])
@login_required
def image_status(job_id):
    if JOB_OWNERS.get(job_id) != current_user_id():
        return jsonify(status='not_found'), 404
    return jsonify(JOBS.get(job_id, {'status': 'not_found'}))


# ── generate ──────────────────────────────────────────────────
@app.route('/api/generate', methods=['POST'])
@login_required
@model_access_required('video')
def generate():
    body      = request.json or {}
    script    = body.get('script', '')
    images    = body.get('images', [])   # [{url, role_label}]
    audio_url = body.get('audio_url')
    video_url = body.get('video_url')
    first_frame_url = body.get('first_frame_url')
    last_frame_url  = body.get('last_frame_url')
    storyboard_ref_url = body.get('storyboard_ref_url')
    style_id   = body.get('style_id')

    ratio       = body.get('ratio', '9:16')
    duration    = int(body.get('duration', 5))
    resolution  = body.get('resolution', '720p')
    optimize    = body.get('optimize_prompt', True)
    selected_model = body.get('video_model', 'seedance')
    force_personal = bool(body.get('use_personal_api')) or selected_model == 'personal-api'
    if not force_personal and selected_model not in ALL_MODELS:
        return jsonify(error='无效的视频模型'), 400
    if selected_model == AGNES_VIDEO_MODEL_ID:
        builtin = {'provider': 'agnes', 'base_url': AGNES_API_BASE, 'api_key': AGNES_API_KEY, 'model': AGNES_VIDEO_MODEL_ID}
    elif selected_model in NANO_GPT_NAMES:
        builtin = {'provider': 'nano', 'base_url': NANO_GPT_BASE, 'api_key': NANO_GPT_API_KEY, 'model': selected_model}
    elif selected_model == THIRD_PARTY_MODEL_ID:
        builtin = {'provider': 'generic', 'base_url': THIRD_PARTY_API_BASE.rstrip('/'), 'api_key': THIRD_PARTY_API_KEY, 'model': selected_model}
    else:
        builtin = {'provider': 'ark', 'base_url': 'https://ark.cn-beijing.volces.com/api/v3', 'api_key': ARK_API_KEY, 'model': MODEL_ID}
    host_url    = request.host_url.rstrip('/')
    job_id      = uuid.uuid4().hex
    user_id     = current_user_id()
    try:
        video_cfg = resolve_api(
            'video', builtin, user_id, force_personal=force_personal,
            profile_id=body.get('api_profile_id'),
            strict_builtin='use_personal_api' in body and not force_personal
        )
    except QuotaError as e:
        return jsonify(error=str(e)), 402
    video_model = video_cfg['model']
    video_caps = MODEL_CAPS.get(video_model) or MODEL_CAPS.get(selected_model)
    if force_personal and video_cfg.get('provider') == 'nano':
        discovered_caps = nano_video_model_capabilities(video_cfg)
        if not discovered_caps and (first_frame_url or last_frame_url):
            return jsonify(error=f'无法读取 {video_model} 的 Nano 视频能力，暂不能安全提交首尾帧'), 400
        if discovered_caps:
            video_caps = discovered_caps
    elif force_personal and video_cfg.get('provider') == 'atlas':
        discovered_caps = atlas_video_model_capabilities(video_cfg)
        if not discovered_caps and (first_frame_url or last_frame_url):
            return jsonify(error=f'无法读取 {video_model} 的 Atlas 视频能力，暂不能安全提交首尾帧'), 400
        if discovered_caps:
            video_caps = discovered_caps
    elif force_personal and video_cfg.get('provider') == 'ark':
        video_caps = MODEL_CAPS.get('seedance')
    if first_frame_url and video_caps and not video_caps.get('supports_first_frame', False):
        return jsonify(error=f'{video_model} 的 API 不支持首帧输入，请更换支持图生视频的模型'), 400
    if last_frame_url and video_caps and not video_caps.get('supports_last_frame', False):
        return jsonify(error=f'{video_model} 的 API 没有尾帧参数，不能严格保持尾帧，请更换支持尾帧的模型'), 400
    if video_caps:
        supported_ratios = video_caps.get('ratios') or []
        supported_resolutions = video_caps.get('resolutions') or []
        if supported_ratios and ratio not in supported_ratios:
            return jsonify(error=f'{video_model} 不支持比例 {ratio}'), 400
        if supported_resolutions and resolution not in supported_resolutions:
            return jsonify(error=f'{video_model} 不支持分辨率 {resolution}，请选择 {" / ".join(supported_resolutions)}'), 400
        min_duration = video_caps.get('min_duration')
        max_duration = video_caps.get('max_duration')
        if min_duration is not None and duration < min_duration or max_duration is not None and duration > max_duration:
            return jsonify(error=f'{video_model} 时长仅支持 {min_duration}-{max_duration} 秒'), 400
        if not video_caps.get('supports_reference_images', True):
            images = []
            storyboard_ref_url = None
        if not video_caps.get('supports_reference_audio', True):
            audio_url = None
        if not video_caps.get('supports_reference_video', True):
            video_url = None
    if (selected_model == 'seedance' or video_cfg.get('provider') == 'ark') and last_frame_url and not first_frame_url:
        return jsonify(error='Seedance 严格尾帧模式必须同时提供首帧图'), 400

    reference_images = list(images)
    if storyboard_ref_url:
        reference_images.append({'url': storyboard_ref_url, 'role_label': '分镜构图参考'})
    # Inject style
    if style_id:
        style = load_prompt_style(style_id, user_id)
        if style:
            # A storyboard prompt may already contain the complete live-action block.
            if LIVE_ACTION_GLOBAL_MARKER not in script:
                if style.get('prompt'):
                    prefix = LIVE_ACTION_GLOBAL_MARKER + '\n' if style_id == LIVE_ACTION_STYLE_ID else ''
                    script = prefix + style.get('prompt', '') + '\n' + script
                if style.get('negative_prompt'):
                    script += '\n【风格约束】避免：' + style.get('negative_prompt', '')

    original_script = script
    if optimize and script.strip():
        script = refine_prompt(script, images, ratio, duration, user_id=user_id, style_id=style_id)

    try:
        point_cost = reserve_model_points('video', selected_model, user_id, quantity=duration, personal=force_personal)
    except QuotaError as e:
        return jsonify(error=str(e)), 402
    JOB_OWNERS[job_id] = user_id
    JOBS[job_id] = {'status': 'pending', 'video_url': None, 'error': None}

    if os.environ.get('MANGA_MOCK_GENERATION') == '1':
        _mock_generation_job(job_id, 'video')
        return jsonify(job_id=job_id)

    # ── Atlas Cloud path ──
    if video_cfg.get('provider') == 'atlas':
        start_metered_job(atlas_video_generate, (
            job_id, script, reference_images, audio_url, video_url, first_frame_url, last_frame_url,
            ratio, duration, resolution,
            host_url, video_model, video_cfg['api_key'], video_cfg['base_url'],
            original_script, optimize, user_id
        ), job_id, user_id, point_cost)
        return jsonify(job_id=job_id)

    # ── Agnes path ──
    if video_cfg.get('provider') == 'agnes':
        start_metered_job(agnes_video_generate, (
            job_id, script, ratio, duration, resolution, original_script, optimize,
            video_cfg['api_key'], user_id, video_cfg['base_url'], video_model
        ), job_id, user_id, point_cost)
        return jsonify(job_id=job_id)

    # ── Nano-GPT path ──
    if video_cfg.get('provider') == 'nano':
        if not video_cfg.get('api_key'):
            JOBS[job_id] = {'status': 'failed', 'video_url': None,
                            'error': 'NANO_GPT_API_KEY 未设置。请 export NANO_GPT_API_KEY=sk-nano-xxx 后重启服务'}
            refund_model_points(user_id, point_cost)
            return jsonify(job_id=job_id)
        start_metered_job(nano_gpt_generate, (
            job_id, video_model, script, reference_images, audio_url, video_url, first_frame_url, last_frame_url, ratio, duration, host_url,
            resolution, original_script, optimize, video_cfg['api_key'], user_id, video_cfg['base_url']
        ), job_id, user_id, point_cost)
        return jsonify(job_id=job_id)

    # ── MiniMax H3 path ──
    if video_cfg.get('provider') == 'minimax':
        start_metered_job(minimax_video_generate, (
            job_id, script, reference_images, audio_url, video_url, first_frame_url, last_frame_url,
            ratio, duration, resolution, host_url, video_model, video_cfg['api_key'],
            video_cfg['base_url'], original_script, optimize, user_id
        ), job_id, user_id, point_cost)
        return jsonify(job_id=job_id)

    # ── Third-party path ──
    if video_cfg.get('provider') not in ('nano', 'agnes', 'ark'):
        if not video_cfg.get('api_key') or not video_cfg.get('base_url'):
            JOBS[job_id] = {'status': 'failed', 'video_url': None,
                            'error': '第三方模型未配置。请设置 THIRD_PARTY_API_BASE 和 THIRD_PARTY_API_KEY 环境变量后重启服务'}
            refund_model_points(user_id, point_cost)
            return jsonify(job_id=job_id)
        start_metered_job(third_party_video_adapter, (
            job_id, script, reference_images, audio_url, video_url, first_frame_url, last_frame_url, ratio, duration, host_url,
            video_model, resolution, original_script, optimize,
            video_cfg['base_url'], video_cfg['api_key'], user_id,
            video_cfg.get('submit_url'), video_cfg.get('status_url')
        ), job_id, user_id, point_cost)
        return jsonify(job_id=job_id)

    # ── Ark Seedance path (default) ──
    def run():
        try:
            # Ark documents strict first/last-frame generation as a dedicated
            # image-to-video mode. Mixing reference_image/audio/video items into
            # the same request changes it into all-modal reference generation,
            # where the boundary frames are guidance instead of hard anchors.
            strict_frame_mode = bool(first_frame_url)
            ark_reference_images = [] if strict_frame_mode else reference_images
            # build role description
            lines = ['【严格参考说明，必须遵守】']
            if strict_frame_mode:
                lines.append('- 首帧图必须作为视频 00:00 的实际起始画面，不得重绘、替换人物或改变构图。')
                if last_frame_url:
                    lines.append('- 尾帧图必须作为视频结束时的实际画面，所有动作和运镜必须自然收束到该画面。')
                    lines.append('- 只生成首帧到尾帧之间的连续过渡，不得生成与首尾帧无关的新人物、新场景或新构图。')
                else:
                    lines.append('- 从首帧图自然延续动作和镜头，不得另起画面。')
            for i, img in enumerate(ark_reference_images, 1):
                lines.append(f'- 图{i+1}：{img["role_label"]}，仅参考此用途，不得混用')
            if ark_reference_images:
                lines.append('请严格区分各参考图用途，不得将穿搭图用于脸部，不得将场景图用于人物。')
            # Storyboard reference instructions
            if storyboard_ref_url:
                lines.append('\n【分镜参考使用规则】')
                lines.append('已提供一张分镜构图参考图。该图只用于指导镜头设计，不用于角色身份识别。')
                lines.append('必须严格参考：镜头景别、机位角度、人物站位、动作关系、构图重心、场景空间关系、前景/中景/背景层次、光线氛围、运镜方向和镜头节奏。')
                lines.append('必须忽略：分镜参考图中的具体脸部细节、五官长相、人物身份、可能出现的文字标记、不准确的服装细节。')
                lines.append('角色脸部身份必须严格参考角色参考图。')
                lines.append('服装款式必须严格参考服装参考图。')
                lines.append('场景细节必须严格参考场景参考图。')
                lines.append('风格质感必须严格参考风格参考图。')
                lines.append('分镜参考图只决定怎么拍，不决定人物长什么样。')
                lines.append('如果分镜参考图是多格分镜板：选择其中最符合当前视频段落的镜头顺序进行参考。不要把多格画面同时塞进同一帧。生成视频时应表现为连续镜头或自然剪辑，而不是拼贴画面。')
                if first_frame_url:
                    lines.append('注意：首帧图优先决定第一帧画面，分镜参考用于约束后续镜头构图和运镜。')
                else:
                    lines.append('注意：分镜参考图用于约束视频开头画面构图。')
            prefix = '\n'.join(lines) + '\n' + QUALITY_PROMPT

            content = [{'type': 'text', 'text': prefix + '\n' + script}]
            for img in ark_reference_images:
                url = img['url']
                if url.startswith('/static/'):
                    url = host_url + url
                content.append({'type':'image_url','image_url':{'url':url},'role':'reference_image'})
            if audio_url and not strict_frame_mode:
                final_audio_url = host_url + audio_url if audio_url.startswith('/static/') else audio_url
                content.append({'type':'audio_url','audio_url':{'url':final_audio_url},'role':'reference_audio'})
            if video_url and not strict_frame_mode:
                content.append({'type':'video_url','video_url':{'url':video_url},'role':'reference_video'})
            if first_frame_url:
                url = first_frame_url
                if url.startswith('/static/'): url = host_url + url
                content.append({'type':'image_url','image_url':{'url':url},'role':'first_frame'})
            if last_frame_url:
                url = last_frame_url
                if url.startswith('/static/'): url = host_url + url
                content.append({'type':'image_url','image_url':{'url':url},'role':'last_frame'})

            try:
                if 'content' in locals() and isinstance(content, list):
                    for item in content:
                        if isinstance(item, dict) and item.get('type') == 'image_url':
                            if 'image_url' in item and 'url' in item['image_url']:
                                orig_url = item['image_url']['url']
                                if orig_url:
                                    item['image_url']['url'] = wash_seedance_image_by_url(orig_url)
            except Exception as wash_err:
                print(f"专属洗图执行出现异常，已降级回原图: {wash_err}")

            client = Ark(api_key=video_cfg['api_key'])
            JOBS[job_id]['status'] = 'running'
            res = client.content_generation.tasks.create(
                model=video_model, content=content,
                generate_audio=True,
                ratio='adaptive' if strict_frame_mode else ratio,
                duration=duration,
                resolution=resolution,
                watermark=False)
            task_id = res.id

            while True:
                r = client.content_generation.tasks.get(task_id=task_id)
                if r.status == 'succeeded':
                    vurl = r.content.video_url
                    stored_vurl, _ = download_and_save_video(vurl)
                    save_video_history(
                        stored_vurl, script,
                        original_script=original_script,
                        refined_script=script if optimize else original_script,
                        model=video_model,
                        ratio=ratio,
                        duration=duration,
                        resolution=resolution,
                        ref_count=len(reference_images),
                        user_id=user_id
                    )
                    JOBS[job_id] = {'status':'succeeded','video_url':stored_vurl,'source_video_url':vurl,'error':None}
                    break
                elif r.status == 'failed':
                    JOBS[job_id] = {'status':'failed','video_url':None,'error':str(r.error)}
                    break
                else:
                    time.sleep(5)
        except Exception as e:
            JOBS[job_id] = {'status':'failed','video_url':None,'error':str(e)}

    start_metered_job(run, (), job_id, user_id, point_cost)
    return jsonify(job_id=job_id)

@app.route('/api/status/<job_id>')
@login_required
def status(job_id):
    if JOB_OWNERS.get(job_id) != current_user_id():
        return jsonify(status='not_found'), 404
    job = JOBS.get(job_id, {'status': 'not_found'})
    return jsonify(job)

@app.route('/api/model-caps')
@login_required
def model_caps():
    return jsonify(MODEL_CAPS)

# ── Styles CRUD ───────────────────────────────────────────────
@app.route('/api/styles', methods=['GET'])
@login_required
def get_styles():
    styles, changed = ensure_default_styles(load_json(styles_path(), []))
    if changed:
        save_json(styles_path(), styles)
    return jsonify(styles)

@app.route('/api/styles', methods=['POST'])
@login_required
def save_styles():
    data = request.json
    if isinstance(data, list):
        save_json(styles_path(), data)
    return jsonify(ok=True)

@app.route('/api/styles/<style_id>', methods=['DELETE'])
@login_required
def delete_style(style_id):
    styles = load_json(styles_path(), [])
    styles = [s for s in styles if s.get('id') != style_id]
    save_json(styles_path(), styles)
    return jsonify(ok=True)

# ── Frame extraction ──────────────────────────────────────────
@app.route('/api/extract-frame', methods=['POST'])
@login_required
def extract_frame():
    body = request.json
    video_url = body.get('video_url', '')
    position = body.get('position', 'first')

    if not video_url:
        return jsonify(error='video_url required'), 400

    import shutil, subprocess
    if not shutil.which('ffmpeg') or not shutil.which('ffprobe'):
        return jsonify(error='未检测到 ffmpeg/ffprobe，无法截帧。Railway 需要重新部署包含 ffmpeg 的版本。'), 500

    if video_url.startswith('/'):
        video_url = request.host_url.rstrip('/') + video_url

    name = uuid.uuid4().hex + '.png'
    out_path = os.path.join(UPLOAD, name)
    tmp_dir = os.path.join(BASE, 'tmp_frames')
    os.makedirs(tmp_dir, exist_ok=True)
    tmp_input = os.path.join(tmp_dir, uuid.uuid4().hex + '.mp4')

    try:
        r = requests.get(video_url, timeout=180, stream=True, headers={'User-Agent': 'Mozilla/5.0'})
        if r.status_code != 200:
            return jsonify(error=f'下载视频失败: {r.status_code}'), 500
        with open(tmp_input, 'wb') as f:
            for chunk in r.iter_content(8192):
                if chunk:
                    f.write(chunk)

        # Use ffmpeg to extract frame
        if position == 'last':
            # Get video duration first, extract from near end
            probe = subprocess.run(['ffprobe', '-v', 'quiet', '-print_format', 'json',
                                    '-show_format', tmp_input],
                                    capture_output=True, text=True, timeout=30)
            duration = 5
            if probe.returncode == 0:
                import json as _json
                info = _json.loads(probe.stdout)
                duration = float(info.get('format', {}).get('duration', 5))
            seek_time = max(0, duration - 2)
            result = subprocess.run(['ffmpeg', '-y', '-ss', str(seek_time), '-i', tmp_input,
                                     '-vframes', '1', '-q:v', '2', out_path],
                                    capture_output=True, text=True, timeout=60)
        else:
            result = subprocess.run(['ffmpeg', '-y', '-i', tmp_input, '-vframes', '1',
                                     '-q:v', '2', out_path],
                                    capture_output=True, text=True, timeout=60)

        if result.returncode != 0:
            err = result.stderr[-500:] if result.stderr else '无输出'
            return jsonify(error=f'截帧失败: {err}'), 500

        if not os.path.exists(out_path):
            return jsonify(error='截帧失败'), 500

        # Upload to TOS
        with open(out_path, 'rb') as f:
            img_bytes = f.read()
        public_url, ok = upload_to_tos(img_bytes, name, 'image/png')
        if ok:
            return jsonify(url=public_url, name=name)
        return jsonify(url=f'/static/uploads/{name}', name=name)

    except Exception as e:
        return jsonify(error=f'截帧异常: {str(e)}'), 500
    finally:
        try:
            if os.path.exists(tmp_input):
                os.remove(tmp_input)
        except Exception:
            pass


def run_upscale_job(job_id, video_url, ratio):
    import subprocess, shutil
    JOBS[job_id] = {'status': 'running', 'url': None, 'error': None, 'progress': '检查 ffmpeg'}

    # Check ffmpeg
    if not shutil.which('ffmpeg') or not shutil.which('ffprobe'):
        JOBS[job_id] = {
            'status': 'failed',
            'url': None,
            'error': '未检测到 ffmpeg/ffprobe。Railway 需要部署包含 ffmpeg 的 nixpacks.toml；本机需要先安装 ffmpeg。'
        }
        return

    # Ratio → target resolution (1080p)
    RATIO_SIZE = {
        '9:16': (1080, 1920),
        '16:9': (1920, 1080),
        '1:1': (1080, 1080),
        '4:3': (1440, 1080),
        '3:4': (1080, 1440),
        '21:9': (1920, 822),
    }
    tw, th = RATIO_SIZE.get(ratio, (1080, 1920))

    tmp_dir = os.path.join(BASE, 'tmp_upscale')
    os.makedirs(tmp_dir, exist_ok=True)
    tmp_input = os.path.join(tmp_dir, uuid.uuid4().hex + '_in.mp4')
    tmp_output = os.path.join(tmp_dir, uuid.uuid4().hex + '_out.mp4')

    try:
        # Download input
        JOBS[job_id]['progress'] = '下载原视频'
        r = requests.get(video_url, timeout=180, stream=True, headers={'User-Agent': 'Mozilla/5.0'})
        if r.status_code != 200:
            JOBS[job_id] = {'status': 'failed', 'url': None, 'error': f'下载视频失败: {r.status_code}'}
            return
        ct = r.headers.get('Content-Type', '')
        total = 0
        with open(tmp_input, 'wb') as f:
            for chunk in r.iter_content(8192):
                if not chunk:
                    continue
                f.write(chunk)
                total += len(chunk)
        if total < 1024:
            JOBS[job_id] = {'status': 'failed', 'url': None, 'error': f'下载的视频太小（{total}字节，Content-Type: {ct}），链接可能已过期'}
            return
        if 'html' in ct.lower() or (total < 1000 and not ct.startswith('video/')):
            JOBS[job_id] = {'status': 'failed', 'url': None, 'error': f'视频链接无效（{total}字节，Content-Type: {ct}），可能是HTML页面'}
            return

        # ffmpeg: lanczos scale + unsharp sharpen + re-encode h264
        JOBS[job_id]['progress'] = 'ffmpeg 转码到 1080p'
        cmd = [
            'ffmpeg', '-y', '-i', tmp_input,
            '-vf', f'scale={tw}:{th}:flags=lanczos,unsharp=5:5:0.5:3:3:0.25',
            '-c:v', 'libx264', '-crf', '18', '-preset', 'medium',
            '-c:a', 'aac', '-b:a', '128k',
            '-movflags', '+faststart',
            tmp_output
        ]
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
        if result.returncode != 0:
            err = result.stderr[-500:] if result.stderr else '无输出'
            JOBS[job_id] = {'status': 'failed', 'url': None, 'error': f'ffmpeg 处理失败: {err}'}
            return

        # Upload to TOS
        JOBS[job_id]['progress'] = '上传 1080p 视频'
        name = uuid.uuid4().hex + '.mp4'
        with open(tmp_output, 'rb') as f:
            vid_bytes = f.read()
        public_url, ok = upload_to_tos(vid_bytes, name, 'video/mp4')
        if ok:
            out_url = public_url
        else:
            out_path = os.path.join(UPLOAD, name)
            shutil.copy(tmp_output, out_path)
            out_url = f'/static/uploads/{name}'

        JOBS[job_id] = {
            'status': 'succeeded',
            'url': out_url,
            'width': tw,
            'height': th,
            'ratio': ratio,
            'target': '1080p',
            'error': None
        }

    except Exception as e:
        JOBS[job_id] = {'status': 'failed', 'url': None, 'error': str(e)}
    finally:
        import shutil as _shutil
        for f in [tmp_input, tmp_output]:
            try: os.remove(f)
            except: pass
        try: _shutil.rmtree(tmp_dir, ignore_errors=True)
        except: pass


# ── Local upscale to 1080p ───────────────────────────────────
@app.route('/api/upscale-local', methods=['POST'])
@login_required
def upscale_local():
    body = request.json
    video_url = body.get('video_url', '')
    ratio = body.get('ratio', '9:16')
    if not video_url:
        return jsonify(error='video_url required'), 400
    if video_url.startswith('/'):
        video_url = request.host_url.rstrip('/') + video_url

    job_id = uuid.uuid4().hex
    JOB_OWNERS[job_id] = current_user_id()
    JOBS[job_id] = {'status': 'pending', 'url': None, 'error': None}
    threading.Thread(target=run_upscale_job, args=(job_id, video_url, ratio), daemon=True).start()
    return jsonify(job_id=job_id)


@app.route('/api/upscale-status/<job_id>')
@login_required
def upscale_status(job_id):
    if JOB_OWNERS.get(job_id) != current_user_id():
        return jsonify(status='not_found'), 404
    return jsonify(JOBS.get(job_id, {'status': 'not_found'}))


# ── Script text model helper ─────────────────────────────────
def call_script_text_model(model_key, system_prompt, user_content, temperature=0.7, max_tokens=4000, api_cfg=None):
    """Use the built-in text model while points remain, otherwise the user's API."""
    cfg = api_cfg or resolve_api('text', builtin_text_api(model_key))
    return call_platform_text(system_prompt, user_content, temperature=temperature, max_tokens=max_tokens, cfg=cfg)


def _extract_json_array(text):
    raw = (text or '').strip()
    if raw.startswith('```'):
        raw = raw.split('\n', 1)[1] if '\n' in raw else raw
    if raw.endswith('```'):
        raw = raw.rsplit('\n', 1)[0]
    start = raw.find('[')
    end = raw.rfind(']')
    if start >= 0 and end > start:
        raw = raw[start:end + 1]
    return raw.strip()


def parse_script_shots_json(raw, script_model, api_cfg=None):
    candidate = _extract_json_array(raw)
    try:
        shots = json.loads(candidate)
    except json.JSONDecodeError:
        repair_prompt = (
            '你是 JSON 修复器。用户会给你一段本应为 JSON 数组的文本，'
            '请只修复语法错误，保留原字段和中文内容，不要改写剧情。'
            '只输出合法 JSON 数组，不要 Markdown，不要解释。'
        )
        fixed = call_script_text_model(script_model, repair_prompt, raw, temperature=0, max_tokens=4000, api_cfg=api_cfg)
        shots = json.loads(_extract_json_array(fixed))
    if not isinstance(shots, list):
        raise ValueError('模型输出不是 JSON 数组')
    return shots


# ── Script workspace endpoints ────────────────────────────────
@app.route('/api/script/import', methods=['POST'])
@login_required
def script_import():
    f = request.files.get('file')
    if not f: return jsonify(error='no file'), 400
    ext = os.path.splitext(secure_filename(f.filename))[1].lower()
    try:
        if ext == '.txt' or ext == '.md':
            text = f.read().decode('utf-8')
        elif ext == '.docx':
            try:
                from docx import Document
                import io
                doc = Document(io.BytesIO(f.read()))
                text = '\n'.join([p.text for p in doc.paragraphs])
            except ImportError:
                return jsonify(error='python-docx 未安装，请 pip install python-docx'), 500
        elif ext == '.pdf':
            return jsonify(error='PDF 暂不支持，请转为 txt 或 md 后上传'), 400
        else:
            return jsonify(error=f'不支持的文件格式: {ext}'), 400
        return jsonify(text=text, length=len(text))
    except Exception as e:
        return jsonify(error=str(e)), 500


@app.route('/api/script/brainstorm', methods=['POST'])
@login_required
@model_access_required('text')
def script_brainstorm():
    body = request.json or {}
    topic = body.get('topic', '').strip()
    script_model = body.get('script_model', SCRIPT_MODEL_DEFAULT)
    force_personal = bool(body.get('use_personal_api')) or script_model == 'personal-api'
    if not force_personal and script_model not in SCRIPT_MODELS:
        return jsonify(error='无效的文本模型'), 400
    if not topic:
        return jsonify(error='请输入主题或关键词'), 400

    style_rule = script_style_rule(body.get('style_id'), target='script', user_id=current_user_id())
    system_prompt = (
        '你是一个专业的短剧编剧。根据用户提供的主题或关键词，'
        '创作一个1-2分钟的短剧片段。输出格式：\n'
        '【场景】xxx\n'
        '【人物】xxx\n'
        '【剧情】xxx\n\n'
        '直接输出剧本文本，不要markdown格式，不要角色列表。\n'
        + style_rule
    )
    point_cost = 0
    try:
        api_cfg = resolve_api(
            'text', builtin_text_api(SCRIPT_MODEL_DEFAULT if force_personal else script_model),
            current_user_id(), force_personal=force_personal,
            profile_id=body.get('api_profile_id'),
            strict_builtin='use_personal_api' in body and not force_personal
        )
        point_cost = reserve_model_points('text', script_model, current_user_id(), personal=force_personal)
        text = call_script_text_model(script_model, system_prompt, topic, temperature=0.8, max_tokens=2000, api_cfg=api_cfg)
        return jsonify(text=text)
    except QuotaError as e:
        return jsonify(error=str(e)), 402
    except Exception as e:
        refund_model_points(current_user_id(), point_cost)
        return jsonify(error=f'生成失败: {str(e)}'), 500


@app.route('/api/script/split', methods=['POST'])
@login_required
@model_access_required('text')
def script_split():
    body = request.json or {}
    script_text = body.get('script', '').strip()
    if not script_text:
        return jsonify(error='请输入剧本文本'), 400

    job_id = uuid.uuid4().hex
    user_id = current_user_id()
    script_model = body.get('script_model', SCRIPT_MODEL_DEFAULT)
    force_personal = bool(body.get('use_personal_api')) or script_model == 'personal-api'
    if not force_personal and script_model not in SCRIPT_MODELS:
        return jsonify(error='无效的文本模型'), 400
    point_cost = 0
    try:
        api_cfg = resolve_api(
            'text', builtin_text_api(SCRIPT_MODEL_DEFAULT if force_personal else script_model),
            user_id, force_personal=force_personal, profile_id=body.get('api_profile_id'),
            strict_builtin='use_personal_api' in body and not force_personal
        )
        point_cost = reserve_model_points('text', script_model, user_id, personal=force_personal)
    except QuotaError as e:
        return jsonify(error=str(e)), 402
    JOB_OWNERS[job_id] = user_id
    SCRIPT_JOBS[job_id] = {
        'status': 'running',
        'created_at': time.time()
    }

    def worker():
        try:
            shots = build_script_shots(body, api_cfg=api_cfg, user_id=user_id)
            SCRIPT_JOBS[job_id].update({
                'status': 'done',
                'shots': shots,
                'count': len(shots),
                'finished_at': time.time()
            })
        except Exception as e:
            refund_model_points(user_id, point_cost)
            SCRIPT_JOBS[job_id].update({
                'status': 'error',
                'error': str(e),
                'finished_at': time.time()
            })

    threading.Thread(target=worker, daemon=True).start()
    return jsonify(job_id=job_id, status='running')


@app.route('/api/script/split/<job_id>', methods=['GET'])
@login_required
def script_split_status(job_id):
    if JOB_OWNERS.get(job_id) != current_user_id():
        return jsonify(error='分镜任务不存在或已过期'), 404
    now = time.time()
    for old_id, job in list(SCRIPT_JOBS.items()):
        if now - job.get('created_at', now) > 3600:
            SCRIPT_JOBS.pop(old_id, None)
    job = SCRIPT_JOBS.get(job_id)
    if not job:
        return jsonify(error='分镜任务不存在或已过期'), 404
    return jsonify(job)


def script_style_rule(style_id, target='video_prompt', user_id=None):
    if not style_id:
        if target == 'script':
            return '用户未选择风格：不要添加真人拍摄、真实摄影、动漫、漫画、3D或其他视觉风格描述。'
        return '【风格规则】用户未选择风格。所有 video_prompt 只写剧情、镜头、动作和技术质量，不得主动添加真人拍摄、真实摄影、动漫、漫画、3D或其他视觉风格词。'

    style = load_prompt_style(style_id, user_id)
    if not style:
        if target == 'script':
            return '所选风格不存在：不要添加任何视觉风格描述。'
        return '【风格规则】所选风格不存在。所有 video_prompt 不得添加任何视觉风格词。'

    name = style.get('name', '用户选择的风格')
    prompt = style.get('prompt', '').strip()
    if target == 'script':
        return f'【用户已选风格：{name}】场景氛围和视觉描述必须遵守以下要求，不得改成其他风格：{prompt}'
    if style_id == LIVE_ACTION_STYLE_ID:
        return (
            f'【用户已选风格：{name}】每个段落必须遵守以下全局风格：{prompt}\n'
            f'【负面词】{style.get("negative_prompt", "").strip()}\n'
            '每个 JSON 对象必须额外输出 timeline 字段，值为字符串数组。'
            '根据该对象的 duration，从00:00开始到结束，每1秒一条，时间必须连续，不得缺口、重叠或超时。'
            '每条严格使用“00:00-00:01｜1秒｜景别，机位/运镜，人物表情、视线、动作、情绪和画面重点”格式。'
            '每秒都要写出具体变化，不得只写“保持”。video_prompt 只写整体连续镜头指令和技术约束，不要重复 timeline。'
        )
    return f'【用户已选风格：{name}】每个 video_prompt 必须遵守以下风格要求，不得添加其他风格：{prompt}'


def format_shot_timeline(shot):
    timeline = shot.get('timeline')
    if not timeline:
        timeline = shot.get('beats') or []
    if isinstance(timeline, str):
        return timeline.strip()
    if not isinstance(timeline, list):
        return ''

    lines = []
    for item in timeline:
        if isinstance(item, str):
            if item.strip():
                lines.append(item.strip())
            continue
        if not isinstance(item, dict):
            continue
        time_range = item.get('time') or item.get('range') or ''
        if not time_range and item.get('start') and item.get('end'):
            time_range = f'{item["start"]}-{item["end"]}'
        seconds = item.get('duration') or item.get('seconds') or ''
        if isinstance(seconds, (int, float)):
            seconds = f'{seconds:g}秒'
        details = '，'.join(str(value).strip() for value in (
            item.get('shot'), item.get('camera'), item.get('action'), item.get('expression'),
            item.get('gaze'), item.get('emotion'), item.get('product_state'),
            item.get('focus'), item.get('continuity')
        ) if value)
        prefix = '｜'.join(part for part in (str(time_range).strip(), str(seconds).strip()) if part)
        line = f'{prefix}｜{details}' if prefix and details else (prefix or details)
        if line:
            lines.append(line)
    return '\n'.join(lines)


def timeline_is_complete(shot):
    try:
        duration = max(1, int(shot.get('duration') or 0))
    except (TypeError, ValueError):
        return False
    timeline = shot.get('timeline')
    if not isinstance(timeline, list) or len(timeline) != duration:
        return False
    required = ('time', 'shot', 'camera', 'action', 'expression', 'focus', 'continuity')
    for second, item in enumerate(timeline):
        if not isinstance(item, dict) or not all(str(item.get(key) or '').strip() for key in required):
            return False
        if item.get('time') != f'00:{second:02d}-00:{second + 1:02d}':
            return False
    return True


def normalize_shot_timeline(shot):
    timeline = shot.get('timeline')
    if not isinstance(timeline, list):
        return shot
    for second, item in enumerate(timeline):
        if isinstance(item, dict):
            item['time'] = f'00:{second:02d}-00:{second + 1:02d}'
            item['duration'] = '1秒'
    return shot


def complete_shot_timeline(shot):
    duration = max(1, min(15, int(shot.get('duration') or 5)))
    original = shot.get('timeline') if isinstance(shot.get('timeline'), list) else []
    scene = str(shot.get('scene') or '当前场景').strip()
    characters = shot.get('characters') or []
    if isinstance(characters, list):
        character = '、'.join(str(item) for item in characters if item) or '画面主体'
    else:
        character = str(characters).strip() or '画面主体'
    story_action = str(shot.get('story_action') or shot.get('action') or '完成当前剧情动作').strip()
    emotion = str(shot.get('emotion') or '克制自然').strip()
    product_state = '产品处于动作涉及的位置，包装正面与手部接触关系清楚，不变形' if '产品' in story_action else '无'
    shots = ('中近景，交代人物与场景关系', '面部特写，突出眼神与嘴部微表情', '近景，兼顾面部和手部动作', '半身近景，呈现身体重心与动作方向')
    cameras = ('平视固定机位，轻微缓慢推近', '略低机位，焦点从环境平稳移到人物眼睛', '侧前方机位，小幅跟随手部与身体位移', '平视机位，镜头缓慢收束到动作终点')
    phases = ('从初始姿态起势，明确动作方向', '动作开始推进，肩颈与手部出现可见位移', '动作进入中段，身体重心随目标方向移动', '完成关键动作，手部接触和物体位置明确', '动作结果显现，面部反应逐步发生变化', '收回多余动作，让姿态和视线自然落定')
    completed = []
    for second in range(duration):
        source = original[second] if second < len(original) else {}
        if isinstance(source, str):
            source = {'action': source.strip()}
        if not isinstance(source, dict):
            source = {}
        phase_index = min(len(phases) - 1, int(second * len(phases) / duration))
        progress = f'第{second + 1}秒：{phases[phase_index]}，围绕“{story_action}”形成明确的起点、过程和落点'
        continuity = (
            '从本段初始姿态进入动作，末端姿势为下一秒的动作起点'
            if second == 0 else
            '准确承接上一秒末端的身体位置、手部位置与视线方向，并把动作推进到下一秒起点'
            if second < duration - 1 else
            '承接上一秒动作轨迹，在本段结束前收束表情、身体重心与物体位置'
        )
        completed.append({
            'time': f'00:{second:02d}-00:{second + 1:02d}',
            'duration': '1秒',
            'shot': str(source.get('shot') or shots[min(len(shots) - 1, int(second * len(shots) / duration))]),
            'camera': str(source.get('camera') or cameras[min(len(cameras) - 1, int(second * len(cameras) / duration))]),
            'action': str(source.get('action') or f'{character}{progress}'),
            'expression': str(source.get('expression') or f'{character}呈现“{emotion}”的渐进微表情，眼睑、眉间和嘴角均有细小变化，面部结构稳定'),
            'gaze': str(source.get('gaze') or '视线先落在当前动作目标上，再随动作进度缓慢移动，眼球方向与头部转动一致'),
            'emotion': str(source.get('emotion') or emotion),
            'product_state': str(source.get('product_state') or product_state),
            'focus': str(source.get('focus') or f'焦点落在{character}的眼神和关键动作上，{scene}作为空间层次，前中后景关系清楚'),
            'continuity': str(source.get('continuity') or continuity),
        })
    shot['duration'] = duration
    shot['timeline'] = completed
    return shot


def ensure_detailed_timelines(shots, script_model, api_cfg=None):
    completed = []
    for original_shot in shots:
        if not isinstance(original_shot, dict):
            raise ValueError('模型返回了无效的分镜对象')
        shot = normalize_shot_timeline(dict(original_shot))
        if timeline_is_complete(shot):
            shot['timeline_text'] = format_shot_timeline(shot)
            completed.append(shot)
            continue
        try:
            duration = max(1, min(15, int(shot.get('duration') or 5)))
        except (TypeError, ValueError):
            duration = 5
        shot['duration'] = duration
        repair_prompt = (
        f'你是严格的逐秒分镜补全器。只输出仅含1个对象的合法 JSON 数组，不要 Markdown，不要解释。'
        f'该段时长固定为{duration}秒，timeline 必须恰好输出{duration}个对象，少一个或多一个都不合格。'
        '保留原有剧情与字段，重写 video_prompt 并补全 timeline。timeline 从00:00开始每1秒一条，'
        '连续、无缺口、无重叠、不得超时。每条必须完整包含：'
        'time（00:00-00:01格式）、duration（固定1秒）、shot（景别与构图）、camera（机位、镜头高度、焦段感和运镜）、'
        'action（这一秒开始到结束的具体动作变化）、expression（面部肌理、眼神、嘴部与微表情变化）、'
        'gaze（视线落点及变化）、emotion（外显与内在情绪）、product_state（产品位置、朝向、可见信息和手部接触；无产品写“无”）、'
        'focus（主体、前中后景、景深与光线重点）、continuity（与上一秒和下一秒如何连续）。'
        '禁止使用“保持”“继续”“同上”“自然动作”等省略描述；每秒都必须可独立执行且具体。'
        '同时重写 video_prompt：先写整体场景、人物、光线、镜头路线与动作目标，再写连续性和质量约束；不要把 timeline 压缩成一句话。'
        )
        for _ in range(1):
            repaired = call_script_text_model(
                script_model, repair_prompt, json.dumps([shot], ensure_ascii=False),
                temperature=0.25, max_tokens=min(12000, max(5000, duration * 850)), api_cfg=api_cfg
            )
            repaired_shots = parse_script_shots_json(repaired, script_model, api_cfg=api_cfg)
            if repaired_shots and isinstance(repaired_shots[0], dict):
                candidate = dict(shot)
                candidate.update(repaired_shots[0])
                candidate['duration'] = duration
                shot = normalize_shot_timeline(candidate)
                if timeline_is_complete(shot):
                    break
        if not timeline_is_complete(shot):
            shot = complete_shot_timeline(shot)
        shot['timeline_text'] = format_shot_timeline(shot)
        completed.append(shot)
    return completed


def apply_live_action_prompt_blocks(shots, style_id, user_id=None):
    if style_id != LIVE_ACTION_STYLE_ID:
        return shots
    style = load_prompt_style(style_id, user_id)
    if not style:
        return shots
    for shot in shots:
        if not isinstance(shot, dict):
            continue
        duration = shot.get('duration') or 5
        timeline_text = format_shot_timeline(shot)
        base_prompt = (shot.get('video_prompt') or shot.get('visual_prompt') or shot.get('prompt') or '').strip()
        if LIVE_ACTION_GLOBAL_MARKER in base_prompt:
            shot['video_prompt'] = base_prompt
            continue
        shot['timeline_text'] = timeline_text
        shot['video_prompt'] = (
            f'{LIVE_ACTION_GLOBAL_MARKER}\n{style.get("prompt", "").strip()}\n'
            f'【负面词】\n{style.get("negative_prompt", "").strip()}\n'
            f'【{duration}秒精准分镜时序脚本】\n{timeline_text}\n'
            f'【连续性与质量约束】\n{base_prompt}'
        ).strip()
    return shots


def build_script_shots(body, api_cfg=None, user_id=None):
    script_text = body.get('script', '').strip()
    mode = body.get('mode', 'smart')  # smart | short | long
    script_model = body.get('script_model', SCRIPT_MODEL_DEFAULT)
    if not script_text:
        raise Exception('请输入剧本文本')
    compact_request = len(script_text) < 120 and mode != 'long'
    duration_match = re.search(r'(?<!\d)([4-9]|1[0-5])\s*(?:秒|s(?:ec(?:onds?)?)?)(?!\w)', script_text, re.IGNORECASE)
    requested_duration = int(duration_match.group(1)) if duration_match else None
    style_rule = script_style_rule(body.get('style_id'), user_id=user_id)

    mode_instructions = {
        'short': '\n当前模式：短镜头模式。每段时长 4-7 秒，逐秒设计景别、动作和情绪变化，适合特写、反应和动作切点；禁止用一个笼统 beat 概括整段。',
        'long': '\n当前模式：长段落模式。每个输出项优先包含 2-3 个 beat，时长 8-15 秒，适合连续剧情推进。',
        'smart': '\n当前模式：智能段落模式（默认）。能合并就合并，该拆才拆，同一场景、同一情绪、同一动作线尽量合成一个段落。'
    }

    if compact_request:
        system_prompt = (
            '你是AI短剧分镜导演。把用户短句扩成1个可直接生成的视频段落，不要过度拆分。\n'
            '只输出 JSON 数组，数组内只放1个对象。字段必须有：segment_no、duration、scene、characters、emotion、story_action、beats、timeline、dialogue、video_prompt。\n'
            + (f'duration 必须严格使用用户指定的{requested_duration}秒。' if requested_duration else 'duration 用6-8秒。')
            + 'timeline 必须恰好等于 duration 条，从00:00开始每1秒一条，连续且不超时。\n'
            'timeline 每条必须完整包含：time、duration（固定1秒）、shot、camera、action、expression、gaze、emotion、product_state、focus、continuity。\n'
            '每秒写清构图、机位高度、焦段感、运镜、人物姿势与位移、手部动作、面部微表情、视线落点、产品位置和前中后景变化。禁止“保持”“继续”“同上”等省略词。\n'
            'video_prompt 必须是详细的直接命令式整体指令，写清场景空间、人物状态、光线、完整镜头路线、动作目标、衔接和质量约束，不得只写一句剧情概述。\n'
            + style_rule + '\n'
            'video_prompt 结尾只加技术质量约束：人物身份稳定，表情自然，动作连贯，手部和肢体正常，服装不穿模，场景保持一致，画面清晰，无字幕，无水印。除上述风格规则外不得自行补充风格词。\n'
            '不要 Markdown，不要解释，只输出 JSON 数组。'
        )
        max_tokens = min(12000, max(6000, (requested_duration or 8) * 750))
    else:
        system_prompt = (
            '你是专业AI短剧分镜导演。用户会提供剧本、剧情梗概或灵感，你要把它拆成"可直接生成的视频段落"，而不是机械拆成单个镜头。\n\n'
            '核心目标：\n'
            '每个输出项是一段可直接交给视频生成模型生成的视频段落，时长 4-15 秒。一个视频段落内部可以包含 1-3 个连续镜头变化或运镜阶段，但必须发生在同一场景、同一时间、同一情绪推进中，保证场景统一、人物连续、动作连贯。\n\n'
            '拆分原则：\n'
            '1. 优先按"剧情动作段落"拆分，不按句子拆分。\n'
            '2. 同一场景、同一人物、同一情绪连续推进的内容，尽量合成一个视频段落。\n'
            '3. 一个视频段落内部允许包含 1-3 个镜头 beat。\n'
            '4. 只有在场景改变、时间跳跃、人物关系改变、情绪爆点需要单独强调、动作无法在一个连续视频里自然完成、需要明显转场时才拆成独立短段落。\n'
            '5. 不要把简单动作拆成多个独立视频段落。不要为了凑数量拆分。\n'
            '6. 每段默认 6-10 秒；动作简单可 4-6 秒；信息密集或包含 2-3 个 beat 可 10-15 秒。\n'
            '7. 每段必须能单独生成，不依赖上一段才能看懂画面。\n'
            '8. 台词要短，符合短剧口语，不能像旁白作文。\n\n'
            '每个视频段落必须包含：segment_no、duration、scene、characters（数组）、emotion、story_action、beats、timeline、dialogue、video_prompt。\n\n'
            'beats 写法：每个 beat 写清楚时间比例（如"前3秒"）、景别、机位/运镜、人物动作、画面重点。\n\n'
            'timeline 写法（强制）：必须恰好等于 duration 条，从00:00开始每1秒一条，连续、无缺口、无重叠、不得超时。每条是对象并完整包含：time、duration（固定1秒）、shot（景别与构图）、camera（机位高度、焦段感和运镜）、action（这一秒内动作的起点、过程、终点）、expression（面部肌理、眼神、嘴部和微表情变化）、gaze（视线落点）、emotion、product_state（产品位置、朝向、可见信息和手部接触，无产品写“无”）、focus（主体、前中后景、景深和光线重点）、continuity（前后秒衔接）。禁止“保持”“继续”“同上”“自然动作”等省略描述。\n\n'
            'video_prompt 写法：\n'
            '1. 必须写成可直接执行的详细连续视频指令，先说明场景空间、人物初始状态和光线，再说明完整镜头路线与动作目标。\n'
            '2. 写清楚机位如何变化、人物如何运动、台词什么时候说。\n'
            '3. 同一场景内多个 beat 要强调"同一连续镜头感"或"自然剪辑感"。\n'
            '4. 保持人物、服装、场景一致。\n'
            '5. 质量约束放在最后：人物身份稳定，表情自然，动作连贯，手部和肢体正常，服装不穿模，场景保持一致，画面清晰，无字幕，无水印。除风格规则外不得自行补充风格词。\n\n'
            + style_rule + '\n\n'
            '如果用户输入很短或没灵感：自动补成一个短剧冲突片段，优先使用强冲突场景：误会、重逢、隐瞒、摊牌、反转、救场、背叛、追问。\n\n'
            + mode_instructions.get(mode, mode_instructions['smart']) + '\n\n'
            '输出格式必须是 JSON 数组，不要 Markdown，不要解释，只输出 JSON 数组。'
        )
        if body.get('style_id') == LIVE_ACTION_STYLE_ID:
            max_tokens = 7000 if len(script_text) < 800 else 10000
        else:
            max_tokens = 6000 if len(script_text) < 800 else 9000
    try:
        raw = call_script_text_model(script_model, system_prompt, script_text, temperature=0.6, max_tokens=max_tokens, api_cfg=api_cfg)
        if raw.startswith('```'):
            raw = raw.split('\n', 1)[1]
        if raw.endswith('```'):
            raw = raw.rsplit('\n', 1)[0]
        shots = parse_script_shots_json(raw, script_model, api_cfg=api_cfg)
        if compact_request and requested_duration and shots:
            shots = shots[:1]
            shots[0]['duration'] = requested_duration
        shots = ensure_detailed_timelines(shots, script_model, api_cfg=api_cfg)
        shots = apply_live_action_prompt_blocks(shots, body.get('style_id'), user_id)
        insert_history({
            'time': datetime.now().strftime('%Y-%m-%d %H:%M'),
            'type': 'script',
            'script_text': script_text[:200],
            'original_script': script_text,
            'segment_count': len(shots),
            'shots': shots
        }, user_id=user_id)
        return shots
    except json.JSONDecodeError as e:
        raise Exception(f'解析分镜结果失败: {str(e)}\n模型原文：{raw[:800]}')
    except Exception as e:
        raise Exception(f'拆分失败: {str(e)}')

threading.Thread(target=migrate_tos_history_videos_to_r2, daemon=True).start()

if __name__ == '__main__':
    port = int(os.environ.get('PORT', '5001'))
    debug = os.environ.get('FLASK_DEBUG', '1') == '1'
    app.run(debug=debug, port=port, threaded=True)
import io
import requests
from PIL import Image, ImageFilter, ImageEnhance

def wash_seedance_image_by_url(url):
    """Pass-through for now, can add image processing later."""
    return url

def wash_tos_image(tos_url):
    """
    黑箱洗图机：输入火山 TOS 链接，在内存里洗完后，上传一个处理后的临时字节流
    """
    try:
        # 1. 从你的 TOS 链接把原图下载到内存里
        response = requests.get(tos_url, timeout=5)
        if response.status_code != 200:
            return tos_url
            
        # 2. 用 PIL 打开图片并进行微观脱敏
        img = Image.open(io.BytesIO(response.content))
        if img.mode == 'RGBA':
            img = img.convert('RGB')
            
        # 轻微模糊，打碎面部特征
        img = img.filter(ImageFilter.GaussianBlur(radius=0.8))
        enhancer = ImageEnhance.Contrast(img)
        img = enhancer.enhance(1.03)
        
        # 3. 重新导出为二进制字节
        output_buffer = io.BytesIO()
        img.save(output_buffer, format='JPEG', quality=85)
        washed_bytes = output_buffer.getvalue()
        
        # 4. 这里需要调用你现有的 upload_to_tos 函数，把洗好的图重新传一份
        # 为了防止文件名冲突，加个前缀
        import os
        from datetime import datetime
        new_name = "washed_" + datetime.now().strftime("%H%M%S") + ".jpg"
        
        # 借用你写好的 upload_to_tos 函数（在图三能看到这个函数定义）
        new_url, ok = upload_to_tos(washed_bytes, new_name, 'image/jpeg')
        if ok:
            print(f"--- 后端：Seedance 专属洗图成功！新链接已生成: {new_url} ---")
            return new_url
            
    except Exception as e:
        print(f"洗图失败，降级使用原链接: {e}")
        
    return tos_url
